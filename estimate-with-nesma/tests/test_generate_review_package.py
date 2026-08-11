import json
import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SCRIPT = Path(__file__).parents[1] / "scripts" / "generate_review_package.py"


def model(method="overview"):
    return {
        "subject": "名片管理系统",
        "method": method,
        "count_type": "application",
        "user": "名片管理操作员",
        "scope": "当前明确需求",
        "boundary": "单机名片管理应用",
        "evidence": [{"id": "E1", "role": "规格", "source": "需求.md", "version": "v1", "location": "原始需求", "fact": "操作员维护并查询联系人", "claims": "D1,T1", "traceable": True}],
        "data_functions": [{"id": "D1", "name": "联系人", "type": "ILF", "change": "UNCHANGED", "ret": 1, "det": 9, "evidence_ids": ["E1"]}],
        "transactions": [{"id": "T1", "name": "录入联系人", "type": "EI", "change": "UNCHANGED", "det": 9, "ftr": 1, "evidence_ids": ["E1"]}],
        "non_counted": [{"requirement": "硬盘尽可能小", "reason": "非功能约束且无阈值"}],
        "assumptions": [],
    }


class GeneratorTest(unittest.TestCase):
    def generate(self, value):
        temp = tempfile.TemporaryDirectory()
        root = Path(temp.name)
        input_path = root / "input.json"
        input_path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(["python", str(SCRIPT), str(input_path), str(root / "out")], text=True, capture_output=True)
        self.addCleanup(temp.cleanup)
        self.assertEqual(result.returncode, 0, result.stderr)
        files = sorted(p.name for p in (root / "out").iterdir())
        self.assertEqual(files, ["估算审阅报告.md", "功能点计数证据.xlsx", "规范化功能清单.md"])
        return root / "out"

    def test_ready_overview_uses_fixed_weights(self):
        out = self.generate(model())
        wb = load_workbook(out / "功能点计数证据.xlsx", data_only=False)
        self.assertEqual(wb["数据功能"]["K2"].value, "=IF(J2,7,0)")
        self.assertEqual(wb["事务功能"]["L2"].value, "=IF(K2,4,0)")
        self.assertIn("待复核", (out / "估算审阅报告.md").read_text())

    def test_ready_detailed_applies_complexity_matrices(self):
        value = model("detailed")
        value["data_functions"][0].update(ret=2, det=30)
        value["transactions"][0].update(det=20, ftr=3)
        out = self.generate(value)
        wb = load_workbook(out / "功能点计数证据.xlsx", data_only=False)
        self.assertEqual(wb["数据功能"]["H2"].value, "中")
        self.assertEqual(wb["数据功能"]["I2"].value, 10)
        self.assertEqual(wb["事务功能"]["I2"].value, "高")
        self.assertEqual(wb["事务功能"]["J2"].value, 6)

    def test_missing_boundary_blocks_and_excludes_rows(self):
        value = model()
        value["boundary"] = ""
        out = self.generate(value)
        wb = load_workbook(out / "功能点计数证据.xlsx", data_only=False)
        self.assertFalse(wb["数据功能"]["J2"].value)
        self.assertNotIn("正式功能点总数", (out / "估算审阅报告.md").read_text())

    def test_untraceable_derived_prd_blocks(self):
        value = model()
        value["evidence"][0].update(role="意图", derived=True, traceable=False)
        out = self.generate(value)
        self.assertIn("派生材料不可追溯", (out / "估算审阅报告.md").read_text())

    def test_evidence_conflict_blocks_affected_rows(self):
        value = model()
        value["evidence"][0]["conflict"] = "联系人由外部系统维护"
        out = self.generate(value)
        wb = load_workbook(out / "功能点计数证据.xlsx", data_only=False)
        self.assertFalse(wb["数据功能"]["J2"].value)
        self.assertIn("证据冲突", (out / "估算审阅报告.md").read_text())

    def test_enhancement_formulas_and_protection(self):
        value = model("detailed")
        value.update(count_type="enhancement", baseline_ufp=100)
        value["data_functions"] = [
            {"id": "D1", "name": "新增", "type": "ILF", "change": "ADD", "ret": 1, "det": 9, "evidence_ids": ["E1"]},
            {"id": "D2", "name": "删除", "type": "ILF", "change": "DEL", "ret": 1, "det": 9, "before_fp": 7, "evidence_ids": ["E1"]},
            {"id": "D3", "name": "修改", "type": "ILF", "change": "CHG", "ret": 2, "det": 30, "before_fp": 7, "evidence_ids": ["E1"]},
        ]
        out = self.generate(value)
        wb = load_workbook(out / "功能点计数证据.xlsx", data_only=False)
        summary = wb["汇总"]
        self.assertEqual(summary["B7"].value, "=B3+B4+B5")
        self.assertEqual(summary["B8"].value, "=B2+B3-B4+(B5-B6)")
        self.assertTrue(wb["数据功能"].protection.sheet)
        self.assertTrue(wb["数据功能"]["K2"].protection.locked)
        self.assertFalse(wb["数据功能"]["L2"].protection.locked)


if __name__ == "__main__":
    unittest.main()
