#!/usr/bin/env python3
"""Generate the fixed four-sheet NESMA package from a confirmed count model."""

import argparse
import json
import unicodedata
from copy import copy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = ("汇总", "数据功能", "事务功能", "证据清单")
DET_SIDES = {"输入", "输出", "启动触发", "消息"}
PROJECT_TYPES = {"新项目", "存量应用"}
METHODS = {"NESMA 详细 FPA", "NESMA 概要 FPA"}
DATA_TYPES = {"ILF", "ELF"}
TRANSACTION_TYPES = {"EI", "EO", "EQ"}
OVERVIEW_WEIGHTS = {"ILF": 7, "ELF": 5, "EI": 4, "EO": 5, "EQ": 4}
WEIGHTS = {
    "ILF": {"低": 7, "中": 10, "高": 15},
    "ELF": {"低": 5, "中": 7, "高": 10},
    "EI": {"低": 3, "中": 4, "高": 6},
    "EO": {"低": 4, "中": 5, "高": 7},
    "EQ": {"低": 3, "中": 4, "高": 6},
}


class ModelError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise ModelError(message)


def read_text(value, label):
    require(isinstance(value, str) and value.strip(), f"{label} 必须是非空字符串")
    return value.strip()


def read_list(value, label):
    require(isinstance(value, list) and value, f"{label} 必须是非空列表")
    return value


def read_names(value, label):
    items = read_list(value, label)
    require(all(isinstance(item, str) and item.strip() for item in items), f"{label} 必须是非空字符串列表")
    require(len(items) == len(set(items)), f"{label} 不得重复")
    return [item.strip() for item in items]


def data_complexity(ret, det):
    row = 0 if ret == 1 else 1 if ret <= 5 else 2
    col = 0 if det <= 19 else 1 if det <= 50 else 2
    return (("低", "低", "中"), ("低", "中", "高"), ("中", "高", "高"))[row][col]


def transaction_complexity(kind, ftr, det):
    if kind == "EI":
        row = 0 if ftr <= 1 else 1 if ftr == 2 else 2
        col = 0 if det <= 4 else 1 if det <= 15 else 2
    else:
        row = 0 if ftr <= 1 else 1 if ftr <= 3 else 2
        col = 0 if det <= 5 else 1 if det <= 19 else 2
    return (("低", "低", "中"), ("低", "中", "高"), ("中", "高", "高"))[row][col]


def check_computed(row, computed, label):
    for key, value in computed.items():
        if key in {"dets_text", "ftrs_text"}:
            row[key] = value
            continue
        if key in row and row[key] != value:
            raise ModelError(f"{label} 的 {key} 应为 {value}，实际为 {row[key]}")
        row[key] = value


def validate_evidence(model):
    evidence = read_list(model.get("evidence"), "evidence")
    ids = []
    by_id = {}
    for index, item in enumerate(evidence, 1):
        require(isinstance(item, dict), f"evidence[{index}] 必须是对象")
        evidence_id = read_text(item.get("id"), f"evidence[{index}].id")
        require(evidence_id not in by_id, f"证据编号重复：{evidence_id}")
        for key in ("role", "source", "version", "location", "fact", "status"):
            read_text(item.get(key), f"{evidence_id}.{key}")
        supports = item.get("supports")
        require(isinstance(supports, list), f"{evidence_id}.supports 必须是列表")
        ids.append(evidence_id)
        by_id[evidence_id] = item
    return ids, by_id


def validate_data_functions(model, evidence_ids, method):
    rows = read_list(model.get("data_functions"), "data_functions")
    ids = []
    for index, row in enumerate(rows, 1):
        require(isinstance(row, dict), f"data_functions[{index}] 必须是对象")
        ident = read_text(row.get("id"), f"data_functions[{index}].id")
        require(ident not in ids, f"数据功能编号重复：{ident}")
        read_text(row.get("module"), f"{ident}.module")
        read_text(row.get("name"), f"{ident}.name")
        require(row.get("type") in DATA_TYPES, f"{ident}.type 必须是 ILF 或 ELF")
        read_text(row.get("confirmation"), f"{ident}.confirmation")
        row_evidence = read_list(row.get("evidence"), f"{ident}.evidence")
        missing = [item for item in row_evidence if item not in evidence_ids]
        require(not missing, f"{ident} 引用不存在的证据：{','.join(missing)}")

        if method == "NESMA 详细 FPA":
            rets = read_names(row.get("rets"), f"{ident}.rets")
            dets = read_names(row.get("dets"), f"{ident}.dets")
            ret = len(rets)
            det = len(dets)
            complexity = data_complexity(ret, det)
            fp = WEIGHTS[row["type"]][complexity]
            computed = {
                "ret_count": ret,
                "det_count": det,
                "complexity": complexity,
                "fp": fp,
            }
        else:
            rets = row.get("rets") or []
            dets = row.get("dets") or []
            require(isinstance(rets, list) and isinstance(dets, list), f"{ident} 的 rets/dets 必须是列表")
            ret = row.get("ret_count")
            det = row.get("det_count")
            require(isinstance(ret, int) and ret >= 1, f"{ident}.ret_count 必须是正整数")
            require(isinstance(det, int) and det >= 1, f"{ident}.det_count 必须是正整数")
            computed = {
                "ret_count": ret,
                "det_count": det,
                "complexity": "概要",
                "fp": OVERVIEW_WEIGHTS[row["type"]],
            }
        check_computed(row, computed, ident)
        ids.append(ident)
    return rows, ids


def validate_transactions(model, evidence_ids, data_ids, method):
    rows = read_list(model.get("transactions"), "transactions")
    ids = []
    for index, row in enumerate(rows, 1):
        require(isinstance(row, dict), f"transactions[{index}] 必须是对象")
        ident = read_text(row.get("id"), f"transactions[{index}].id")
        require(ident not in ids, f"事务功能编号重复：{ident}")
        read_text(row.get("module"), f"{ident}.module")
        read_text(row.get("name"), f"{ident}.name")
        read_text(row.get("processing"), f"{ident}.processing")
        require(row.get("type") in TRANSACTION_TYPES, f"{ident}.type 必须是 EI、EO 或 EQ")
        read_text(row.get("confirmation"), f"{ident}.confirmation")
        row_evidence = read_list(row.get("evidence"), f"{ident}.evidence")
        missing = [item for item in row_evidence if item not in evidence_ids]
        require(not missing, f"{ident} 引用不存在的证据：{','.join(missing)}")

        if method == "NESMA 详细 FPA":
            dets = read_list(row.get("dets"), f"{ident}.dets")
            require(all(isinstance(item, dict) for item in dets), f"{ident}.dets 必须是对象列表")
            names = []
            for item in dets:
                name = read_text(item.get("name"), f"{ident}.dets.name")
                side = item.get("side")
                require(side in DET_SIDES, f"{ident}.dets.side 必须是输入、输出、启动触发或消息")
                names.append((name, side))
            require(len(names) == len(set(names)), f"{ident}.dets 不得重复")
            require(sum(side == "启动触发" for _, side in names) <= 1, f"{ident} 启动触发超过一行")
            require(sum(side == "消息" for _, side in names) <= 1, f"{ident} 消息超过一行")
            require(all(name == "启动触发" for name, side in names if side == "启动触发"), f"{ident} 启动触发行名称必须是启动触发")
            require(all(name == "消息" for name, side in names if side == "消息"), f"{ident} 消息行名称必须是消息")
            ftrs = read_list(row.get("ftrs"), f"{ident}.ftrs")
            unknown = [item for item in ftrs if not isinstance(item, str) or item not in data_ids]
            require(not unknown, f"{ident}.ftrs 引用未登记的数据功能：{', '.join(str(item) for item in unknown)}")
            require(len(ftrs) == len(set(ftrs)), f"{ident}.ftrs 不得重复")
            det = len(dets)
            ftr = len(ftrs)
            complexity = transaction_complexity(row["type"], ftr, det)
            fp = WEIGHTS[row["type"]][complexity]
            computed = {
                "det_count": det,
                "ftr_count": ftr,
                "complexity": complexity,
                "fp": fp,
                "dets_text": "\n".join(f"{side}：{name}" for name, side in names),
                "ftrs_text": ",".join(ftrs),
            }
        else:
            det = row.get("det_count")
            ftr = row.get("ftr_count")
            require(isinstance(det, int) and det >= 1, f"{ident}.det_count 必须是正整数")
            require(isinstance(ftr, int) and ftr >= 0, f"{ident}.ftr_count 必须是非负整数")
            computed = {
                "det_count": det,
                "ftr_count": ftr,
                "complexity": "概要",
                "fp": OVERVIEW_WEIGHTS[row["type"]],
                "dets_text": row.get("dets_text", ""),
                "ftrs_text": row.get("ftrs_text", ""),
            }
        check_computed(row, computed, ident)
        ids.append(ident)
    return rows


def validate_model(model):
    require(isinstance(model, dict), "计数模型必须是 JSON 对象")
    require("assumptions" not in model, "最终计数模型不得包含 assumptions")
    require("unresolved" not in model, "最终计数模型不得包含 unresolved")
    package = model.get("package")
    require(isinstance(package, dict), "package 必须是对象")
    for key in ("name", "version", "output_directory", "status"):
        read_text(package.get(key), f"package.{key}")
    counting = model.get("counting")
    require(isinstance(counting, dict), "counting 必须是对象")
    require(counting.get("project_type") in PROJECT_TYPES, "计数类型必须是新项目或存量应用")
    require(counting.get("method") in METHODS, "计数方法必须是 NESMA 详细 FPA 或 NESMA 概要 FPA")
    for key in ("application_boundary", "scope"):
        read_text(counting.get(key), f"counting.{key}")
    require(isinstance(counting.get("users"), list) and counting["users"], "counting.users 必须是非空列表")
    confirmed = model.get("confirmed_decisions", [])
    require(isinstance(confirmed, list) and all(isinstance(item, str) and item.strip() for item in confirmed), "confirmed_decisions 必须是非空字符串列表")
    excluded = model.get("excluded", [])
    require(isinstance(excluded, list), "excluded 必须是列表")
    evidence_ids, evidence_by_id = validate_evidence(model)
    data_rows, data_ids = validate_data_functions(model, evidence_ids, counting["method"])
    transaction_rows = validate_transactions(model, evidence_ids, data_ids, counting["method"])
    for item in excluded:
        require(isinstance(item, dict), "excluded 项必须是对象")
        read_text(item.get("item"), "excluded.item")
        read_text(item.get("reason"), "excluded.reason")
        item_evidence = item.get("evidence", [])
        require(isinstance(item_evidence, list), "excluded.evidence 必须是列表")
        require(all(evidence in evidence_by_id for evidence in item_evidence), "excluded 引用不存在的证据")
    require(len(set(data_ids + [row["id"] for row in transaction_rows])) == len(data_ids) + len(transaction_rows), "数据功能和事务功能编号不得重复")
    return data_rows, transaction_rows


def set_widths(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    for index, cells in enumerate(ws.iter_cols(), 1):
        width = 0
        for cell in cells:
            if cell.value is None:
                continue
            visual = max(
                (
                    sum(2 if unicodedata.east_asian_width(char) in "WFA" else 1 for char in line)
                    for line in str(cell.value).splitlines()
                ),
                default=0,
            )
            width = max(width, visual + 2)
            if visual + 2 > 46:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        ws.column_dimensions[get_column_letter(index)].width = min(width, 46)


def build_workbook(model, data_rows, transaction_rows, output):
    counts = {kind: sum(row["type"] == kind for row in data_rows + transaction_rows) for kind in ("ILF", "ELF", "EI", "EO", "EQ")}
    data_fp = sum(row["fp"] for row in data_rows)
    transaction_fp = sum(row["fp"] for row in transaction_rows)
    evidence_ids = [item["id"] for item in model["evidence"]]
    evidence_ref = ",".join(evidence_ids[:1])
    excluded = "；".join(item["item"] for item in model.get("excluded", [])) or "无"

    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("汇总")
    summary.append(["项目", "值", "说明/公式", "证据编号"])
    rows = [
        ("系统", model["package"]["name"], "计数对象。", evidence_ref),
        ("产物版本", model["package"]["version"], "阶段一固定产物版本。", evidence_ref),
        ("后端基线", model.get("baseline", {}).get("backend_commit", ""), model.get("baseline", {}).get("backend_tag", ""), evidence_ref),
        ("前端基线", model.get("baseline", {}).get("frontend_commit", ""), model.get("baseline", {}).get("frontend_ref", ""), evidence_ref),
        ("计数类型", model["counting"]["project_type"], model["counting"].get("change_type", ""), evidence_ref),
        ("计数方法", model["counting"]["method"], "由 RET/DET、FTR 和带侧 DET 明细计算。" if model["counting"]["method"] == "NESMA 详细 FPA" else "使用概要 FPA 固定权重。", evidence_ref),
        ("功能用户", "；".join(model["counting"]["users"]), "应用边界内用户。", evidence_ref),
        ("应用边界", model["counting"]["application_boundary"], "应用与用户的概念接口。", evidence_ref),
        ("计数范围", model["counting"]["scope"], "纳入本次计数的功能集合。", evidence_ref),
        ("ILF 数", counts["ILF"], "明细复算。", "明细复算"),
        ("ELF 数", counts["ELF"], "明细复算。", "明细复算"),
        ("EI 数", counts["EI"], "明细复算。", "明细复算"),
        ("EO 数", counts["EO"], "明细复算。", "明细复算"),
        ("EQ 数", counts["EQ"], "明细复算。", "明细复算"),
        ("数据功能 FP", data_fp, "数据功能 FP 合计。", "明细复算"),
        ("事务功能 FP", transaction_fp, "事务功能 FP 合计。", "明细复算"),
        ("UFP 总数", data_fp + transaction_fp, "数据功能 FP + 事务功能 FP。", "明细复算"),
        ("排除项", excluded, "不进入 UFP。", "明细复算"),
        ("证据覆盖率", "100%", "全部正式行引用有效证据。", evidence_ref),
        ("包状态", "完成", "零未决项。", evidence_ref),
    ]
    for row in rows:
        summary.append(row)
    set_widths(summary)

    data = wb.create_sheet("数据功能")
    data.append(["编号", "模块", "功能名称", "类型", "RET 数", "DET 数", "复杂度", "FP", "RET 明细", "DET 明细", "证据编号", "用户确认记录"])
    for row in data_rows:
        data.append([
            row["id"], row["module"], row["name"], row["type"], row["ret_count"], row["det_count"],
            row["complexity"], row["fp"], "\n".join(row.get("rets", [])), "\n".join(row.get("dets", [])),
            ",".join(row["evidence"]), row["confirmation"],
        ])
    set_widths(data)

    transaction = wb.create_sheet("事务功能")
    transaction.append(["编号", "模块", "功能名称", "类型", "DET 数", "FTR 数", "复杂度", "FP", "处理说明", "DET 明细（带侧）", "FTR 引用", "证据编号", "用户确认记录"])
    for row in transaction_rows:
        transaction.append([
            row["id"], row["module"], row["name"], row["type"], row["det_count"], row["ftr_count"],
            row["complexity"], row["fp"], row["processing"], row["dets_text"], row["ftrs_text"],
            ",".join(row["evidence"]), row["confirmation"],
        ])
    set_widths(transaction)

    evidence = wb.create_sheet("证据清单")
    evidence.append(["证据编号", "角色", "来源路径或接口", "版本/commit/时间", "精确定位", "观察事实", "支持的数据/事务编号", "状态", "排除理由"])
    for item in model["evidence"]:
        evidence.append([
            item["id"], item["role"], item["source"], item["version"], item["location"], item["fact"],
            "；".join(item["supports"]), item["status"], item.get("exclusion", ""),
        ])
    set_widths(evidence)
    wb.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("model", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    model = json.loads(args.model.read_text(encoding="utf-8"))
    data_rows, transaction_rows = validate_model(model)
    if args.output.exists() and (not args.output.is_dir() or any(args.output.iterdir())):
        parser.error("output directory is not empty; write a new version directory")
    if args.output.resolve() == args.model.parent.resolve():
        parser.error("output directory must differ from the input model directory")
    args.output.mkdir(parents=True, exist_ok=True)
    model["package"]["output_directory"] = str(args.output.resolve())
    (args.output / "计数模型.json").write_text(json.dumps(model, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    build_workbook(model, data_rows, transaction_rows, args.output / "功能点计数证据.xlsx")
    print(f"Generated {args.output / '计数模型.json'}")
    print(f"Generated {args.output / '功能点计数证据.xlsx'}")


if __name__ == "__main__":
    main()
