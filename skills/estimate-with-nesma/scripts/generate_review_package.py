#!/usr/bin/env python3
"""Generate the three-file NESMA estimation review package from transient JSON."""

import argparse
import json
import unicodedata
from copy import copy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter


SCHEMA_VERSION = "1"
GENERATOR_VERSION = "1"
OVERVIEW_WEIGHTS = {"ILF": 7, "ELF": 5, "EI": 4, "EO": 5, "EQ": 4}
WEIGHTS = {
    "ILF": {"低": 7, "中": 10, "高": 15},
    "ELF": {"低": 5, "中": 7, "高": 10},
    "EI": {"低": 3, "中": 4, "高": 6},
    "EO": {"低": 4, "中": 5, "高": 7},
    "EQ": {"低": 3, "中": 4, "高": 6},
}


def evidence_is_usable(evidence, evidence_id):
    item = evidence.get(evidence_id, {})
    return bool(item) and not item.get("conflict") and not (item.get("derived") and not item.get("traceable"))


def data_complexity(ret, det):
    row = 0 if ret == 1 else 1 if ret <= 5 else 2
    col = 0 if det <= 19 else 1 if det <= 50 else 2
    return (("低", "低", "中"), ("低", "中", "高"), ("中", "高", "高"))[row][col]


def transaction_complexity(kind, ftr, det):
    if kind == "EI":
        row, col = (0 if ftr <= 1 else 1 if ftr == 2 else 2), (0 if det <= 4 else 1 if det <= 15 else 2)
    else:
        row, col = (0 if ftr <= 1 else 1 if ftr <= 3 else 2), (0 if det <= 5 else 1 if det <= 19 else 2)
    return (("低", "低", "中"), ("低", "中", "高"), ("中", "高", "高"))[row][col]


def readiness(model):
    reasons = []
    for field, label in (("user", "用户"), ("scope", "计数范围"), ("boundary", "应用边界")):
        if not model.get(field):
            reasons.append(f"缺少{label}")
    evidence = {e.get("id"): e for e in model.get("evidence", [])}
    rows = model.get("data_functions", []) + model.get("transactions", [])
    formal_evidence_ids = {i for row in rows for i in row.get("evidence_ids", []) + (row.get("detail_evidence_ids", []) if row.get("change") != "DEL" else []) + (row.get("before_evidence_ids", []) if row.get("change") == "DEL" else [])}
    for item in model.get("evidence", []):
        if item.get("id") in formal_evidence_ids and item.get("derived") and not item.get("traceable"):
            reasons.append(f"{item.get('id', '未知证据')}：派生材料不可追溯")
        if item.get("id") in formal_evidence_ids and item.get("conflict"):
            reasons.append(f"{item.get('id', '未知证据')}：证据冲突：{item['conflict']}")
    for row in rows:
        if not row.get("evidence_ids") or any(i not in evidence for i in row.get("evidence_ids", [])):
            reasons.append(f"{row.get('id', '未知项')}：缺少可定位证据")
    if model.get("method") == "detailed":
        for row in model.get("data_functions", []):
            if row.get("change") != "DEL" and (not isinstance(row.get("ret"), int) or not isinstance(row.get("det"), int)):
                reasons.append(f"{row.get('id', '未知项')}：缺少 RET/DET")
            if row.get("change") != "DEL" and (not row.get("detail_evidence_ids") or any(i not in evidence for i in row.get("detail_evidence_ids", []))):
                reasons.append(f"{row.get('id', '未知项')}：缺少 RET/DET 证据")
        for row in model.get("transactions", []):
            if row.get("change") != "DEL" and (not isinstance(row.get("det"), int) or not isinstance(row.get("ftr"), int)):
                reasons.append(f"{row.get('id', '未知项')}：缺少 DET/FTR")
            if row.get("change") != "DEL" and (not row.get("detail_evidence_ids") or any(i not in evidence for i in row.get("detail_evidence_ids", []))):
                reasons.append(f"{row.get('id', '未知项')}：缺少 DET/FTR 证据")
    if model.get("count_type") == "enhancement":
        for row in rows:
            if row.get("change") == "DEL" and (not isinstance(row.get("before_fp"), (int, float)) or not row.get("before_evidence_ids") or any(i not in evidence for i in row.get("before_evidence_ids", []))):
                reasons.append(f"{row.get('id', '未知项')}：缺少变更前规模证据")
    return list(dict.fromkeys(reasons))


def row_is_counted(row, evidence, globally_blocked, method, count_type):
    ids = row.get("evidence_ids", [])
    if globally_blocked or not ids or any(not evidence_is_usable(evidence, i) for i in ids):
        return False
    if method == "detailed" and row.get("change") != "DEL":
        counts = (row.get("ret"), row.get("det")) if row.get("type") in ("ILF", "ELF") else (row.get("det"), row.get("ftr"))
        if any(not isinstance(value, int) or isinstance(value, bool) for value in counts) or not row.get("detail_evidence_ids") or any(not evidence_is_usable(evidence, i) for i in row["detail_evidence_ids"]):
            return False
    if count_type == "enhancement" and row.get("change") == "DEL" and (not isinstance(row.get("before_fp"), (int, float)) or isinstance(row.get("before_fp"), bool) or not row.get("before_evidence_ids") or any(not evidence_is_usable(evidence, i) for i in row["before_evidence_ids"])):
        return False
    return True


def style_sheet(ws, review_columns=()):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    for column in review_columns:
        for cell in ws[column][1:]:
            cell.protection = Protection(locked=False)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for index, cells in enumerate(ws.iter_cols(), 1):
        width = 0
        for cell in cells:
            if cell.value is None or (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            visual_width = max((sum(2 if unicodedata.east_asian_width(char) in "WFA" else 1 for char in line) for line in str(cell.value).splitlines()), default=0)
            width = max(width, visual_width + 2)
            if visual_width + 2 > 50:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        ws.column_dimensions[get_column_letter(index)].width = min(width, 50)
    ws.protection.sheet = True
    ws.protection.enable()


def build_workbook(model, reasons, output):
    wb = Workbook()
    wb.remove(wb.active)
    evidence = {e.get("id"): e for e in model.get("evidence", [])}
    blocked = bool(reasons)
    globally_blocked = any(not model.get(field) for field in ("user", "scope", "boundary"))

    info = wb.create_sheet("计数说明")
    info.append(["字段", "值"])
    for pair in (("Schema 版本", SCHEMA_VERSION), ("Generator 版本", GENERATOR_VERSION), ("审阅包版本", model.get("package_version", 1)), ("计数对象", model.get("subject", "")), ("方法", model.get("method", "")), ("状态", "计数阻断" if blocked else "待复核"), ("用户", model.get("user", "")), ("范围", model.get("scope", "")), ("边界", model.get("boundary", ""))):
        info.append(pair)
    style_sheet(info)

    ev = wb.create_sheet("证据登记")
    ev.append(["证据编号", "证据角色", "来源", "版本/实测时间", "具体定位", "观察事实", "支持主张", "派生材料", "可追溯", "冲突/限制"])
    for e in model.get("evidence", []):
        ev.append([e.get(k, "") for k in ("id", "role", "source", "version", "location", "fact", "claims", "derived", "traceable", "conflict")])
    style_sheet(ev)

    data = wb.create_sheet("数据功能")
    data.append(["编号", "名称", "类型", "变更类型", "RET", "DET", "证据编号", "复杂度", "变更后FP", "纳入公式", "计数FP", "输入事实复核", "方法复核", "复核意见", "变更前FP", "详细计数证据", "变更前证据"])
    for index, row in enumerate(model.get("data_functions", []), 2):
        complexity = data_complexity(row["ret"], row["det"]) if isinstance(row.get("ret"), int) and isinstance(row.get("det"), int) else "未决"
        fp = OVERVIEW_WEIGHTS.get(row.get("type")) if model["method"] == "overview" else WEIGHTS.get(row.get("type"), {}).get(complexity)
        included = row_is_counted(row, evidence, globally_blocked, model["method"], model["count_type"])
        data.append([row.get("id"), row.get("name"), row.get("type"), row.get("change", "UNCHANGED"), row.get("ret"), row.get("det"), ",".join(row.get("evidence_ids", [])), complexity, fp, included, f"=IF(J{index},{fp or 0},0)", "待复核", "待复核", "", row.get("before_fp"), ",".join(row.get("detail_evidence_ids", [])), ",".join(row.get("before_evidence_ids", []))])
    style_sheet(data, ("L", "M", "N"))

    tx = wb.create_sheet("事务功能")
    tx.append(["编号", "名称", "类型", "变更类型", "DET", "FTR", "证据编号", "主要处理", "复杂度", "变更后FP", "纳入公式", "计数FP", "输入事实复核", "方法复核", "复核意见", "变更前FP", "详细计数证据", "变更前证据"])
    for index, row in enumerate(model.get("transactions", []), 2):
        complexity = transaction_complexity(row.get("type"), row["ftr"], row["det"]) if isinstance(row.get("ftr"), int) and isinstance(row.get("det"), int) else "未决"
        fp = OVERVIEW_WEIGHTS.get(row.get("type")) if model["method"] == "overview" else WEIGHTS.get(row.get("type"), {}).get(complexity)
        included = row_is_counted(row, evidence, globally_blocked, model["method"], model["count_type"])
        tx.append([row.get("id"), row.get("name"), row.get("type"), row.get("change", "UNCHANGED"), row.get("det"), row.get("ftr"), ",".join(row.get("evidence_ids", [])), row.get("processing", ""), complexity, fp, included, f"=IF(K{index},{fp or 0},0)", "待复核", "待复核", "", row.get("before_fp"), ",".join(row.get("detail_evidence_ids", [])), ",".join(row.get("before_evidence_ids", []))])
    style_sheet(tx, ("M", "N", "O"))

    issues = wb.create_sheet("假设与冲突")
    issues.append(["类型", "内容", "状态", "复核意见"])
    for reason in reasons:
        issues.append(["阻断", reason, "未解决", ""])
    for assumption in model.get("assumptions", []):
        issues.append(["假设", assumption, "待确认", ""])
    for item in model.get("non_counted", []):
        issues.append(["非计数需求", f"{item.get('requirement', '')}：{item.get('reason', '')}", "已排除", ""])
    style_sheet(issues, ("C", "D"))

    summary = wb.create_sheet("汇总")
    summary.append(["项目", "公式/值", "证据编号"])
    summary.append(["增强前应用 UFP", model.get("baseline_ufp", ""), ",".join(model.get("baseline_evidence_ids", []))])
    summary.append(["ADD", '=SUMIF(数据功能!D:D,"ADD",数据功能!K:K)+SUMIF(事务功能!D:D,"ADD",事务功能!L:L)'])
    summary.append(["DEL", '=SUMIFS(数据功能!O:O,数据功能!D:D,"DEL",数据功能!J:J,TRUE)+SUMIFS(事务功能!P:P,事务功能!D:D,"DEL",事务功能!K:K,TRUE)'])
    summary.append(["CHGA", '=SUMIF(数据功能!D:D,"CHG",数据功能!K:K)+SUMIF(事务功能!D:D,"CHG",事务功能!L:L)'])
    summary.append(["CHGB", '=SUMIFS(数据功能!O:O,数据功能!D:D,"CHG",数据功能!J:J,TRUE)+SUMIFS(事务功能!P:P,事务功能!D:D,"CHG",事务功能!K:K,TRUE)'])
    summary.append(["EFP", "=B3+B4+B5" if model.get("count_type") == "enhancement" and not blocked else "阻断"])
    before_rows = [r for r in model.get("data_functions", []) + model.get("transactions", []) if r.get("change") in ("DEL", "CHG")]
    has_before_evidence = all(isinstance(r.get("before_fp"), (int, float)) and not isinstance(r.get("before_fp"), bool) and r.get("before_evidence_ids") and all(evidence_is_usable(evidence, i) for i in r["before_evidence_ids"]) for r in before_rows)
    has_baseline_evidence = model.get("baseline_evidence_ids") and all(evidence_is_usable(evidence, i) for i in model["baseline_evidence_ids"])
    summary.append(["修改后应用 UFP", "=B2+B3-B4+(B5-B6)" if model.get("count_type") == "enhancement" and not blocked and model.get("baseline_ufp") is not None and has_baseline_evidence and has_before_evidence else "证据不足"])
    summary.append(["UFP", "不适用于增强项目" if model.get("count_type") == "enhancement" else "=SUM(数据功能!K:K)+SUM(事务功能!L:L)" if not blocked else "阻断"])
    style_sheet(summary)

    review = wb.create_sheet("复核记录")
    review.append(["版本", "复核类型", "复核人", "结论", "意见", "日期"])
    version = model.get("package_version", 1)
    review.append([version, "规则/公式检查", "生成器", "已执行", "仅证明计算规则，不代替输入事实或方法复核", ""])
    review.append([version, "输入事实复核", "", "待复核", "", ""])
    review.append([version, "NESMA/FPA 方法复核", "", "待复核", "", ""])
    style_sheet(review, ("C", "D", "E", "F"))
    wb.save(output)


def write_markdown(model, reasons, directory):
    status = "计数阻断" if reasons else "待复核"
    report = [f"# {model.get('subject', '未命名计数对象')}估算审阅报告", "", f"- 状态：{status}", f"- 方法：{'详细 FPA' if model.get('method') == 'detailed' else '概要 FPA'}", f"- 用户：{model.get('user') or '未提供'}", f"- 计数范围：{model.get('scope') or '未提供'}", f"- 应用边界：{model.get('boundary') or '未提供'}", ""]
    if reasons:
        report += ["## 计数阻断", ""] + [f"- {r}" for r in reasons]
    else:
        report += ["## 正式功能点总数", "", "以 `功能点计数证据.xlsx` 的受控公式为准；结果仍须输入事实复核与 NESMA/FPA 方法复核。"]
    report += ["", "## 复核分工", "", "- 规则/公式检查：生成器已执行。", "- 输入事实复核：待熟悉业务的用户完成。", "- NESMA/FPA 方法复核：待具备方法能力的审阅者完成。", "", "> 本包不是认证结论或可直接交付客户的终稿。", ""]
    (directory / "估算审阅报告.md").write_text("\n".join(report), encoding="utf-8")

    inventory = [f"# {model.get('subject', '未命名计数对象')}规范化功能清单", "", "## 数据功能", "", "| 编号 | 功能 | 类型 | 证据 |", "| --- | --- | --- | --- |"]
    inventory += [f"| {r.get('id','')} | {r.get('name','')} | {r.get('type','未决')} | {','.join(r.get('evidence_ids', []))} |" for r in model.get("data_functions", [])]
    inventory += ["", "## 事务功能", "", "| 编号 | 基本过程 | 类型 | 证据 |", "| --- | --- | --- | --- |"]
    inventory += [f"| {r.get('id','')} | {r.get('name','')} | {r.get('type','未决')} | {','.join(r.get('evidence_ids', []))} |" for r in model.get("transactions", [])]
    inventory += ["", "## 非计数需求", "", "| 需求 | 排除理由 |", "| --- | --- |"]
    inventory += [f"| {r.get('requirement','')} | {r.get('reason','')} |" for r in model.get("non_counted", [])]
    (directory / "规范化功能清单.md").write_text("\n".join(inventory) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    model = json.loads(args.input.read_text(encoding="utf-8"))
    if model.get("method") not in ("overview", "detailed"):
        parser.error("method must be overview or detailed")
    if model.get("count_type") not in ("application", "new", "enhancement"):
        parser.error("count_type must be application, new, or enhancement")
    for row in model.get("data_functions", []):
        if row.get("type") not in ("ILF", "ELF"):
            parser.error("data function type must be ILF or ELF")
        if any(value is not None and (not isinstance(value, int) or isinstance(value, bool) or value < 1) for value in (row.get("ret"), row.get("det"))):
            parser.error("data function RET/DET must be positive integers")
    for row in model.get("transactions", []):
        if row.get("type") not in ("EI", "EO", "EQ"):
            parser.error("transaction function type must be EI, EO, or EQ")
        if row.get("det") is not None and (not isinstance(row["det"], int) or isinstance(row["det"], bool) or row["det"] < 1):
            parser.error("transaction DET must be a positive integer")
        if row.get("ftr") is not None and (not isinstance(row["ftr"], int) or isinstance(row["ftr"], bool) or row["ftr"] < 0):
            parser.error("transaction FTR must be a non-negative integer")
    for row in model.get("data_functions", []) + model.get("transactions", []):
        if row.get("change", "UNCHANGED") not in ("ADD", "DEL", "CHG", "UNCHANGED"):
            parser.error("change must be ADD, DEL, CHG, or UNCHANGED")
        before_fp = row.get("before_fp")
        if before_fp is not None and (isinstance(before_fp, bool) or not isinstance(before_fp, (int, float)) or before_fp <= 0):
            parser.error("before_fp must be a positive number")
    baseline = model.get("baseline_ufp")
    if baseline is not None and (isinstance(baseline, bool) or not isinstance(baseline, (int, float)) or baseline < 0):
        parser.error("baseline_ufp must be a non-negative number")
    if args.output.exists() and any(args.output.iterdir()):
        parser.error("output directory is not empty; create a new version directory")
    args.output.mkdir(parents=True, exist_ok=True)
    reasons = readiness(model)
    build_workbook(model, reasons, args.output / "功能点计数证据.xlsx")
    write_markdown(model, reasons, args.output)


if __name__ == "__main__":
    main()
