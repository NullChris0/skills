import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


REPO = Path(__file__).parents[3]
SCRIPT = Path(__file__).parents[1] / "scripts" / "generate_count_package.py"
FIXTURE = REPO / "tests/fixtures/nesma-employee-assessment-v3/count/计数模型.json"


class CountPackageTest(unittest.TestCase):
    def model(self):
        return json.loads(FIXTURE.read_text(encoding="utf-8"))

    def run_script(self, value):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        model_path = root / "input.json"
        model_path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        output = root / "out"
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(model_path), str(output)],
            text=True,
            capture_output=True,
        )
        return result, output

    def test_golden_fixture_produces_four_sheets_and_267_ufp(self):
        result, output = self.run_script(self.model())
        self.assertEqual(result.returncode, 0, result.stderr)
        workbook = load_workbook(output / "功能点计数证据.xlsx", data_only=True)
        self.assertEqual(workbook.sheetnames, ["汇总", "数据功能", "事务功能", "证据清单"])
        summary = {row[0]: row[1] for row in workbook["汇总"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(summary["ILF 数"], 11)
        self.assertEqual(summary["ELF 数"], 2)
        self.assertEqual(summary["EI 数"], 22)
        self.assertEqual(summary["EO 数"], 10)
        self.assertEqual(summary["EQ 数"], 5)
        self.assertEqual(summary["数据功能 FP"], 87)
        self.assertEqual(summary["事务功能 FP"], 180)
        self.assertEqual(summary["UFP 总数"], 267)
        self.assertEqual(workbook["数据功能"].max_row - 1, 13)
        self.assertEqual(workbook["事务功能"].max_row - 1, 37)
        self.assertEqual(workbook["证据清单"].max_row - 1, 27)
        generated = json.loads((output / "计数模型.json").read_text(encoding="utf-8"))
        self.assertNotIn("assumptions", generated)
        self.assertEqual(generated["package"]["output_directory"], str(output.resolve()))

    def test_assumptions_are_rejected(self):
        value = self.model()
        value["assumptions"] = ["待确认"]
        result, _ = self.run_script(value)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("assumptions", result.stderr)

    def test_unknown_ftr_is_rejected(self):
        value = self.model()
        value["transactions"][0]["ftrs"] = ["DF-999"]
        result, _ = self.run_script(value)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("DF-999", result.stderr)

    def test_wrong_fp_is_rejected(self):
        value = self.model()
        value["data_functions"][0]["fp"] = 99
        result, _ = self.run_script(value)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("fp", result.stderr)

    def test_missing_confirmation_is_rejected(self):
        value = self.model()
        value["transactions"][0]["confirmation"] = ""
        result, _ = self.run_script(value)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("confirmation", result.stderr)

    def test_nonempty_output_directory_is_rejected(self):
        value = self.model()
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        model_path = root / "input.json"
        model_path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        output = root / "out"
        output.mkdir()
        (output / "keep").write_text("keep", encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(model_path), str(output)],
            text=True,
            capture_output=True,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual((output / "keep").read_text(encoding="utf-8"), "keep")


if __name__ == "__main__":
    unittest.main()
