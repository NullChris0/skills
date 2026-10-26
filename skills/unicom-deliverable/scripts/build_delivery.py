#!/usr/bin/env python3
"""Build a delivery from a stage-two model."""

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from copy import copy
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from delivery_quality import (
    REFERENCE_DOCUMENTS,
    check_document_model,
    check_feature_granularity,
    compact,
    descendants,
)


SKILL = Path(__file__).parents[1]
TEMPLATE = SKILL / "assets/联通功能清单模板.xlsx"
STYLE = SKILL / "assets/api-card.sty"
DOCX = SKILL / "scripts/docx"
RENDER_CARDS = SKILL / "scripts/render_cards.py"
VALIDATE = SKILL / "scripts/validate_delivery.py"
UPDATE_TOC = SKILL / "scripts/update_toc_pages.py"
MAPPING_HEADERS = ["NESMA编号", "类型", "名称", "FP", "对应功能清单行", "对应正文章节/页面/数据模型", "对应接口YAML/卡片", "覆盖状态", "说明"]


def require(condition, message):
    if not condition:
        raise SystemExit(f"FAIL: {message}")


def count_rows(model):
    data = model.get("data_functions")
    transactions = model.get("transactions")
    require(isinstance(data, list) and data, "计数模型缺少 data_functions")
    require(isinstance(transactions, list) and transactions, "计数模型缺少 transactions")
    return data + transactions, data, transactions


def count_summary(rows):
    return {
        "ilf": sum(row["type"] == "ILF" for row in rows),
        "elf": sum(row["type"] == "ELF" for row in rows),
        "ei": sum(row["type"] == "EI" for row in rows),
        "eo": sum(row["type"] == "EO" for row in rows),
        "eq": sum(row["type"] == "EQ" for row in rows),
        "ufp": sum(row["fp"] for row in rows),
    }


def section_by_prefix(sections, prefix):
    for section in sections:
        if compact(section.get("heading")).startswith(compact(prefix)):
            return section
    return None


def strip_direct_content(section):
    section["paras"] = []
    section["tables"] = []
    section["images"] = []
    section["meta"] = []


def normalize_fixed_sections(document, sections):
    title = compact(document.get("title"))
    date_match = re.search(r"(\d{4})年(\d{1,2})月", str(document.get("date", "")))
    if date_match:
        document["date"] = f"{date_match.group(1)}年{date_match.group(2)}月"
    purpose = section_by_prefix(sections, "2 文档目的")
    if purpose:
        purpose["paras"] = [
            f"本文基于开发人员的已核验 NESMA 详细 FPA 结论，用于说明{title}的已交付业务范围、"
            "数据模型、页面功能、接口契约和可验收约束，从而为需求支撑提供依据。"
        ]
        purpose["tables"] = []
        purpose["images"] = []
        purpose["meta"] = []

    scope = section_by_prefix(sections, "3 文档范围")
    if scope:
        scope["paras"] = [item for item in scope.get("paras", []) if compact(item)][:1]
        scope["tables"] = []
        scope["images"] = []
        scope["meta"] = []

    terms = section_by_prefix(sections, "4 术语和缩写解释")
    if terms:
        terms["paras"] = []
        terms["images"] = []
        terms["meta"] = []
        terms["tables"] = terms.get("tables", [])[:1]

    references = section_by_prefix(sections, "5 参考文档")
    if references:
        references["paras"] = list(REFERENCE_DOCUMENTS)
        references["tables"] = []
        references["images"] = []
        references["meta"] = []

    definition = section_by_prefix(sections, "6 范围定义")
    if definition:
        definition["tables"] = []
        definition["images"] = []
        definition["meta"] = []

    for index, section in enumerate(sections):
        if list(descendants(sections, index)):
            strip_direct_content(section)


def add_table_references(sections):
    for index, section in enumerate(sections):
        identifiers = []
        for image in section.get("images", []):
            match = re.search(r"表\s*(\d+(?:[.\-]\d+)*)", compact(image.get("caption")))
            if match:
                identifiers.append(f"表{match.group(1)}")
        if not identifiers:
            continue
        nested = [sections[index], *descendants(sections, index)]
        body = " ".join(
            str(part)
            for item in nested
            for part in item.get("paras", [])
        ).replace(" ", "")
        missing = [ident for ident in identifiers if ident not in body]
        if not missing:
            continue
        targets = [item for item in nested if any(compact(part) for part in item.get("paras", []))]
        require(targets, f"章节「{section.get('heading')}」有接口表但没有任何正文可挂引用")
        target = targets[-1]
        target["paras"][-1] = (
            str(target["paras"][-1]).rstrip()
            + f" 接口或数据契约见{'、'.join(missing)}。"
        )


def normalize_data_mappings(delivery):
    features = delivery.get("features")
    mappings = delivery.get("mappings")
    require(isinstance(features, list) and features, "交付模型缺少 features")
    require(isinstance(mappings, list), "交付模型缺少 mappings")
    by_data_function = {}
    for feature in features:
        for ident in feature.get("df", []):
            by_data_function.setdefault(ident, []).append(feature["row_id"])
    for mapping in mappings:
        if mapping.get("type") not in {"ILF", "ELF"}:
            continue
        rows = by_data_function.get(mapping.get("id"), [])
        require(rows, f"{mapping.get('id')} 未被任何功能清单行引用")
        mapping["feature_rows"] = "、".join(rows)


def clean_data_model_ids(value):
    return re.sub(r"\s*DF-\d+\s*", " ", str(value or "")).strip()


def strip_data_model_ids(sections):
    start = next(
        (
            index
            for index, section in enumerate(sections)
            if compact(section.get("heading")).startswith("7.1 数据模型")
        ),
        None,
    )
    require(start is not None, "缺少 7.1 数据模型章节")
    for section in [sections[start], *descendants(sections, start)]:
        section["heading"] = clean_data_model_ids(section.get("heading"))
        section["paras"] = [clean_data_model_ids(item) for item in section.get("paras", [])]
        for table in section.get("tables", []):
            table["headers"] = [clean_data_model_ids(item) for item in table.get("headers", [])]
            table["rows"] = [
                [clean_data_model_ids(item) for item in row]
                for row in table.get("rows", [])
            ]
        for image in section.get("images", []):
            image["caption"] = clean_data_model_ids(image.get("caption"))


def normalize_delivery(delivery):
    document = delivery.get("document")
    require(isinstance(document, dict), "交付模型缺少 document")
    sections = document.get("sections")
    require(isinstance(sections, list) and sections, "delivery document 缺少 sections")
    normalize_fixed_sections(document, sections)
    strip_data_model_ids(sections)
    add_table_references(sections)
    normalize_data_mappings(delivery)


def check_model(delivery):
    require(isinstance(delivery, dict), "交付模型必须是 JSON 对象")
    model = delivery.get("model")
    document = delivery.get("document")
    require(isinstance(model, dict), "交付模型缺少 model")
    require("assumptions" not in model and "unresolved" not in model, "阶段一模型含假设或未决项")
    require(isinstance(document, dict), "交付模型缺少 document")
    rows, data, transactions = count_rows(model)
    summary = count_summary(rows)
    require(document.get("counts") == summary, "document.counts 与计数模型不一致")
    by_id = {row["id"]: row for row in rows}
    require(len(by_id) == len(rows), "计数模型编号重复")

    mappings = delivery.get("mappings")
    require(isinstance(mappings, list), "交付模型缺少 mappings")
    require([row.get("id") for row in mappings] == [row["id"] for row in rows], "映射必须按计数模型顺序全覆盖")
    feature_ids = {feature.get("row_id") for feature in delivery.get("features", [])}
    feature_by_transaction = {
        feature.get("tf"): feature.get("row_id")
        for feature in delivery.get("features", [])
        if feature.get("tf")
    }
    for mapping in mappings:
        row = by_id[mapping["id"]]
        require(mapping.get("type") == row["type"], f"{mapping['id']} 类型与计数模型不一致")
        require(mapping.get("name") == row["name"], f"{mapping['id']} 名称与计数模型不一致")
        require(mapping.get("fp") == row["fp"], f"{mapping['id']} FP 与计数模型不一致")
        require(mapping.get("status") == "已覆盖", f"{mapping['id']} 尚未覆盖")
        require(mapping.get("feature_rows") and mapping.get("section_page_model"), f"{mapping['id']} 缺少正文归属")
        mapping_rows = [item.strip() for item in re.split(r"[、,，]", str(mapping["feature_rows"])) if item.strip()]
        require(mapping_rows and all(item in feature_ids for item in mapping_rows), f"{mapping['id']} 的 feature_rows 不是功能清单行")
        if row["type"] in {"EI", "EO", "EQ"}:
            require(mapping_rows == [feature_by_transaction.get(mapping["id"])], f"{mapping['id']} 必须唯一对应一行功能清单")

    features = delivery.get("features")
    require(isinstance(features, list) and features, "交付模型缺少 features")
    remarks = "\n".join(str(row.get("remark", "")) for row in features)
    for transaction in transactions:
        require(remarks.count(transaction["id"]) == 1, f"功能清单未唯一覆盖 {transaction['id']}")
    check_document_model(model, document, mappings)
    check_feature_granularity(features, transactions)
    return model, document, rows, data, transactions


def write_mapping_xlsx(path, mappings, rows):
    thin = Side(style="thin", color="BFC8C3")
    header_fill = PatternFill("solid", fgColor="DDE7E1")
    header_font = Font(name="宋体", size=11, bold=True, color="24312E")
    body_font = Font(name="宋体", size=10, color="24312E")
    wrap = Alignment(vertical="top", wrap_text=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NESMA映射"
    sheet.append(MAPPING_HEADERS)
    for row in mappings:
        sheet.append([
            row.get("id"),
            row.get("type"),
            row.get("name"),
            row.get("fp"),
            row.get("feature_rows"),
            row.get("section_page_model"),
            row.get("interfaces"),
            row.get("status"),
            row.get("note"),
        ])
    for index, width in enumerate([12, 8, 30, 8, 28, 48, 42, 12, 60], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap
            cell.border = Border(bottom=thin)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    summary = workbook.create_sheet("覆盖汇总")
    summary.append(["项目", "值", "说明"])
    types = Counter(row["type"] for row in rows)
    metrics = [
        ("数据功能", sum(row["type"] in {"ILF", "ELF"} for row in rows), f"{types['ILF']} ILF + {types['ELF']} ELF"),
        ("事务功能", sum(row["type"] in {"EI", "EO", "EQ"} for row in rows), f"{types['EI']} EI + {types['EO']} EO + {types['EQ']} EQ"),
        ("数据功能FP", sum(row["fp"] for row in rows if row["type"] in {"ILF", "ELF"}), "计数模型"),
        ("事务功能FP", sum(row["fp"] for row in rows if row["type"] in {"EI", "EO", "EQ"}), "计数模型"),
        ("UFP总数", sum(row["fp"] for row in rows), "数据功能FP + 事务功能FP"),
        ("未覆盖", sum(row.get("status") != "已覆盖" for row in mappings), "应为0"),
    ]
    for metric in metrics:
        summary.append(metric)
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 42
    workbook.save(path)


def write_feature_xlsx(path, features):
    workbook = load_workbook(TEMPLATE)
    sheet = workbook.active
    style_row = [copy(sheet.cell(2, column)._style) for column in range(1, 9)]
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for index, item in enumerate(features, 1):
        values = [index, item["system"], item["module_1"], item["module_2"], item["description"], item["external"], item["change"], item["remark"]]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index + 1, column, value)
            cell._style = copy(style_row[column - 1])
            cell.alignment = Alignment(horizontal="center" if column in {1, 7} else "left", vertical="top", wrap_text=True)
        sheet.row_dimensions[index + 1].height = 72
    for column in (2, 3):
        start = 2
        while start <= sheet.max_row:
            value = sheet.cell(start, column).value
            end = start
            while end < sheet.max_row and sheet.cell(end + 1, column).value == value:
                end += 1
            if end > start:
                sheet.merge_cells(start_row=start, start_column=column, end_row=end, end_column=column)
                sheet.cell(start, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            start = end + 1
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    workbook.save(path)


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9_.-]+", "-", str(value)).strip("-")


def normalize_fields(fields):
    normalized = []
    for field in fields or []:
        if isinstance(field, list) and len(field) == 3 and all(isinstance(item, str) for item in field):
            normalized.append({"name": field[0], "type": field[1], "label": field[2]})
        else:
            normalized.append(field)
    return normalized


def http_fields(spec):
    if "request" in spec:
        return normalize_fields(spec["request"])
    if spec.get("method", "").upper() in {"GET", "HEAD"}:
        return normalize_fields((spec.get("path_params") or []) + (spec.get("query") or []))
    return normalize_fields((spec.get("path_params") or []) + (spec.get("body") or spec.get("form") or []))


def http_response(spec):
    if "response" in spec:
        return normalize_fields(spec["response"])
    if spec.get("data"):
        return spec["data"]
    if spec.get("response_kind") == "file":
        return []
    return []


def write_yaml(root, delivery):
    records = []
    for spec in delivery.get("specs", []):
        kind = spec.get("kind", "http")
        record = {
            "id": spec["id"],
            "kind": kind,
            "title": spec["title"],
            "nesma": spec.get("nesma"),
            "section": spec.get("section"),
            "table": spec.get("table"),
            "caption": spec.get("caption", spec["title"]),
        }
        for key in ("method", "path", "controller", "auth", "envelope", "rules", "rule"):
            if spec.get(key) not in (None, "", []):
                record[key] = spec[key]
        if spec.get("response_kind") == "file":
            record["file"] = {"kind": spec.get("file_kind", "file")}
        record["request"] = http_fields(spec)
        record["response"] = http_response(spec)
        records.append((kind, record))

    for master in delivery.get("masters", []):
        require(master.get("source") and master.get("direction"), f"{master.get('id')} 缺少 source 或 direction")
        record = {
            "id": master["id"],
            "kind": "elf",
            "title": master["title"],
            "nesma": master.get("nesma"),
            "section": master.get("section"),
            "table": master.get("table"),
            "caption": master.get("caption", master["title"]),
            "source": master["source"],
            "direction": master["direction"],
            "fields": master.get("fields", master.get("data", [])),
            "rules": master.get("rules", master.get("rule")),
        }
        records.append(("elf", record))

    for db in delivery.get("dbs", []):
        columns = [
            {
                "name": column["name"],
                "type": column["type"],
                "nullable": column.get("nullable", not column.get("required", False)),
                "desc": column.get("desc", ""),
            }
            for column in db["columns"]
        ]
        records.append(("db", {
            "id": db["id"],
            "kind": "db",
            "title": db["title"],
            "nesma": db.get("nesma"),
            "section": db.get("section"),
            "table": db.get("table"),
            "caption": db.get("caption", db["title"]),
            "tableName": db.get("tableName", db.get("name")),
            "columns": columns,
            "purpose": db.get("purpose", db.get("rule")),
            "rules": db.get("rules", db.get("rule", db.get("purpose"))),
        }))

    for kind, record in records:
        require(record.get("id") and record.get("title"), f"{kind} 卡片缺少 id 或 title")
        folder = root / kind
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{safe_name(record['id'])}.yaml"
        path.write_text(yaml.safe_dump(record, allow_unicode=True, sort_keys=False, width=100), encoding="utf-8")


def remap_card_images(document, delivery, rendered):
    rendered_by_id = {
        card["id"]: (Path("接口表") / Path(card["image"]).relative_to(Path(rendered["root"]))).as_posix()
        for card in rendered["cards"]
    }
    by_old_name = {}
    for record in delivery.get("specs", []) + delivery.get("masters", []) + delivery.get("dbs", []):
        target = rendered_by_id.get(record.get("id"))
        if not target:
            continue
        for key in ("image_rel", "image_name"):
            value = record.get(key)
            if value:
                by_old_name[Path(value).name] = target
        for value in record.get("image_rels", []):
            by_old_name[Path(value).name] = target
    for section in document.get("sections", []):
        for image in section.get("images", []):
            path = Path(image["path"])
            target = by_old_name.get(path.name)
            if target and path.parts[:2] == ("接口表", "images"):
                image["path"] = target


def copy_document_assets(model_dir, output, document):
    for section in document.get("sections", []):
        for image in section.get("images", []):
            relative = Path(image["path"])
            require(not relative.is_absolute() and ".." not in relative.parts, f"图片路径必须是交付目录内相对路径：{relative}")
            if relative.parts[:2] == ("接口表", "images"):
                continue
            source = model_dir / relative
            require(source.is_file(), f"缺少项目图片：{source}")
            target = output / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)


def find_dotnet():
    candidate = Path.home() / ".dotnet/dotnet"
    if candidate.is_file():
        return candidate
    found = shutil.which("dotnet")
    require(found, "未找到 dotnet")
    return Path(found)


def find_libreoffice():
    found = shutil.which("libreoffice") or shutil.which("soffice")
    require(found, "未找到 libreoffice/soffice，无法回填目录页码")
    return Path(found)


def card_style(delivery, model_path):
    raw = delivery.get("card_style")
    if not raw:
        return STYLE
    path = Path(raw)
    if not path.is_absolute():
        path = model_path.parent / path
    require(path.is_file(), f"卡片样式不存在：{path}")
    require(path.name == "api-card.sty", "卡片样式文件名必须是 api-card.sty")
    return path


def build(model_path, output):
    require(model_path.is_file(), f"交付模型不存在：{model_path}")
    require(not output.exists() or (output.is_dir() and not any(output.iterdir())), f"输出目录必须不存在或为空：{output}")
    delivery = json.loads(model_path.read_text(encoding="utf-8"))
    normalize_delivery(delivery)
    model, document, rows, data, transactions = check_model(delivery)
    style = card_style(delivery, model_path)

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=f".{output.name}-", dir=output.parent) as temporary:
        work = Path(temporary)
        copy_document_assets(model_path.parent, work, document)
        write_mapping_xlsx(work / "映射表.xlsx", delivery["mappings"], rows)
        write_feature_xlsx(work / "功能清单.xlsx", delivery["features"])
        interface_dir = work / "接口表"
        interface_dir.mkdir()
        write_yaml(interface_dir, delivery)
        rendered = subprocess.run(
            [sys.executable, str(RENDER_CARDS), str(interface_dir), str(style)],
            text=True,
            capture_output=True,
        )
        require(rendered.returncode == 0, rendered.stderr or rendered.stdout)
        remap_card_images(document, delivery, json.loads(rendered.stdout))
        (work / "delivery-model.json").write_text(json.dumps(delivery, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        docx_path = work / "需求说明书.docx"
        subprocess.run([
            str(find_dotnet()), "run", "--project", str(DOCX), "--",
            str(work / "delivery-model.json"), str(docx_path),
        ], check=True)
        toc_dir = work / ".toc"
        toc_dir.mkdir()
        subprocess.run([
            str(find_libreoffice()), "--headless", "--convert-to", "pdf",
            "--outdir", str(toc_dir), str(docx_path),
        ], check=True, capture_output=True, text=True)
        toc_pdf = toc_dir / "需求说明书.pdf"
        require(toc_pdf.is_file(), "LibreOffice 未生成目录页码所需的 PDF")
        subprocess.run([
            sys.executable, str(UPDATE_TOC),
            str(work / "delivery-model.json"), str(toc_pdf),
        ], check=True)
        subprocess.run([
            str(find_dotnet()), "run", "--project", str(DOCX), "--",
            str(work / "delivery-model.json"), str(docx_path),
        ], check=True)
        shutil.rmtree(toc_dir)
        count_path = work / ".count-model.json"
        count_path.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
        validation = subprocess.run(
            [sys.executable, str(VALIDATE), str(count_path), str(work)],
            text=True,
            capture_output=True,
        )
        require(validation.returncode == 0, validation.stderr or validation.stdout)
        count_path.unlink()
        if output.exists():
            output.rmdir()
        work.rename(output)

    print(json.dumps({
        "status": "PASS",
        "output": str(output),
        "counts": count_summary(rows),
        "mappings": len(delivery["mappings"]),
        "features": len(delivery["features"]),
    }, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("delivery_model", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.delivery_model, args.output)


if __name__ == "__main__":
    main()
