#!/usr/bin/env python3
"""Validate a stage-two delivery against its stage-one count model."""

import argparse
import json
import posixpath
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import yaml
from openpyxl import load_workbook

from delivery_quality import check_document_model, check_feature_granularity


def fail(message):
    raise SystemExit(f"FAIL: {message}")


def require(condition, message):
    if not condition:
        fail(message)


def expected_ids(model):
    data = model.get("data_functions")
    transactions = model.get("transactions")
    require(isinstance(data, list) and data, "计数模型缺少 data_functions")
    require(isinstance(transactions, list) and transactions, "计数模型缺少 transactions")
    return [row["id"] for row in data] + [row["id"] for row in transactions], data, transactions


def check_mapping(output, model_rows, ids):
    path = output / "映射表.xlsx"
    require(path.exists(), f"缺少 {path.name}")
    workbook = load_workbook(path, data_only=True)
    require("NESMA映射" in workbook.sheetnames, "映射表缺少 NESMA映射 工作表")
    sheet = workbook["NESMA映射"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    mapped = [row[0] for row in rows]
    require(mapped == ids, "映射表编号必须与计数模型完全一致")
    require(len(mapped) == len(set(mapped)), "映射表编号重复")
    require(all(row[7] == "已覆盖" for row in rows), "映射表存在未覆盖项")
    require(all(row[4] and row[5] for row in rows), "映射表功能清单行或正文位置为空")
    by_id = {row["id"]: row for row in model_rows}
    for mapped_row in rows:
        source = by_id[mapped_row[0]]
        require(
            (mapped_row[1], mapped_row[2], mapped_row[3]) == (source["type"], source["name"], source["fp"]),
            f"映射表 {mapped_row[0]} 的类型、名称或 FP 与计数模型不一致",
        )
    return {"rows": len(rows)}


def check_features(output, transactions):
    path = output / "功能清单.xlsx"
    require(path.exists(), f"缺少 {path.name}")
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    expected = ["序号", "系统名称", "一级模块", "二级模块", "功能需求描述", "外部接口需求描述", "功能（新增/修改）", "备注"]
    require(headers == expected, "功能清单列顺序不符合模板")
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    remarks = "\n".join(str(row[7]) for row in rows)
    for row in transactions:
        require(remarks.count(row["id"]) == 1, f"功能清单未唯一覆盖 {row['id']}")
    last_system = ""
    last_module = ""
    for row in rows:
        last_system = row[1] or last_system
        last_module = row[2] or last_module
        require(last_system and last_module and row[3] and row[4], "功能清单存在缺少系统/模块/描述的行")
    return {"rows": len(rows)}


def flatten_nesma(value):
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    require(isinstance(value, list) and all(isinstance(item, str) for item in value), "nesma 必须是字符串或字符串列表")
    return value


def check_interfaces(output, ids):
    root = output / "接口表"
    require(root.exists(), "缺少 接口表 目录")
    yaml_files = sorted(root.rglob("*.yaml"))
    require(yaml_files, "接口表没有 YAML")
    images = [path for path in (root / "images").rglob("*") if path.is_file()] if (root / "images").exists() else []
    require(images, "接口表没有图片")
    require(len(images) >= len(yaml_files), f"YAML 与图片数量不一致：{len(yaml_files)} vs {len(images)}")
    seen = set()
    http = 0
    elf = 0
    db = 0
    known = set(ids)
    for path in yaml_files:
        record = yaml.safe_load(path.read_text(encoding="utf-8"))
        require(isinstance(record, dict), f"{path.name} 不是 YAML 对象")
        ident = record.get("id")
        require(isinstance(ident, str) and ident, f"{path.name} 缺少 id")
        require(ident not in seen, f"接口 id 重复：{ident}")
        seen.add(ident)
        record_type = record.get("kind")
        if record_type is None and "columns" in record:
            record_type = "db"
        require(record_type in {"http", "elf", "db"}, f"{ident} 的 kind 无效")
        if record_type == "http":
            require(all(record.get(key) for key in ("method", "path", "controller")), f"{ident} 缺少 HTTP 契约字段")
            http += 1
        elif record_type == "elf":
            require(record.get("nesma"), f"{ident} 缺少 nesma")
            elf += 1
        else:
            require(record.get("columns"), f"{ident} 缺少 columns")
            db += 1
        for ref in flatten_nesma(record.get("nesma")):
            if ref == "uncounted":
                continue
            require(ref in known, f"{ident} 引用不存在的 NESMA 编号：{ref}")
    return {"yaml": len(yaml_files), "images": len(images), "http": http, "elf": elf, "db": db}


def check_docx(output, ufp, ids, image_count):
    path = output / "需求说明书.docx"
    require(path.exists(), f"缺少 {path.name}")
    with ZipFile(path) as archive:
        require("word/document.xml" in archive.namelist(), "Word 缺少 document.xml")
        root = ET.fromstring(archive.read("word/document.xml"))
        text = "".join(root.itertext())
        ufp_values = {int(value) for value in re.findall(r"(\d+)\s*UFP", text)}
        require(ufp_values == {ufp}, f"Word UFP 与计数模型不一致：{sorted(ufp_values)}")
        nesma_values = set(re.findall(r"\b(?:DF|TF)-\d+\b", text))
        unknown = sorted(nesma_values - set(ids))
        require(not unknown, f"Word 含非阶段一 NESMA 编号：{', '.join(unknown)}")
        rels = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        media = []
        for relation in rels:
            if not relation.attrib.get("Type", "").endswith("/image"):
                continue
            target = relation.attrib["Target"].lstrip("/")
            candidates = [target, posixpath.normpath(posixpath.join("word", target))]
            require(any(candidate in archive.namelist() for candidate in candidates), f"Word 图片关系目标不存在：{target}")
            media.append(target)
        require(len(media) >= image_count, f"Word 图片数少于接口表图片数：{len(media)} vs {image_count}")
    return {"media": len(media), "chars": len(text)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("count_model", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--delivery-model", type=Path)
    args = parser.parse_args()
    model = json.loads(args.count_model.read_text(encoding="utf-8"))
    require("assumptions" not in model and "unresolved" not in model, "阶段一模型含假设或未决项")
    ids, data, transactions = expected_ids(model)
    ufp = sum(row["fp"] for row in data + transactions)
    delivery_path = args.delivery_model or args.output / "delivery-model.json"
    require(delivery_path.exists(), "缺少 delivery-model.json")
    delivery = json.loads(delivery_path.read_text(encoding="utf-8"))
    check_document_model(model, delivery.get("document"), delivery.get("mappings"))
    check_feature_granularity(delivery.get("features"), transactions)
    mapping = check_mapping(args.output, data + transactions, ids)
    features = check_features(args.output, transactions)
    interfaces = check_interfaces(args.output, ids)
    docx = check_docx(args.output, ufp, ids, interfaces["images"])
    print(json.dumps({
        "status": "PASS",
        "counts": {"data_functions": len(data), "transactions": len(transactions), "ufp": ufp},
        "mapping": mapping,
        "features": features,
        "interfaces": interfaces,
        "docx": docx,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
