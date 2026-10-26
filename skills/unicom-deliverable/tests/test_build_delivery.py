import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from docx import Document

REPO = Path(__file__).parents[3]
SCRIPTS = Path(__file__).parents[1] / "scripts"
BUILD = SCRIPTS / "build_delivery.py"
VALIDATE = SCRIPTS / "validate_delivery.py"
sys.path.insert(0, str(SCRIPTS))
from build_delivery import add_table_references

COUNT = REPO / "tests/fixtures/nesma-employee-assessment-v3/count/计数模型.json"
MODEL = REPO / "tests/fixtures/nesma-employee-assessment-v3/delivery/delivery-model.json"


@unittest.skipUnless(
    shutil.which("xelatex")
    and Path.home().joinpath(".dotnet/dotnet").exists()
    and shutil.which("libreoffice")
    and shutil.which("pdftotext")
    and shutil.which("pdfinfo"),
    "render toolchain unavailable",
)
class DeliveryBuildTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.output = Path(cls.temp.name) / "delivery"
        result = subprocess.run(
            [sys.executable, str(BUILD), str(MODEL), str(cls.output)],
            text=True,
            capture_output=True,
        )
        if result.returncode:
            raise AssertionError(result.stderr or result.stdout)
        cls.validation = subprocess.run(
            [sys.executable, str(VALIDATE), str(COUNT), str(cls.output)],
            text=True,
            capture_output=True,
        )

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_golden_delivery_builds_and_passes(self):
        self.assertEqual(self.validation.returncode, 0, self.validation.stderr)
        payload = json.loads(self.validation.stdout)
        self.assertEqual(payload["status"], "PASS")
        self.assertEqual(payload["counts"]["ufp"], 267)
        self.assertEqual(payload["interfaces"]["yaml"], 73)
        self.assertEqual(payload["docx"]["media"], 74)
        self.assertGreater(payload["toc"]["first"], 1)
        self.assertGreaterEqual(payload["toc"]["last"], payload["toc"]["first"])

    def copy_output(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        output = Path(temp.name) / "delivery"
        shutil.copytree(self.output, output)
        return output

    def replace_docx_text(self, output, old, new):
        path = output / "需求说明书.docx"
        rewritten = path.with_suffix(".new")
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/document.xml":
                    data = data.decode("utf-8").replace(old, new, 1).encode("utf-8")
                target.writestr(item, data)
        rewritten.replace(path)

    def test_missing_card_image_is_rejected(self):
        output = self.copy_output()
        next((output / "接口表/images").glob("*.png")).unlink()
        result = subprocess.run(
            [sys.executable, str(VALIDATE), str(COUNT), str(output)],
            text=True,
            capture_output=True,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("YAML 与图片数量不一致", result.stderr)

    def test_old_ufp_is_rejected(self):
        output = self.copy_output()
        self.replace_docx_text(output, "267 UFP", "267 UFP 旧版 233 UFP")
        result = subprocess.run(
            [sys.executable, str(VALIDATE), str(COUNT), str(output)],
            text=True,
            capture_output=True,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("UFP", result.stderr)

    def test_unknown_nesma_id_in_word_is_rejected(self):
        output = self.copy_output()
        self.replace_docx_text(output, "UFP", "UFP TF-99")
        result = subprocess.run(
            [sys.executable, str(VALIDATE), str(COUNT), str(output)],
            text=True,
            capture_output=True,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("TF-99", result.stderr)

    def test_cover_has_no_version_or_day(self):
        with zipfile.ZipFile(self.output / "需求说明书.docx") as archive:
            text = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("员工考核系统改造", text)
        self.assertIn("需求说明书", text)
        self.assertIn("2026年10月", text)
        self.assertNotIn("v3-new-v2", text)
        self.assertNotIn("2026年10月23日", text)

    def test_data_function_mapping_uses_feature_rows(self):
        workbook = load_workbook(self.output / "映射表.xlsx", data_only=True)
        rows = list(workbook["NESMA映射"].iter_rows(min_row=2, values_only=True))
        data_rows = [row for row in rows if row[1] in {"ILF", "ELF"}]
        self.assertTrue(data_rows)
        for row in data_rows:
            self.assertTrue(str(row[4]).startswith("F-"), row[0])

    def test_fixed_front_sections_are_normalized(self):
        delivery = json.loads((self.output / "delivery-model.json").read_text(encoding="utf-8"))
        sections = {
            section["heading"]: section
            for section in delivery["document"]["sections"]
        }
        self.assertEqual(sections["1 项目概述"]["paras"], [])
        self.assertEqual(len(sections["4 术语和缩写解释"]["paras"]), 0)
        self.assertEqual(sections["6 范围定义"]["tables"], [])
        page = next(
            section
            for section in delivery["document"]["sections"]
            if section["heading"].startswith("7.2.1 ")
        )
        self.assertEqual(page["paras"], [])
        self.assertEqual(
            sections["5 参考文档"]["paras"],
            [
                "《系统与软件工程功能规模测量 NESMA 方法 GB/T42588—2023》",
                "《2025年中国软件行业基准数据（CSBMK® -202510）》",
                "GB∕T36964-2018 软件工程 软件开发成本度量规范",
            ],
        )

    def test_missing_table_reference_is_added(self):
        sections = [{
            "heading": "7.2.1.1 查询记录",
            "level": 4,
            "paras": ["查询记录。"],
            "images": [{"path": "接口表/images/item.list.png", "caption": "表7-1 查询记录"}],
            "tables": [],
        }]
        add_table_references(sections)
        self.assertIn("接口或数据契约见表7-1", sections[0]["paras"][-1])

    def test_nonleaf_headings_have_no_blank_paragraph(self):
        document = Document(self.output / "需求说明书.docx")
        for heading in ("1 项目概述", "7.1 数据模型", "7.2.1 考核模板页面"):
            index = next(
                index
                for index, paragraph in enumerate(document.paragraphs)
                if paragraph.text == heading
            )
            paragraph = document.paragraphs[index]
            following = document.paragraphs[index + 1]
            self.assertTrue(following.style.name.startswith("Heading"), heading)
            self.assertEqual(paragraph.paragraph_format.space_after.pt, 0)

    def test_data_model_has_no_df_ids(self):
        delivery = json.loads((self.output / "delivery-model.json").read_text(encoding="utf-8"))
        data_model = [
            section
            for section in delivery["document"]["sections"]
            if section["heading"].startswith("7.1")
        ]
        self.assertTrue(data_model)
        self.assertFalse(any("DF-" in json.dumps(section, ensure_ascii=False) for section in data_model))


if __name__ == "__main__":
    unittest.main()
