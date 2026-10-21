#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import re
import shutil
import subprocess
import textwrap
from collections import OrderedDict
from copy import copy
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageChops, ImageDraw, ImageFont


WORK = Path("/tmp/hx-delivery-work")
OUT = Path("/home/nc/projects/work/hx-employee-assessment-delivery-v3-new-v2")
COUNT_DIR = Path(
    "/home/nc/projects/work/cuc-ah-ucac-pso-hx-repository/"
    "doc/hx-assessment/NESMA估算/员工年度考核-v3-new-v2"
)
BACKEND = WORK / "backend-baseline/hx_assessment"
FRONTEND = Path("/home/nc/projects/work/hx-emp-assessment-view-repo")
TEMPLATE_XLSX = Path("/home/nc/Downloads/需求规格说明书附件功能清单.xlsx")

JSON_PATH = COUNT_DIR / "计数模型.json"
XLSX_PATH = COUNT_DIR / "功能点计数证据.xlsx"

FONT_REGULAR = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FONT_BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
if not Path(FONT_REGULAR).exists():
    FONT_REGULAR = "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"
if not Path(FONT_BOLD).exists():
    FONT_BOLD = FONT_REGULAR

COLORS = {
    "primary": "#344B46",
    "secondary": "#5E746D",
    "accent": "#B69563",
    "dark": "#24312E",
    "mid": "#58635F",
    "light": "#8A9792",
    "border": "#D7DDD9",
    "panel": "#F3F6F4",
    "white": "#FFFFFF",
}


def load_json() -> dict[str, Any]:
    return json.loads(JSON_PATH.read_text(encoding="utf-8"))


def java_source(class_name: str) -> Path:
    matches = list(BACKEND.glob(f"src/main/java/**/{class_name}.java"))
    if not matches:
        raise FileNotFoundError(class_name)
    return matches[0]


def clean_comment(comment: str) -> str:
    lines = []
    for line in comment.splitlines():
        line = re.sub(r"^\s*/\*\*?", "", line)
        line = re.sub(r"\*/\s*$", "", line)
        line = re.sub(r"^\s*\*\s?", "", line)
        line = line.strip()
        if line and not line.startswith("@"):
            lines.append(line)
    text = " ".join(lines)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("{@code ", "").replace("}", "")
    text = re.sub(r"\s+", " ", text).strip()
    first = re.split(r"[。；;]", text, maxsplit=1)[0].strip()
    return first or "业务字段"


def parse_java_fields(class_name: str) -> list[dict[str, Any]]:
    source = java_source(class_name).read_text(encoding="utf-8")
    nested = re.search(r"\n\s*(?:@\w+\s*)*(?:public\s+)?static\s+class\s+\w+", source)
    if nested:
        source = source[: nested.start()]
    pattern = re.compile(
        r"(?P<comment>/\*\*.*?\*/\s*)?"
        r"(?P<annotations>(?:@\w+(?:\([^)]*\))?\s*)*)"
        r"private\s+(?:final\s+)?(?P<type>[\w<>, ?.\[\]]+?)\s+"
        r"(?P<name>\w+)(?:\s*=\s*[^;]+)?;",
        re.S,
    )
    fields = []
    for match in pattern.finditer(source):
        type_name = re.sub(r"\s+", " ", match.group("type")).strip()
        if type_name in {"static", "final"}:
            continue
        fields.append(
            {
                "name": match.group("name"),
                "java_type": type_name,
                "label": clean_comment(match.group("comment") or ""),
                "comment": (match.group("comment") or "").strip(),
            }
        )
    return fields


LABEL_OVERRIDES = {
    "id": "主键",
    "planId": "计划主键",
    "templateId": "模板主键",
    "assessmentYear": "考核年度",
    "templateName": "模板名称",
    "description": "说明",
    "status": "状态",
    "enabled": "是否可用于新计划",
    "createdBy": "创建人主键",
    "createdAt": "创建时间",
    "updatedBy": "最后修改人主键",
    "updatedAt": "最后修改时间",
    "draftRevision": "草稿版本",
    "expectedRevision": "客户端读取到的草稿版本",
    "remark": "备注",
    "createdByName": "发起人",
    "publishedAt": "发布时间",
    "completedAt": "完成时间",
    "terminatedAt": "终止时间",
    "employeeId": "员工编号",
    "personName": "姓名",
    "employeeCode": "员工编号",
    "departmentName": "考核部门名称",
    "organizationName": "现单位",
    "positionType": "考核职级编码",
    "positionTypeLabel": "考核职级展示名",
    "planDepartmentId": "计划部门主键",
    "planPersonId": "计划人员主键",
    "assignmentId": "计划任职主键",
    "assessmentTarget": "是否为受考核任职",
    "annualWorkSummary": "年度主要工作摘要",
    "assessmentGrade": "年度考核等次",
    "submittedAt": "首次提交时间",
    "readOnly": "是否只读",
    "totalScore": "评分总分",
    "summaryScore": "评分汇总分",
    "ratingGroupId": "评分组主键",
    "groupType": "评分组类型",
    "groupTypeLabel": "评分组中文名",
    "weight": "权重百分比",
    "averageScore": "组内平均分",
    "weightedScore": "组加权分",
    "taskId": "评分任务主键",
    "seatId": "评分席位主键",
    "indicatorId": "叶子指标主键",
    "indicatorName": "指标名称",
    "maxScore": "指标满分",
    "parentId": "父分组主键",
    "parentName": "父分组名称",
    "score": "指标所得分",
    "accountCode": "匿名账号",
    "accountType": "匿名账号类型",
    "password": "密码",
    "outcome": "兑换结果",
    "workbench": "工作台初始化数据",
    "tasks": "评分任务",
    "seatQuotas": "席位评分配额边界",
    "total": "总记录数",
    "rows": "记录列表",
    "errors": "可定位的错误列表",
    "accepted": "是否通过校验并完成写入",
    "draftVersion": "草稿版本",
    "exists": "是否已有持久化登记表",
    "ready": "是否满足发布就绪条件",
    "years": "可选年度集合",
    "defaultYear": "默认年度",
    "primary": "主考核卡片",
    "scope": "当前页分管范围",
    "scopes": "全部领导任职的最终分管范围",
    "publicationBlockers": "待配置发布事项",
    "template": "模板基础信息",
    "sheetName": "工作表名称",
    "rowNumber": "Excel 物理行号",
    "column": "Excel 列标识",
    "message": "校验失败说明",
}


def field_specs(
    class_name: str,
    *,
    skip: set[str] | None = None,
    recursive: bool = True,
    depth: int = 0,
) -> list[dict[str, Any]]:
    skip = skip or set()
    result = []
    for item in parse_java_fields(class_name):
        if item["name"] in skip:
            continue
        name = item["name"]
        item["label"] = LABEL_OVERRIDES.get(name, item["label"])
        item.pop("comment", None)
        type_name = item["java_type"]
        base_match = re.search(r"(?:List<)?([A-Z]\w+)(?:>)?", type_name)
        base_type = base_match.group(1) if base_match else ""
        scalar = base_type in {
            "String",
            "Long",
            "Integer",
            "Boolean",
            "BigDecimal",
            "Date",
            "Object",
        }
        result.append(item)
        if (
            recursive
            and depth < 1
            and base_type
            and not scalar
            and base_type not in {"MultipartFile"}
        ):
            try:
                item["fields"] = field_specs(
                    base_type, recursive=True, depth=depth + 1
                )
            except FileNotFoundError:
                pass
    return result


def nested_field_specs(class_name: str, nested_name: str) -> list[dict[str, Any]]:
    source = java_source(class_name).read_text(encoding="utf-8")
    marker = source.find(f"class {nested_name}")
    if marker < 0:
        return []
    tail = source[marker:]
    # The next nested class starts a new scope in the DTO/VO files used here.
    next_match = re.search(r"\n\s*(?:public\s+)?static\s+class\s+\w+", tail[1:])
    if next_match:
        tail = tail[: next_match.start() + 1]
    pattern = re.compile(
        r"(?P<comment>/\*\*.*?\*/\s*)?"
        r"(?P<annotations>(?:@\w+(?:\([^)]*\))?\s*)*)"
        r"private\s+(?P<type>[\w<>, ?.\[\]]+?)\s+(?P<name>\w+)(?:\s*=\s*[^;]+)?;",
        re.S,
    )
    out = []
    for match in pattern.finditer(tail):
        out.append(
            {
                "name": match.group("name"),
                "java_type": re.sub(r"\s+", " ", match.group("type")).strip(),
                "label": LABEL_OVERRIDES.get(
                    match.group("name"),
                    clean_comment(match.group("comment") or ""),
                ),
            }
        )
    return out


def api(
    api_id: str,
    resource: str,
    category: str,
    category_no: int,
    title: str,
    nesma: str | list[str],
    controller: str,
    method: str,
    path: str,
    consumes: str,
    auth: str,
    envelope: str,
    page_section: str,
    rule: str,
    *,
    response_java: str = "",
    data_java: str = "",
    body_java: str = "",
    query_java: str = "",
    path_params: list[tuple[str, str, str]] | None = None,
    form: list[dict[str, Any]] | None = None,
    response_kind: str = "json",
    file_kind: str = "",
    same_origin: bool = False,
    data_fields: list[dict[str, Any]] | None = None,
    excluded_reason: str = "",
) -> dict[str, Any]:
    return {
        "id": api_id,
        "kind": "http",
        "resource": resource,
        "category": category,
        "category_no": category_no,
        "title": title,
        "nesma": nesma,
        "controller": controller,
        "method": method,
        "path": path,
        "consumes": consumes,
        "auth": auth,
        "envelope": envelope,
        "response_java": response_java,
        "data_java": data_java,
        "body_java": body_java,
        "query_java": query_java,
        "path_params": path_params or [],
        "form": form or [],
        "page_section": page_section,
        "rule": rule,
        "response_kind": response_kind,
        "file_kind": file_kind,
        "same_origin": same_origin,
        "data_fields": data_fields or [],
        "excluded_reason": excluded_reason,
    }


def endpoint_specs() -> list[dict[str, Any]]:
    specs = [
        api(
            "template.create",
            "template",
            "考核模板",
            1,
            "新建考核模板",
            "TF-01",
            "TemplateController#createTemplate",
            "POST",
            "/ea/template/templates",
            "json",
            "admin",
            "ajax-result",
            "8.2.2",
            "管理员提交模板名称和说明；系统创建草稿并记录当前管理员审计信息。",
            response_java="Template",
            data_java="Template",
            body_java="Template",
        ),
        api(
            "template.update",
            "template",
            "考核模板",
            1,
            "修改考核模板资料",
            "TF-02",
            "TemplateController#updateTemplate",
            "PUT",
            "/ea/template/templates/{templateId}",
            "json",
            "admin",
            "ajax-result",
            "8.2.3",
            "仅草稿模板允许修改名称和说明；模板路径参数必须存在且为草稿。",
            body_java="Template",
            path_params=[("templateId", "Long", "模板主键")],
        ),
        api(
            "template.configure-indicators",
            "template",
            "考核模板",
            1,
            "配置考核指标",
            "TF-03",
            "TemplateController#configureIndicators",
            "PUT",
            "/ea/template/templates/{templateId}/indicators",
            "json",
            "admin",
            "ajax-result",
            "8.2.5",
            "管理员整体替换草稿模板的两层指标结构；叶指标满分合计必须为100。",
            body_java="TemplateIndicatorBatchDTO",
            path_params=[("templateId", "Long", "模板主键")],
        ),
        api(
            "template.publish",
            "template",
            "考核模板",
            1,
            "发布考核模板",
            "TF-04",
            "TemplateController#publishTemplate",
            "POST",
            "/ea/template/templates/{templateId}/publish",
            "none",
            "admin",
            "ajax-result",
            "8.2.6",
            "模板必须具有完整两层指标且叶指标满分合计为100；发布后模板只读。",
            path_params=[("templateId", "Long", "模板主键")],
        ),
        api(
            "template.disable",
            "template",
            "考核模板",
            1,
            "停用考核模板",
            "TF-05",
            "TemplateController#disableTemplate",
            "POST",
            "/ea/template/templates/{templateId}/disable",
            "none",
            "admin",
            "ajax-result",
            "8.2.7",
            "停用只影响新计划选择，不改变已引用该模板的计划。",
            path_params=[("templateId", "Long", "模板主键")],
        ),
        api(
            "template.copy",
            "template",
            "考核模板",
            1,
            "复制考核模板",
            "TF-06",
            "TemplateController#copyTemplate",
            "POST",
            "/ea/template/templates/{templateId}/copy",
            "json",
            "admin",
            "ajax-result",
            "8.2.8",
            "系统从来源模板复制完整指标结构，并创建相互独立的新草稿。",
            response_java="Template",
            data_java="Template",
            body_java="TemplateCopyDTO",
            path_params=[("templateId", "Long", "源模板主键")],
        ),
        api(
            "template.delete",
            "template",
            "考核模板",
            1,
            "删除草稿考核模板",
            "TF-07",
            "TemplateController#deleteTemplate",
            "DELETE",
            "/ea/template/templates/{templateId}",
            "none",
            "admin",
            "ajax-result",
            "8.2.9",
            "仅从未发布的草稿模板允许删除；已发布模板只能停用。",
            path_params=[("templateId", "Long", "模板主键")],
        ),
        api(
            "template.list",
            "template",
            "考核模板",
            1,
            "查询考核模板",
            "TF-08",
            "TemplateController#listTemplates",
            "GET",
            "/ea/template/templates",
            "none",
            "admin",
            "table-data-info",
            "8.2.1",
            "按模板名称和状态筛选，按平台分页约定返回模板列表。",
            response_java="TemplateListVO",
            data_java="TemplateListVO",
            query_java="TemplateQueryDTO",
        ),
        api(
            "template.detail",
            "template",
            "考核模板",
            1,
            "查询考核模板详情",
            "TF-08",
            "TemplateController#getTemplateDetail",
            "GET",
            "/ea/template/templates/{templateId}",
            "none",
            "admin",
            "ajax-result",
            "8.2.4",
            "返回模板基本信息及完整两层指标树，不执行维护。",
            response_java="TemplateDetailVO",
            data_java="TemplateDetailVO",
            path_params=[("templateId", "Long", "模板主键")],
        ),
        api(
            "plan.create",
            "plan",
            "年度考核计划",
            2,
            "新建年度考核计划",
            "TF-09",
            "PlanController#createPlan",
            "POST",
            "/ea/plan/plans",
            "json",
            "admin",
            "ajax-result",
            "8.3.2",
            "同一年度最多存在一个未终止计划；模板必须已发布且可用于新计划。",
            response_java="Plan",
            data_java="Plan",
            body_java="PlanCreateDTO",
        ),
        api(
            "plan.update-draft",
            "plan",
            "年度考核计划",
            2,
            "修改草稿计划",
            "TF-10",
            "PlanController#updateDraft",
            "PUT",
            "/ea/plan/plans/{planId}",
            "json",
            "admin",
            "ajax-result",
            "8.4.1",
            "客户端必须提交读取时的草稿版本；模板或备注变化后重建未发布结构。",
            response_java="Plan",
            data_java="Plan",
            body_java="PlanDraftUpdateDTO",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.delete",
            "plan",
            "年度考核计划",
            2,
            "删除草稿计划",
            "TF-11",
            "PlanController#deleteDraft",
            "DELETE",
            "/ea/plan/plans/{planId}",
            "none",
            "admin",
            "ajax-result",
            "8.3.3",
            "仅从未发布的草稿计划允许删除。",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.list",
            "plan",
            "年度考核计划",
            2,
            "查询年度考核计划",
            "TF-12",
            "PlanController#listPlans",
            "GET",
            "/ea/plan/plans",
            "none",
            "admin",
            "table-data-info",
            "8.3.1",
            "按考核年度和生命周期状态筛选，返回计划及时间事实。",
            response_java="PlanListVO",
            data_java="PlanListVO",
            query_java="PlanQueryDTO",
        ),
        api(
            "plan.personnel-import-template",
            "plan",
            "计划人员",
            3,
            "下载计划人员导入模板",
            "uncounted",
            "PlanController#downloadPersonnelImportTemplate",
            "GET",
            "/ea/plan/plans/personnel-import-template",
            "none",
            "admin",
            "file-stream",
            "8.5.1",
            "输出两张固定表头的空Excel模板，不读取或计算持久业务数据。",
            response_kind="file",
            file_kind="xlsx",
            excluded_reason="人员导入模板下载；固定空模板，按阶段一决定不单列。",
        ),
        api(
            "plan.personnel-precheck",
            "plan",
            "计划人员",
            3,
            "核对人员文件",
            "TF-13",
            "PlanController#precheckPersonnel",
            "POST",
            "/ea/plan/plans/{planId}/personnel-precheck",
            "multipart",
            "admin",
            "ajax-result",
            "8.5.2",
            "只读解析完整人员文件，返回校验、替换影响和结构预览；不写库。",
            response_java="PlanPersonnelPrecheckResult",
            data_java="PlanPersonnelPrecheckResult",
            path_params=[("planId", "Long", "计划主键")],
            form=[{"name": "file", "java_type": "MultipartFile", "label": "人员文件", "required": True}],
        ),
        api(
            "plan.personnel-import",
            "plan",
            "计划人员",
            3,
            "确认替换人员配置",
            "TF-14",
            "PlanController#confirmPersonnelImport",
            "POST",
            "/ea/plan/plans/{planId}/personnel-import",
            "multipart",
            "admin",
            "ajax-result",
            "8.5.3",
            "重新读取同一文件，摘要与草稿版本均匹配后原子替换人员快照并重建草稿结构。",
            response_java="PlanPersonnelImportResult",
            data_java="PlanPersonnelImportResult",
            path_params=[("planId", "Long", "计划主键")],
            form=[
                {"name": "expectedRevision", "java_type": "Long", "label": "客户端读取到的草稿版本", "required": True},
                {"name": "precheckedContentDigest", "java_type": "String", "label": "预检文件摘要", "required": True},
                {"name": "file", "java_type": "MultipartFile", "label": "人员文件", "required": True},
            ],
        ),
        api(
            "plan.personnel",
            "plan",
            "计划人员",
            3,
            "查询任职清单",
            "TF-15",
            "PlanController#listPersonnel",
            "GET",
            "/ea/plan/plans/{planId}/personnel",
            "none",
            "admin",
            "ajax-result",
            "8.6.1",
            "按计划部门和考核职级过滤，返回当前页任职快照及部门选项。",
            response_java="PlanPersonnelSnapshotVO",
            data_java="PlanPersonnelSnapshotVO",
            query_java="PlanPersonnelQueryDTO",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.roster-workbench",
            "plan",
            "主考核与分管",
            4,
            "查询主考核与分管",
            "TF-16",
            "PlanController#getRosterWorkbench",
            "GET",
            "/ea/plan/plans/{planId}/roster-workbench",
            "none",
            "admin",
            "ajax-result",
            "8.7.1",
            "在同一只读快照中返回全部总监主考核卡片及当前页领导的最终员工范围。",
            response_java="PlanRosterWorkbenchVO",
            data_java="PlanRosterWorkbenchVO",
            query_java="PlanRosterQueryDTO",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.assessment-scope-save",
            "plan",
            "分管范围",
            5,
            "保存分管范围",
            "TF-17",
            "PlanController#saveAssessmentScope",
            "PUT",
            "/ea/plan/plans/{planId}/assessment-scopes/{assessorAssignmentId}",
            "json",
            "admin",
            "ajax-result",
            "8.7.2",
            "显式员工成员集合替换指定领导任职的默认普通员工分管范围，并重建草稿结构。",
            response_java="PlanAssessmentScopeConfigVO",
            data_java="PlanAssessmentScopeConfigVO",
            body_java="PlanAssessmentScopeSaveDTO",
            path_params=[
                ("planId", "Long", "计划主键"),
                ("assessorAssignmentId", "Long", "评分主体任职主键"),
            ],
        ),
        api(
            "plan.readiness",
            "plan",
            "评分结构",
            6,
            "核对发布就绪状态",
            "uncounted",
            "PlanController#inspectReadiness",
            "GET",
            "/ea/plan/plans/{planId}/readiness",
            "none",
            "admin",
            "ajax-result",
            "8.8.1",
            "全计划结构、账号、授权和预期任务守恒的派生查询，作为发布门禁支撑。",
            response_java="PlanReadinessSummaryVO",
            data_java="PlanReadinessSummaryVO",
            path_params=[("planId", "Long", "计划主键")],
            excluded_reason="发布就绪独立查询；已并入门禁与结构核对，不单列。",
        ),
        api(
            "plan.structure-preview",
            "plan",
            "评分结构",
            6,
            "查询评分结构",
            "TF-18",
            "PlanController#previewStructure",
            "GET",
            "/ea/plan/plans/{planId}/structure-preview",
            "none",
            "admin",
            "ajax-result",
            "8.8.2",
            "按考核归属部门分页读取被考核人结构、评分来源、有效权重和缺组转移。",
            response_java="PlanStructurePageVO",
            data_java="PlanStructurePageVO",
            path_params=[("planId", "Long", "计划主键")],
            form=[],
        ),
        api(
            "plan.publish",
            "plan",
            "计划发布",
            7,
            "发布计划",
            "TF-19",
            "PlanController#publish",
            "POST",
            "/ea/plan/plans/{planId}/publish",
            "json",
            "admin",
            "file-stream",
            "8.4.2",
            "校验版本和全计划就绪门禁，发布时激活账号、生成任务、冻结结构并一次性流式交付明文凭证。",
            body_java="PlanPublishDTO",
            path_params=[("planId", "Long", "计划主键")],
            response_kind="file",
            file_kind="xlsx",
        ),
        api(
            "plan.credentials-reset",
            "plan",
            "匿名评分",
            8,
            "批量重置未提交凭证",
            "TF-20",
            "PlanController#resetPendingCredentials",
            "POST",
            "/ea/plan/plans/{planId}/credentials/reset",
            "none",
            "admin",
            "file-stream",
            "8.4.3",
            "仅进行中计划可重置；同一事务重置全部未最终提交账号，撤销会话并一次性交付新凭证。",
            path_params=[("planId", "Long", "计划主键")],
            response_kind="file",
            file_kind="xlsx",
        ),
        api(
            "plan.terminate",
            "plan",
            "计划生命周期",
            9,
            "终止计划",
            "TF-21",
            "PlanController#terminate",
            "POST",
            "/ea/plan/plans/{planId}/terminate",
            "none",
            "admin",
            "ajax-result",
            "8.4.5",
            "仅进行中计划可终止；终止立即阻断后续评分并撤销全部会话，不覆盖账号生命周期事实。",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.complete",
            "plan",
            "计划生命周期",
            9,
            "完成计划",
            "TF-22",
            "PlanController#complete",
            "POST",
            "/ea/plan/plans/{planId}/complete",
            "none",
            "admin",
            "ajax-result",
            "8.4.6",
            "全部应有评分任务正式提交后，原子固化组级和个人汇总并结束计划。",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.rating-progress",
            "plan",
            "评分汇总",
            10,
            "查询评分进度",
            "TF-23",
            "PlanController#ratingProgress",
            "GET",
            "/ea/plan/plans/{planId}/rating-progress",
            "none",
            "admin",
            "ajax-result",
            "8.4.4",
            "返回任务和账号进度、是否满足完成条件及阻断说明。",
            response_java="PlanRatingProgressVO",
            data_java="PlanRatingProgressVO",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.rating-results",
            "plan",
            "评分汇总",
            10,
            "查询评分汇总",
            "TF-24",
            "PlanController#ratingResults",
            "GET",
            "/ea/plan/plans/{planId}/rating-results",
            "none",
            "admin",
            "ajax-result",
            "8.9.1",
            "仅已完成计划可查询；返回已固化汇总分、等次和四个稳定评分组槽位。",
            response_java="PlanRatingResultPageVO",
            data_java="PlanRatingResultPageVO",
            query_java="PlanRatingResultsQueryDTO",
            path_params=[("planId", "Long", "计划主键")],
        ),
        api(
            "plan.rating-summary-detail",
            "plan",
            "评分汇总",
            10,
            "查询评分汇总明细",
            "TF-25",
            "PlanController#ratingSummaryDetail",
            "GET",
            "/ea/plan/plans/{planId}/rating-results/{targetId}",
            "none",
            "admin",
            "ajax-result",
            "8.10.1",
            "读取时校验完整审计链并重算总监员工组组成项均分；不披露账号、席位和人员身份。",
            response_java="PlanRatingSummaryDetailVO",
            data_java="PlanRatingSummaryDetailVO",
            path_params=[
                ("planId", "Long", "计划主键"),
                ("targetId", "Long", "被考核对象主键"),
            ],
        ),
        api(
            "plan.grade-fill-export",
            "plan",
            "评分汇总",
            10,
            "导出年度考核等次填写表",
            "uncounted",
            "PlanController#exportGradeFill",
            "GET",
            "/ea/plan/plans/{planId}/grade-fill-export",
            "none",
            "admin",
            "file-stream",
            "8.9.2",
            "从已固化评分汇总生成固定六列Excel，供线下评定年度考核等次。",
            path_params=[("planId", "Long", "计划主键")],
            response_kind="file",
            file_kind="xlsx",
            excluded_reason="年度考核等次填写表导出；已归入评分汇总查询及导入流程，不单列。",
        ),
        api(
            "plan.grade-import",
            "plan",
            "评分汇总",
            10,
            "导入年度考核等次",
            "TF-26",
            "PlanController#importGrades",
            "POST",
            "/ea/plan/plans/{planId}/grade-import",
            "multipart",
            "admin",
            "ajax-result",
            "8.9.3",
            "完整校验表头、范围、唯一性、只读列和五档等次后，原子更新登记表等次。",
            response_java="PlanGradeImportResult",
            data_java="PlanGradeImportResult",
            path_params=[("planId", "Long", "计划主键")],
            form=[{"name": "file", "java_type": "MultipartFile", "label": "年度考核等次文件", "required": True}],
        ),
        api(
            "plan.roster-check-export",
            "plan",
            "主考核与分管",
            4,
            "导出主考核与分管核对表",
            "TF-27",
            "PlanController#exportRosterCheck",
            "GET",
            "/ea/plan/plans/{planId}/roster-check-export",
            "none",
            "admin",
            "file-stream",
            "8.7.3",
            "从完整一致快照导出人员结构与领导分管两张核对表，不受页面筛选影响。",
            path_params=[("planId", "Long", "计划主键")],
            response_kind="file",
            file_kind="xlsx",
        ),
        api(
            "registration.years",
            "registration",
            "年度考核登记表",
            11,
            "查询可选年度",
            "TF-28",
            "EmployeeRegistrationController#getAvailableYears",
            "GET",
            "/ea/my/employee-registrations/years",
            "none",
            "employee",
            "ajax-result",
            "8.11.1",
            "按当前年度、上一年度和本人历史档案返回可选年度及默认年度。",
            response_java="EmployeeRegistrationYearsVO",
            data_java="EmployeeRegistrationYearsVO",
        ),
        api(
            "registration.detail",
            "registration",
            "年度考核登记表",
            11,
            "查询年度考核登记表",
            "TF-29",
            "EmployeeRegistrationController#getDetail",
            "GET",
            "/ea/my/employee-registrations/{assessmentYear}",
            "none",
            "employee",
            "ajax-result",
            "8.11.2",
            "只允许当前实名员工读取本人数据；无档案时返回不落库的临时空表。",
            response_java="EmployeeRegistrationDetailVO",
            data_java="EmployeeRegistrationDetailVO",
            path_params=[("assessmentYear", "Integer", "考核年度")],
        ),
        api(
            "registration.save-draft",
            "registration",
            "年度考核登记表",
            11,
            "保存登记表草稿",
            "TF-30",
            "EmployeeRegistrationController#saveDraft",
            "PUT",
            "/ea/my/employee-registrations/{assessmentYear}/draft",
            "json",
            "employee",
            "ajax-result",
            "8.11.3",
            "首次保存固化当前员工快照；之后按草稿版本更新本人摘要，不覆盖已导入等次。",
            response_java="EmployeeRegistrationDetailVO",
            data_java="EmployeeRegistrationDetailVO",
            body_java="EmployeeRegistrationDraftDTO",
            path_params=[("assessmentYear", "Integer", "考核年度")],
        ),
        api(
            "registration.submit",
            "registration",
            "年度考核登记表",
            11,
            "提交登记表",
            "TF-31",
            "EmployeeRegistrationController#submit",
            "POST",
            "/ea/my/employee-registrations/{assessmentYear}/submit",
            "json",
            "employee",
            "ajax-result",
            "8.11.4",
            "首次提交必须具有工作摘要；提交后员工填写内容永久只读，网络重试返回已提交事实。",
            response_java="EmployeeRegistrationDetailVO",
            data_java="EmployeeRegistrationDetailVO",
            body_java="EmployeeRegistrationSubmitDTO",
            path_params=[("assessmentYear", "Integer", "考核年度")],
        ),
        api(
            "registration.export",
            "registration",
            "年度考核登记表",
            11,
            "导出登记表Word",
            "TF-32",
            "EmployeeRegistrationController#export",
            "GET",
            "/ea/my/employee-registrations/{assessmentYear}/export",
            "none",
            "employee",
            "file-stream",
            "8.11.5",
            "仅导出当前员工已提交的登记表，并替换客户固定Word模板中的身份、摘要和等次标记。",
            path_params=[("assessmentYear", "Integer", "考核年度")],
            response_kind="file",
            file_kind="docx",
        ),
        api(
            "rating.session.exchange",
            "rating",
            "匿名评分",
            12,
            "匿名登录",
            "TF-33",
            "RatingSessionController#exchange",
            "POST",
            "/open/ea/rating/session/exchange",
            "json",
            "rating-session",
            "ajax-result",
            "8.12.1",
            "校验匿名账号密码，建立独立数据库会话并通过HttpOnly Cookie传递原始令牌。",
            response_java="RatingSessionExchangeResult",
            data_java="RatingSessionExchangeResult",
            body_java="RatingSessionExchangeDTO",
            same_origin=True,
        ),
        api(
            "rating.session.current",
            "rating",
            "匿名评分",
            12,
            "恢复匿名评分工作台",
            "TF-34",
            "RatingSessionController#current",
            "GET",
            "/open/ea/rating/session/current",
            "none",
            "rating-session",
            "ajax-result",
            "8.12.1",
            "从有效评分Cookie恢复账号获授权的全部任务和席位配额边界。",
            response_java="RatingWorkbenchVO",
            data_java="RatingWorkbenchVO",
        ),
        api(
            "rating.task.allocate",
            "rating",
            "匿名评分",
            12,
            "按比例分配",
            "TF-35",
            "RatingDraftController#allocate",
            "POST",
            "/open/ea/rating/tasks/{taskId}/allocation",
            "json",
            "rating-session",
            "ajax-result",
            "8.12.4",
            "按叶子指标满分和最大余数法计算建议分，不写库。",
            response_java="RatingDraftVO",
            data_java="RatingDraftVO",
            body_java="RatingScoreAllocationDTO",
            path_params=[("taskId", "Long", "评分任务主键")],
            same_origin=True,
        ),
        api(
            "rating.task.save-draft",
            "rating",
            "匿名评分",
            12,
            "保存手填评分草稿",
            "TF-36",
            "RatingDraftController#saveDraft",
            "PUT",
            "/open/ea/rating/tasks/{taskId}/draft",
            "json",
            "rating-session",
            "ajax-result",
            "8.12.5",
            "整体保存任务指标分值并校验草稿版本、指标范围和总分完整性。",
            response_java="RatingDraftVO",
            data_java="RatingDraftVO",
            body_java="RatingDraftSaveDTO",
            path_params=[("taskId", "Long", "评分任务主键")],
            same_origin=True,
        ),
        api(
            "rating.seat.get-drafts",
            "rating",
            "匿名评分",
            12,
            "查询席位评分草稿",
            "TF-34",
            "RatingSeatDraftController#getSeatDrafts",
            "GET",
            "/open/ea/rating/seats/{seatId}/drafts",
            "none",
            "rating-session",
            "ajax-result",
            "8.12.2",
            "一次读取席位全部任务草稿和公共指标定义，用于评分矩阵初始化。",
            response_java="RatingSeatDraftVO",
            data_java="RatingSeatDraftVO",
            path_params=[("seatId", "Long", "评分席位主键")],
            data_fields=seat_draft_fields(),
        ),
        api(
            "rating.task.get-draft",
            "rating",
            "匿名评分",
            12,
            "查询评分任务草稿",
            "TF-34",
            "RatingDraftController#getDraft",
            "GET",
            "/open/ea/rating/tasks/{taskId}",
            "none",
            "rating-session",
            "ajax-result",
            "8.12.3",
            "单任务恢复时返回完整指标草稿和当前总分。",
            response_java="RatingDraftVO",
            data_java="RatingDraftVO",
            path_params=[("taskId", "Long", "评分任务主键")],
        ),
        api(
            "rating.session.logout",
            "rating",
            "匿名评分",
            12,
            "退出评分",
            "uncounted",
            "RatingSessionController#logout",
            "POST",
            "/open/ea/rating/session/logout",
            "json",
            "rating-session",
            "ajax-result",
            "8.12.6",
            "撤销当前服务端会话并清除评分Cookie；无令牌时幂等成功。",
            same_origin=True,
            excluded_reason="退出评分；按技术会话动作处理，不单列。",
        ),
        api(
            "rating.submission.employee",
            "rating",
            "匿名评分",
            12,
            "员工匿名账号最终提交",
            "TF-37",
            "RatingSubmissionController#submitEmployee",
            "POST",
            "/open/ea/rating/submission/employee",
            "json",
            "rating-session",
            "ajax-result",
            "8.12.7",
            "统一校验账号全部授权席位任务完整性后转为正式评分、停用账号并撤销会话。",
            response_java="RatingSubmissionResult",
            data_java="RatingSubmissionResult",
            same_origin=True,
        ),
        api(
            "rating.submission.leader",
            "rating",
            "匿名评分",
            12,
            "领导匿名账号最终提交",
            "TF-37",
            "RatingSubmissionController#submitLeader",
            "POST",
            "/open/ea/rating/submission/leader",
            "json",
            "rating-session",
            "ajax-result",
            "8.12.7",
            "统一校验领导账号全部授权席位任务完整性后转为正式评分、停用账号并撤销会话。",
            response_java="RatingSubmissionResult",
            data_java="RatingSubmissionResult",
            same_origin=True,
        ),
    ]
    return specs


def seat_draft_fields() -> list[dict[str, Any]]:
    return [
        {"name": "seatId", "java_type": "Long", "label": "评分席位主键"},
        {
            "name": "indicators",
            "java_type": "List<Indicator>",
            "label": "公共叶子指标定义",
            "fields": [
                {"name": "indicatorId", "java_type": "Long", "label": "叶子指标主键"},
                {"name": "indicatorName", "java_type": "String", "label": "指标名称"},
                {"name": "description", "java_type": "String", "label": "指标说明"},
                {"name": "maxScore", "java_type": "BigDecimal", "label": "指标满分"},
                {"name": "parentId", "java_type": "Long", "label": "所属分组主键"},
                {"name": "parentName", "java_type": "String", "label": "所属分组名称"},
            ],
        },
        {
            "name": "tasks",
            "java_type": "List<Task>",
            "label": "席位全部评分任务",
            "fields": [
                {"name": "taskId", "java_type": "Long", "label": "评分任务主键"},
                {"name": "draftVersion", "java_type": "Long", "label": "草稿版本"},
                {"name": "totalScore", "java_type": "BigDecimal", "label": "评分总分"},
                {
                    "name": "indicators",
                    "java_type": "List<Score>",
                    "label": "全部叶子指标所得分",
                    "fields": [
                        {"name": "indicatorId", "java_type": "Long", "label": "叶子指标主键"},
                        {"name": "score", "java_type": "BigDecimal", "label": "指标所得分"},
                    ],
                },
            ],
        },
    ]


def enrich_endpoints(specs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    layout = [
        ("template.list", "7.2.1.1"),
        ("template.create", "7.2.1.2"),
        ("template.update", "7.2.1.3"),
        ("template.detail", "7.2.1.4"),
        ("template.configure-indicators", "7.2.1.5"),
        ("template.publish", "7.2.1.6"),
        ("template.disable", "7.2.1.7"),
        ("template.copy", "7.2.1.8"),
        ("template.delete", "7.2.1.9"),
        ("plan.list", "7.2.2.1"),
        ("plan.create", "7.2.2.2"),
        ("plan.delete", "7.2.2.3"),
        ("plan.update-draft", "7.2.3.1"),
        ("plan.readiness", "7.2.3.2"),
        ("plan.publish", "7.2.3.3"),
        ("plan.credentials-reset", "7.2.3.4"),
        ("plan.terminate", "7.2.3.5"),
        ("plan.complete", "7.2.3.6"),
        ("plan.rating-progress", "7.2.3.7"),
        ("plan.personnel-import-template", "7.2.4.1"),
        ("plan.personnel-precheck", "7.2.4.2"),
        ("plan.personnel-import", "7.2.4.3"),
        ("plan.personnel", "7.2.5.1"),
        ("plan.roster-workbench", "7.2.6.1"),
        ("plan.assessment-scope-save", "7.2.6.2"),
        ("plan.roster-check-export", "7.2.6.3"),
        ("plan.structure-preview", "7.2.7.1"),
        ("plan.rating-results", "7.2.8.1"),
        ("plan.grade-fill-export", "7.2.8.2"),
        ("plan.grade-import", "7.2.8.3"),
        ("plan.rating-summary-detail", "7.2.9.1"),
        ("rating.session.exchange", "7.2.10.1"),
        ("rating.session.current", "7.2.10.2"),
        ("rating.session.logout", "7.2.10.3"),
        ("rating.task.get-draft", "7.2.10.4"),
        ("rating.seat.get-drafts", "7.2.10.4"),
        ("rating.task.allocate", "7.2.10.5"),
        ("rating.task.save-draft", "7.2.10.6"),
        ("rating.submission.employee", "7.2.10.7"),
        ("rating.submission.leader", "7.2.10.7"),
        ("registration.years", "7.2.11.1"),
        ("registration.detail", "7.2.11.2"),
        ("registration.save-draft", "7.2.11.3"),
        ("registration.submit", "7.2.11.4"),
        ("registration.export", "7.2.11.5"),
    ]
    order = {api_id: idx for idx, (api_id, _) in enumerate(layout, 1)}
    sections = dict(layout)
    specs = sorted(specs, key=lambda item: order[item["id"]])
    for idx, spec in enumerate(specs, 1):
        if spec["nesma"] != "uncounted":
            setattr_stage = "计数"
        else:
            setattr_stage = "排除"
        spec["table"] = f"7-{idx}"
        spec["section"] = sections[spec["id"]]
        spec["ordinal"] = idx
        spec["stage"] = setattr_stage
        if spec.get("body_java"):
            spec["body"] = field_specs(spec["body_java"])
            if spec["body_java"] == "Template":
                spec["body"] = [
                    item
                    for item in spec["body"]
                    if item["name"] in {"templateName", "description"}
                ]
        if spec.get("query_java"):
            spec["query"] = field_specs(spec["query_java"])
            for item in spec["query"]:
                if item["name"] in {"pageNum", "pageSize"}:
                    item["platform"] = True
        if spec.get("data_fields"):
            spec["data"] = spec["data_fields"]
        elif spec.get("data_java"):
            spec["data"] = field_specs(spec["data_java"])
        else:
            spec["data"] = []
        if spec["path"] == "/ea/plan/plans/{planId}/structure-preview":
            spec["query"] = [
                {"name": "departmentId", "java_type": "Long", "label": "计划部门主键"},
                {"name": "pageNum", "java_type": "int", "label": "页码", "platform": True},
                {"name": "pageSize", "java_type": "int", "label": "每页条数", "platform": True},
            ]
        if spec["data_java"] == "Template":
            spec["data"] = [
                x
                for x in spec["data"]
                if x["name"]
                not in {"status", "enabled", "createdBy", "createdAt", "updatedBy", "updatedAt"}
            ]
        if spec["data_java"] == "RatingSessionExchangeResult":
            spec["data"] = [
                x for x in spec["data"] if x["name"] != "sessionToken"
            ]
        spec["yaml_rel"] = f"http/{spec['id']}.yaml"
        spec["image_name"] = f"api-{spec['id']}.png"
        spec["image_rel"] = f"images/{spec['image_name']}"
    return specs


def parse_db_tables() -> list[dict[str, Any]]:
    sql_path = BACKEND / "sql/assessment-v2/01-rebuild-schema.sql"
    source = sql_path.read_text(encoding="utf-8")
    tables = []
    for match in re.finditer(
        r"CREATE TABLE\s+(\w+)\s*\((.*?)\)\s*COMMENT\s*=",
        source,
        flags=re.S | re.I,
    ):
        name, body = match.group(1), match.group(2)
        columns = []
        for line in body.splitlines():
            line = line.strip().rstrip(",")
            if line.lower().startswith(
                ("primary ", "unique ", "key ", "index ", "constraint ")
            ):
                break
            if not line:
                continue
            col_match = re.match(
                r"`?(\w+)`?\s+([a-zA-Z]+(?:\([^)]+\))?)\s*(.*)", line
            )
            if not col_match:
                continue
            col_name, col_type, rest = col_match.groups()
            comment = ""
            comment_match = re.search(r"COMMENT\s+'([^']*)'", rest, re.I)
            if comment_match:
                comment = comment_match.group(1)
            columns.append(
                {
                    "name": col_name,
                    "type": col_type.lower(),
                    "required": "NOT NULL" in rest.upper(),
                    "desc": comment or title_for_db_column(col_name),
                }
            )
        tables.append(
            {
                "id": f"data.{name}",
                "name": name,
                "title": title_for_db_table(name),
                "columns": columns,
                "yaml_rel": f"data/{name}.yaml",
                "image_name": f"data-{name}.png",
                "image_rel": f"images/data-{name}.png",
            }
        )
    return tables


DB_TITLES = {
    "ea_template": "考核模板",
    "ea_template_indicator": "模板指标",
    "ea_plan_year_guard": "计划年度守卫",
    "ea_plan": "年度考核计划",
    "ea_plan_person": "计划人员快照",
    "ea_plan_department": "计划部门快照",
    "ea_plan_assignment": "计划任职快照",
    "ea_plan_department_weight": "部门职级权重",
    "ea_plan_primary_assessment": "总监主考核BU",
    "ea_plan_primary_roster": "主考核正副职名单",
    "ea_plan_assessment_scope": "领导普通员工分管范围",
    "ea_plan_assessment_scope_config": "分管范围显式配置标记",
    "ea_assessment_target": "被考核人",
    "ea_plan_indicator": "计划指标快照",
    "ea_rating_group": "评分组",
    "ea_rating_seat": "评分席位",
    "ea_rating_relation": "评分关系",
    "ea_anonymous_account": "匿名评分账号",
    "ea_account_seat_grant": "账号席位授权",
    "ea_rating_task": "评分任务",
    "ea_rating_session": "评分会话",
    "ea_rating_record": "指标所得分",
    "ea_rating_group_result": "评分组汇总",
    "ea_rating_result": "被考核人汇总分",
    "ea_employee_registration": "年度考核登记表",
    "ea_operation_log": "员工考核操作日志",
}


def title_for_db_table(name: str) -> str:
    return DB_TITLES.get(name, name)


def title_for_db_column(name: str) -> str:
    mapping = {
        "id": "主键",
        "plan_id": "计划主键",
        "template_id": "模板主键",
        "assessment_year": "考核年度",
        "status": "状态",
        "draft_revision": "草稿版本",
        "created_by": "创建人主键",
        "created_at": "创建时间",
        "updated_by": "最后修改人主键",
        "updated_at": "最后修改时间",
        "published_at": "发布时间",
        "completed_at": "完成时间",
        "terminated_at": "终止时间",
        "employee_code": "员工编号",
        "employee_id": "主系统人员标识",
        "person_name": "姓名",
        "department_name": "考核部门名称",
        "position_type": "考核职级",
        "weight": "权重",
        "plan_person_id": "计划人员主键",
        "plan_department_id": "计划部门主键",
        "assessment_grade": "年度考核等次",
        "annual_work_summary": "年度主要工作摘要",
        "submitted_at": "首次提交时间",
    }
    return mapping.get(name, name.replace("_", " "))


def data_function_db_map() -> dict[str, list[str]]:
    return {
        "DF-01": ["data.ea_plan_person", "data.ea_plan_department", "data.ea_plan_assignment"],
        "DF-02": ["data.ea_plan_department_weight"],
        "DF-03": ["data.ea_plan_primary_assessment", "data.ea_plan_primary_roster"],
        "DF-04": ["data.ea_template", "data.ea_template_indicator"],
        "DF-05": ["data.ea_plan", "data.ea_plan_indicator"],
        "DF-06": ["data.ea_assessment_target", "data.ea_rating_group", "data.ea_rating_relation", "data.ea_rating_seat"],
        "DF-07": ["data.ea_anonymous_account", "data.ea_account_seat_grant"],
        "DF-08": ["data.ea_plan_assessment_scope", "data.ea_plan_assessment_scope_config"],
        "DF-09": ["data.ea_rating_task", "data.ea_rating_record"],
        "DF-10": ["data.ea_rating_group_result", "data.ea_rating_result"],
        "DF-11": ["data.ea_employee_registration"],
        "DF-12": ["master-data.current-employee", "master-data.import-validate"],
        "DF-13": ["master-data.current-employee", "master-data.import-validate"],
    }


def feature_rows(model: dict[str, Any], specs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    api_by_tf: dict[str, list[str]] = {}
    for spec in specs:
        if isinstance(spec["nesma"], str) and spec["nesma"].startswith("TF-"):
            api_by_tf.setdefault(spec["nesma"], []).append(spec["id"])
    page_map = {
        "TF-01": ("7.2.1.2", "考核模板页面"),
        "TF-02": ("7.2.1.3", "考核模板页面"),
        "TF-03": ("7.2.1.5", "考核模板页面"),
        "TF-04": ("7.2.1.6", "考核模板页面"),
        "TF-05": ("7.2.1.7", "考核模板页面"),
        "TF-06": ("7.2.1.8", "考核模板页面"),
        "TF-07": ("7.2.1.9", "考核模板页面"),
        "TF-08": ("7.2.1.1", "考核模板页面"),
        "TF-09": ("7.2.2.2", "考核计划列表"),
        "TF-10": ("7.2.3.1", "计划工作台—概览"),
        "TF-11": ("7.2.2.3", "考核计划列表"),
        "TF-12": ("7.2.2.1", "考核计划列表"),
        "TF-13": ("7.2.4.2", "计划工作台—人员操作"),
        "TF-14": ("7.2.4.3", "计划工作台—人员操作"),
        "TF-15": ("7.2.5.1", "计划工作台—任职清单"),
        "TF-16": ("7.2.6.1", "计划工作台—主考核与分管"),
        "TF-17": ("7.2.6.2", "计划工作台—主考核与分管"),
        "TF-18": ("7.2.7.1", "计划工作台—结构预览"),
        "TF-19": ("7.2.3.3", "计划工作台—概览"),
        "TF-20": ("7.2.3.4", "计划工作台—概览"),
        "TF-21": ("7.2.3.5", "计划工作台—概览"),
        "TF-22": ("7.2.3.6", "计划工作台—概览"),
        "TF-23": ("7.2.3.7", "计划工作台—概览"),
        "TF-24": ("7.2.8.1", "计划工作台—评分汇总"),
        "TF-25": ("7.2.9.1", "评分汇总审计明细"),
        "TF-26": ("7.2.8.3", "计划工作台—评分汇总"),
        "TF-27": ("7.2.6.3", "计划工作台—主考核与分管"),
        "TF-28": ("7.2.11.1", "年度考核登记表"),
        "TF-29": ("7.2.11.2", "年度考核登记表"),
        "TF-30": ("7.2.11.3", "年度考核登记表"),
        "TF-31": ("7.2.11.4", "年度考核登记表"),
        "TF-32": ("7.2.11.5", "年度考核登记表"),
        "TF-33": ("7.2.10.1", "匿名评分工作台"),
        "TF-34": ("7.2.10.4", "匿名评分工作台"),
        "TF-35": ("7.2.10.5", "匿名评分工作台"),
        "TF-36": ("7.2.10.6", "匿名评分工作台"),
        "TF-37": ("7.2.10.7", "匿名评分工作台"),
    }
    primary_module = {
        "考核模板": "考核模板管理",
        "年度考核计划": "年度考核计划",
        "计划人员": "计划人员管理",
        "总监主考核": "主考核与分管",
        "分管范围": "主考核与分管",
        "评分结构": "评分结构管理",
        "匿名评分": "匿名评分",
        "评分汇总": "评分汇总与等次",
        "年度考核登记表": "年度考核登记表",
    }
    description_map = {
        "TF-13": "上传人员任职双表文件，只读核对人员身份、部门职级、权重与总监主考核，并预览替换影响。",
        "TF-14": "确认同一份已核对文件，原子替换计划人员、部门、任职、权重及主考核配置，并重建评分结构。",
        "TF-16": "按完整一致快照查看总监主考核名单、领导普通员工分管范围及来源。",
        "TF-19": "完成发布门禁后激活匿名账号、生成评分任务、冻结评分结构并一次性交付凭证。",
        "TF-20": "对进行中计划全部未最终提交账号批量重置密码、撤销会话并重新交付凭证。",
        "TF-25": "按匿名序号展示一名被考核人的完整评分审计链、组内均分、组加权分和组成项。",
        "TF-33": "使用匿名账号密码进入独立评分入口，只看到账号获授权席位的任务。",
        "TF-34": "读取账号或席位授权范围内的任务清单、指标定义、草稿分值和评分配额边界。",
        "TF-35": "输入期望总分，按最大余数法生成各叶子指标建议分值，不写库。",
        "TF-37": "账号全部授权席位填写完整后一次性提交，提交后不可撤回或修改。",
    }
    rows = []
    for idx, item in enumerate(model["transactions"], 1):
        tf = item["id"]
        section, page = page_map[tf]
        module = primary_module.get(item["module"], item["module"])
        rows.append(
            {
                "row_id": f"F-{idx:02d}",
                "system": "员工年度考核",
                "module_1": module,
                "module_2": item["name"],
                "description": description_map.get(tf, item["processing"]),
                "external": external_description(tf),
                "change": "新增",
                "remark": f"NESMA：{tf}（{item['type']}，{item['fp']} FP）；数据功能：{'、'.join(item['ftrs'])}。",
                "tf": tf,
                "df": item["ftrs"],
                "section": section,
                "page": page,
                "apis": api_by_tf.get(tf, []),
            }
        )
    return rows


def external_description(tf: str) -> str:
    if tf in {"TF-13", "TF-14"}:
        return "读取主系统人员标识；不向主系统写回数据。"
    if tf in {"TF-12", "TF-28", "TF-29", "TF-30", "TF-31", "TF-32"}:
        return "读取主系统实名身份、人员及部门信息；不向主系统写回数据。"
    return "不涉及外部系统接口。"


DATA_MODEL_SECTIONS = {
    "DF-01": "7.1.1.3 组织人员导入快照",
    "DF-02": "7.1.1.4 部门职级权重",
    "DF-03": "7.1.1.5 总监主考核",
    "DF-04": "7.1.1.1 考核模板",
    "DF-05": "7.1.1.2 年度考核计划及计划指标快照",
    "DF-06": "7.1.1.7 被考核对象与评分结构",
    "DF-07": "7.1.2.1 匿名评分账号与席位授权",
    "DF-08": "7.1.1.6 领导普通员工分管范围",
    "DF-09": "7.1.2.2 评分任务与指标所得分",
    "DF-10": "7.1.3.1 评分汇总",
    "DF-11": "7.1.3.2 年度考核登记表",
    "DF-12": "7.1.4.2 主系统人员",
    "DF-13": "7.1.4.3 主系统部门",
}


def mapping_rows(
    model: dict[str, Any], rows: list[dict[str, Any]], specs: list[dict[str, Any]], dbs: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    tf_rows = {r["tf"]: r for r in rows}
    df_db = data_function_db_map()
    df_usage: dict[str, list[str]] = {}
    for row in rows:
        for df in row["df"]:
            df_usage.setdefault(df, []).append(row["row_id"])
    result = []
    for item in model["data_functions"]:
        df = item["id"]
        interfaces = df_db.get(df, [])
        result.append(
            {
                "id": df,
                "type": item["type"],
                "name": item["name"],
                "fp": item["fp"],
                "feature_rows": "、".join(df_usage.get(df, [])),
                "section_page_model": f"{DATA_MODEL_SECTIONS[df]}；数据模型「{item['name']}」",
                "interfaces": "、".join(interfaces),
                "status": "已覆盖",
                "note": f"RET/DET：{', '.join(item['rets'])} / {', '.join(item['dets'])}。",
            }
        )
    for item in model["transactions"]:
        tf = item["id"]
        row = tf_rows[tf]
        result.append(
            {
                "id": tf,
                "type": item["type"],
                "name": item["name"],
                "fp": item["fp"],
                "feature_rows": row["row_id"],
                "section_page_model": f"{row['section']} {row['page']}；功能清单「{row['module_2']}」",
                "interfaces": "、".join(row["apis"]),
                "status": "已覆盖",
                "note": item["processing"],
            }
        )
    return result


def xml_safe(text: Any) -> str:
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_xlsx_outputs(
    model: dict[str, Any],
    rows: list[dict[str, Any]],
    mappings: list[dict[str, Any]],
) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    thin = Side(style="thin", color="BFC8C3")
    header_fill = PatternFill("solid", fgColor="DDE7E1")
    header_font = Font(name="宋体", size=11, bold=True, color="24312E")
    body_font = Font(name="宋体", size=10, color="24312E")
    wrap = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "NESMA映射"
    headers = [
        "NESMA编号",
        "类型",
        "名称",
        "FP",
        "对应功能清单行",
        "对应正文章节/页面/数据模型",
        "对应接口YAML/卡片",
        "覆盖状态",
        "说明",
    ]
    ws.append(headers)
    for item in mappings:
        ws.append(
            [
                item["id"],
                item["type"],
                item["name"],
                item["fp"],
                item["feature_rows"],
                item["section_page_model"],
                item["interfaces"],
                item["status"],
                item["note"],
            ]
        )
    widths = [12, 8, 30, 8, 28, 48, 42, 12, 60]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("覆盖汇总")
    summary.append(["项目", "值", "说明"])
    metrics = [
        ("数据功能", len(model["data_functions"]), "11 ILF + 2 ELF"),
        ("事务功能", len(model["transactions"]), "22 EI + 10 EO + 5 EQ"),
        ("数据功能FP", sum(x["fp"] for x in model["data_functions"]), "明细复算"),
        ("事务功能FP", sum(x["fp"] for x in model["transactions"]), "明细复算"),
        (
            "UFP总数",
            sum(x["fp"] for x in model["data_functions"])
            + sum(x["fp"] for x in model["transactions"]),
            "数据功能FP + 事务功能FP",
        ),
        ("未覆盖", sum(1 for x in mappings if x["status"] != "已覆盖"), "应为0"),
    ]
    for metric in metrics:
        summary.append(metric)
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 42
    wb.save(OUT / "映射表.xlsx")

    wb2 = load_workbook(TEMPLATE_XLSX)
    ws2 = wb2.active
    data_styles = [copy(ws2.cell(2, col)._style) for col in range(1, 9)]
    for merged in list(ws2.merged_cells.ranges):
        ws2.unmerge_cells(str(merged))
    if ws2.max_row > 1:
        ws2.delete_rows(2, ws2.max_row - 1)
    for idx, row in enumerate(rows, 1):
        values = (
            [
                idx,
                row["system"],
                row["module_1"],
                row["module_2"],
                row["description"],
                row["external"],
                row["change"],
                row["remark"],
            ]
        )
        excel_row = idx + 1
        for col, value in enumerate(values, 1):
            cell = ws2.cell(excel_row, col, value)
            cell._style = copy(data_styles[col - 1])
            cell.alignment = Alignment(
                horizontal="center" if col in {1, 7} else "left",
                vertical="top",
                wrap_text=True,
            )
        ws2.row_dimensions[excel_row].height = 72

    def merge_groups(column: int) -> None:
        start = 2
        while start <= ws2.max_row:
            value = ws2.cell(start, column).value
            end = start
            while end + 1 <= ws2.max_row and ws2.cell(end + 1, column).value == value:
                end += 1
            if end > start:
                ws2.merge_cells(
                    start_row=start,
                    start_column=column,
                    end_row=end,
                    end_column=column,
                )
                ws2.cell(start, column).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            start = end + 1

    merge_groups(2)
    merge_groups(3)
    ws2.freeze_panes = "A2"
    ws2.sheet_view.showGridLines = False
    wb2.save(OUT / "功能清单.xlsx")


def yaml_record(spec: dict[str, Any]) -> OrderedDict[str, Any]:
    rec: OrderedDict[str, Any] = OrderedDict()
    for key in [
        "id",
        "kind",
        "resource",
        "title",
        "nesma",
        "table",
        "section",
        "controller",
        "method",
        "path",
        "consumes",
        "auth",
    ]:
        if key in spec:
            rec[key] = spec[key]
    if spec.get("same_origin"):
        rec["same_origin"] = True
    rec["envelope"] = spec["envelope"]
    if spec.get("response_java"):
        rec["response_java"] = spec["response_java"]
    if spec["response_kind"] == "file":
        rec["file"] = {
            "kind": spec["file_kind"],
            "name_pattern": spec["file_kind"],
        }
    if spec.get("path_params"):
        rec["path_params"] = [
            {"name": n, "java_type": t, "label": label, "required": True}
            for n, t, label in spec["path_params"]
        ]
    if spec.get("query_java"):
        rec["query"] = spec.get("query", [])
    if spec.get("body_java"):
        rec["body_java"] = spec["body_java"]
        rec["body"] = spec.get("body", [])
    if spec.get("form"):
        rec["form"] = spec["form"]
    if spec.get("data"):
        rec["data"] = spec["data"]
    rec["rule"] = spec["rule"]
    if spec.get("excluded_reason"):
        rec["notes"] = spec["excluded_reason"]
    return rec


def to_plain(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: to_plain(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [to_plain(item) for item in value]
    return value


def write_yaml(
    model: dict[str, Any],
    specs: list[dict[str, Any]],
    dbs: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    api_dir = OUT / "接口表"
    shutil.rmtree(api_dir, ignore_errors=True)
    (api_dir / "http").mkdir(parents=True)
    (api_dir / "data").mkdir(parents=True)
    (api_dir / "images").mkdir(parents=True)
    (api_dir / "README.md").write_text(
        """# 接口表编辑面

本目录与 `需求说明书.docx` 同批交付。接口事实来自后端固定提交
`c761f4f529a80e2972436292f1aef35828fcf5a9`，计数编号只取同批
`计数模型.json`。旧 `api-tables` 仅用于版式参考，不是本目录事实来源。

## 目录

- `http/*.yaml`：45 个 HTTP 控制器方法、2 个主系统只读契约。
- `data/*.yaml`：26 个当前基线物理表字段卡片。
- `tex/api-card.sty`：与旧需求说明书一致的接口卡排版宏。
- `tex/cards/*.tex`：由 YAML 生成的 TeX 卡片源。
- `images/*.png`：从 TeX 卡片渲染的 73 张接口/数据图片，供 Word 插图使用。

## 记录约定

- HTTP 卡片包含 `id`、`kind`、`resource`、`title`、`nesma`、`controller`、
  `method`、`path`、`consumes`、`auth`、`envelope` 和参数/字段。
- `nesma` 只能引用当前 `DF-01` 至 `DF-13`、`TF-01` 至 `TF-37`，或明确标记
  `uncounted`。`uncounted` 表示按阶段一决定不单列的支撑行为。
- 成功信封只有三种：`ajax-result`、`table-data-info`、`file-stream`。
  文件下载不虚构 JSON；匿名评分令牌只通过 Cookie 传输，不出现在响应字段中。
- 一张业务事务可以对应多个调用入口。例如 `TF-34` 合并工作台、任务和席位读取，
  `TF-37` 含员工和领导两个最终提交入口；计数仍为阶段一的一条事务功能。
- `table` 是需求说明书接口表号，采用 `7-n`；物理表采用 `7.1-n`。
- `section` 是接口所在叶子节号，位于正文 `7.2 界面功能模块` 下。

编辑 YAML 后，应同步重渲染 `images` 并更新 Word 中同号接口表；不能只改其中一侧。
""",
        encoding="utf-8",
    )

    for spec in specs:
        path = api_dir / spec["yaml_rel"]
        path.write_text(
            yaml.safe_dump(
                to_plain(yaml_record(spec)),
                allow_unicode=True,
                sort_keys=False,
                width=100,
            ),
            encoding="utf-8",
        )

    for db in dbs:
        rec = OrderedDict(
            [
                ("id", db["id"]),
                ("name", db["name"]),
                ("title", db["title"]),
                ("columns", db["columns"]),
            ]
        )
        (api_dir / db["yaml_rel"]).write_text(
            yaml.safe_dump(to_plain(rec), allow_unicode=True, sort_keys=False, width=100),
            encoding="utf-8",
        )

    master_records = [
        OrderedDict(
            [
                ("id", "master-data.import-validate"),
                ("kind", "elf"),
                ("resource", "master-data"),
                ("title", "主系统人员与部门共享数据读取"),
                ("nesma", ["DF-12", "DF-13"]),
                ("table", f"7-{len(specs) + 1}"),
                ("section", "7.2.12.1"),
                (
                    "data",
                    [
                        {"name": "employeeCode", "java_type": "String", "label": "员工编号"},
                        {"name": "employeeId", "java_type": "String", "label": "主系统人员标识"},
                        {"name": "personName", "java_type": "String", "label": "姓名"},
                        {"name": "status", "java_type": "String", "label": "人员状态"},
                        {"name": "departmentName", "java_type": "String", "label": "考核部门名称"},
                    ],
                ),
                (
                    "rule",
                    "计划人员导入前只读校验人员与部门；任一标识无效时整次导入零写入。",
                ),
            ]
        ),
        OrderedDict(
            [
                ("id", "master-data.current-employee"),
                ("kind", "elf"),
                ("resource", "master-data"),
                ("title", "当前实名人员与部门读取"),
                ("nesma", ["DF-12", "DF-13"]),
                ("table", f"7-{len(specs) + 2}"),
                ("section", "7.2.12.2"),
                (
                    "data",
                    [
                        {"name": "employeeId", "java_type": "String", "label": "员工编号"},
                        {"name": "personName", "java_type": "String", "label": "姓名"},
                        {"name": "organizationName", "java_type": "String", "label": "现单位"},
                        {"name": "departmentName", "java_type": "String", "label": "现部门"},
                    ],
                ),
                (
                    "rule",
                    "登记表查询、保存和提交读取当前实名身份的人员及部门；字段缺失时拒绝操作。",
                ),
            ]
        ),
    ]
    for idx, rec in enumerate(master_records, 1):
        (api_dir / "http" / f"{rec['id']}.yaml").write_text(
            yaml.safe_dump(to_plain(rec), allow_unicode=True, sort_keys=False, width=100),
            encoding="utf-8",
        )
        rec["image_name"] = f"api-{rec['id']}.png"
        rec["image_rel"] = f"images/{rec['image_name']}"
        rec["ordinal"] = len(specs) + idx
        rec["category"] = "主系统数据接口"
        rec["category_no"] = 13
        rec["kind"] = "elf"
        rec["method"] = "READ"
        rec["path"] = "主系统共享数据"
        rec["auth"] = "系统内部只读"
        rec["envelope"] = "elf"
        rec["data_fields"] = rec["data"]
    return master_records


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    text = str(text or "")
    lines: list[str] = []
    for para in text.splitlines() or [""]:
        current = ""
        for char in para:
            candidate = current + char
            if draw.textbbox((0, 0), candidate, font=font)[2] <= width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = char
        lines.append(current)
    return lines or [""]


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    width: int,
    *,
    line_gap: int = 10,
) -> int:
    x, y = xy
    lines = wrap_text(draw, text, font, width)
    line_h = draw.textbbox((0, 0), "中Ag", font=font)[3] + line_gap
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_h
    return y


def render_card_image(
    *,
    title: str,
    subtitle: str,
    method: str,
    path: str,
    meta: list[tuple[str, str]],
    sections: list[tuple[str, list[dict[str, Any]]]],
    output: Path,
) -> None:
    width = 1680
    margin = 60
    header_h = 250
    font_title = ImageFont.truetype(FONT_BOLD, 48)
    font_sub = ImageFont.truetype(FONT_REGULAR, 25)
    font_method = ImageFont.truetype(FONT_BOLD, 25)
    font_body = ImageFont.truetype(FONT_REGULAR, 25)
    font_small = ImageFont.truetype(FONT_REGULAR, 22)
    font_bold = ImageFont.truetype(FONT_BOLD, 24)

    # Draw on a generous canvas first, then crop to the actual rendered height.
    canvas_height = 6000
    img = Image.new("RGB", (width, canvas_height), COLORS["white"])
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, header_h), fill=COLORS["primary"])
    draw.rectangle((0, 0, 18, canvas_height), fill=COLORS["accent"])
    draw.text((margin, 42), title, font=font_title, fill=COLORS["white"])
    draw.text((margin, 112), subtitle, font=font_sub, fill="#DBE5E0")
    draw.rounded_rectangle(
        (margin, 170, width - margin, 226),
        radius=12,
        outline="#759087",
        width=2,
        fill="#2D433E",
    )
    badge_w = 156 if method not in {"READ", "ELF"} else 150
    draw.text((margin + 16, 183), method, font=font_method, fill="#F1D7A9")
    draw.text((margin + badge_w, 183), path, font=font_method, fill=COLORS["white"])

    y = header_h + 32
    for key, value in meta:
        draw.text((margin, y), key, font=font_small, fill=COLORS["secondary"])
        y = draw_wrapped(
            draw,
            (margin + 170, y),
            value,
            font_small,
            COLORS["dark"],
            width - 2 * margin - 170,
            line_gap=6,
        )
        y += 10
    y += 14

    for heading, fields in sections:
        draw.rectangle((margin, y, width - margin, y + 48), fill=COLORS["panel"])
        draw.text((margin + 16, y + 10), heading, font=font_bold, fill=COLORS["primary"])
        y += 58
        if not fields:
            draw.text((margin + 20, y), "无", font=font_body, fill=COLORS["mid"])
            y += 52
        for f in fields:
            name = f.get("name", "")
            typ = f.get("java_type", "")
            label = f.get("label", "")
            required = " *" if f.get("required") else ""
            platform = " [平台分页]" if f.get("platform") else ""
            text = f"{name}  ·  {typ}  ·  {label}{required}{platform}"
            draw.ellipse((margin + 18, y + 12, margin + 27, y + 21), fill=COLORS["accent"])
            y = draw_wrapped(
                draw,
                (margin + 40, y),
                text,
                font_small,
                COLORS["dark"],
                width - 2 * margin - 60,
                line_gap=7,
            )
            y += 8
            if f.get("fields"):
                child_names = "、".join(x.get("name", "") for x in f["fields"])
                y = draw_wrapped(
                    draw,
                    (margin + 68, y),
                    f"嵌套字段：{child_names}",
                    font_small,
                    COLORS["mid"],
                    width - 2 * margin - 88,
                    line_gap=5,
                )
                y += 7
        y += 12
        draw.line((margin, y, width - margin, y), fill=COLORS["border"], width=2)
        y += 22

    final_height = max(1000, min(canvas_height, y + 100))
    img = img.crop((0, 0, width, final_height))
    draw = ImageDraw.Draw(img)
    draw.text(
        (margin, final_height - 42),
        "员工年度考核 v3-new-v2 · 由接口YAML与固定代码基线生成",
        font=font_small,
        fill=COLORS["light"],
    )
    img.save(output, optimize=True)


def image_for_api(spec: dict[str, Any]) -> None:
    meta = [
        ("接口编号", spec["id"]),
        ("NESMA", str(spec["nesma"])),
        ("正文位置", f"{spec['section']} / 表 {spec['table']}"),
        ("控制器", spec.get("controller", "主系统只读数据接口")),
        ("鉴权", spec.get("auth", "系统内部只读")),
        ("响应", response_text(spec)),
    ]
    if spec.get("excluded_reason"):
        meta.append(("计数说明", spec["excluded_reason"]))
    sections = []
    if spec.get("path_params"):
        sections.append(
            (
                "路径参数",
                [
                    {"name": n, "java_type": t, "label": label, "required": True}
                    for n, t, label in spec["path_params"]
                ],
            )
        )
    if spec.get("query"):
        sections.append(("查询参数", spec["query"]))
    if spec.get("body"):
        sections.append(("请求体", spec["body"]))
    if spec.get("form"):
        sections.append(("表单字段", spec["form"]))
    if spec.get("data") or spec.get("data_fields"):
        sections.append(("成功响应业务字段", spec.get("data") or spec.get("data_fields", [])))
    sections.append(("调用规则", [{"name": "rule", "java_type": "String", "label": spec["rule"]}]))
    render_card_image(
        title=spec["title"],
        subtitle=f"{spec['category']} · 表 {spec['table']}",
        method=spec.get("method", "READ"),
        path=spec.get("path", "主系统共享数据"),
        meta=meta,
        sections=sections,
        output=OUT / "接口表" / spec["image_rel"],
    )


def response_text(spec: dict[str, Any]) -> str:
    if spec.get("envelope") == "file-stream":
        return f"文件流（{spec.get('file_kind', 'file')}）"
    if spec.get("envelope") == "table-data-info":
        return "TableDataInfo：code / msg / total / rows"
    if spec.get("envelope") == "ajax-result":
        return "AjaxResult：code / msg / data"
    return "自然语言业务事实"


def image_for_db(db: dict[str, Any]) -> None:
    fields = [
        {
            "name": c["name"],
            "java_type": c["type"],
            "label": c["desc"],
            "required": c["required"],
        }
        for c in db["columns"]
    ]
    render_card_image(
        title=db["title"],
        subtitle="数据模型 / 物理表",
        method="TABLE",
        path=db["name"],
        meta=[
            ("数据卡片", db["id"]),
            ("字段数", str(len(db["columns"]))),
            ("来源", "assessment-v2/01-rebuild-schema.sql"),
            ("用途", "需求说明书数据模型支撑"),
        ],
        sections=[("字段", fields)],
        output=OUT / "接口表" / db["image_rel"],
    )


def tex_escape(value: Any) -> str:
    text = str(value or "")
    text = re.sub(r"\s+", " ", text)
    text = text.replace("\\", r"\textbackslash{}")
    text = text.replace("&", r"\&")
    text = text.replace("%", r"\%")
    text = text.replace("$", r"\$")
    text = text.replace("#", r"\#")
    text = text.replace("{", r"\{")
    text = text.replace("}", r"\}")
    text = text.replace("~", r"\textasciitilde{}")
    text = text.replace("^", r"\textasciicircum{}")
    text = text.replace("<", r"\textless{}")
    text = text.replace(">", r"\textgreater{}")
    text = text.replace("_", r"\_")
    return text


def tex_name(value: Any) -> str:
    return tex_escape(value).replace(r"\_", r"\_\allowbreak{}")


def tex_field_rows(fields: list[dict[str, Any]]) -> str:
    lines = []
    for idx, field in enumerate(fields):
        macro = r"\EaFieldFirst" if idx == 0 else r"\EaField"
        name = tex_name(field.get("name", ""))
        label = tex_escape(field.get("label", ""))
        desc = tex_escape(field.get("desc", ""))
        if field.get("required"):
            desc = ("必填。" + desc) if desc else "必填"
        if field.get("platform"):
            desc = ("平台分页。" + desc) if desc else "平台分页"
        lines.append(f"{macro}{{{name}}}{{{label}}}{{{desc}}}")
        children = field.get("fields") or []
        if children:
            lines.append(r"\EaKids{%")
            lines.extend(tex_field_rows(children).splitlines())
            lines.append("}")
    return "\n".join(lines)


def tex_http_card(spec: dict[str, Any]) -> str:
    file_info = spec.get("file") or {}
    parts = [
        rf"\EaHttpCard{{{spec['id']}}}{{%",
        f"  title={{{tex_escape(spec['title'])}}},",
        f"  method={{{tex_escape(spec['method'])}}},",
        f"  path={{{tex_escape(spec['path'])}}},",
        f"  table={{{tex_escape(spec['table'])}}},",
        f"  auth={{{tex_escape(spec['auth'])}}},",
        f"  sameorigin={{{'true' if spec.get('same_origin') else 'false'}}},",
        f"  envelope={{{tex_escape(spec['envelope'])}}},",
    ]
    if spec.get("response_java"):
        parts.append(f"  responsejava={{{tex_escape(spec['response_java'])}}},")
    if file_info.get("kind"):
        parts.append(f"  filekind={{{tex_escape(file_info['kind'])}}},")
    if file_info.get("name_pattern"):
        parts.append(f"  filename={{{tex_escape(file_info['name_pattern'])}}},")

    for key, tex_key, rows_key in [
        ("path_params", "pathparams", "pathrows"),
        ("query", "query", "queryrows"),
        ("body", "body", "bodyrows"),
        ("form", "form", "formrows"),
    ]:
        fields = spec.get(key) or []
        if fields:
            if key == "path_params":
                normalized = [
                    {
                        "name": item[0],
                        "java_type": item[1],
                        "label": item[2],
                        "required": True,
                    }
                    for item in fields
                ]
            else:
                normalized = fields
            parts.append(f"  {tex_key}={{")
            parts.append(tex_field_rows(normalized))
            parts.append("  },")
            parts.append(f"  {rows_key}={{1}},")

    if spec.get("data"):
        parts.append("  data={")
        parts.append(tex_field_rows(spec["data"]))
        parts.append("  },")
    parts.append("  responserows={1},")
    if spec.get("rule"):
        parts.append(f"  rule={{{tex_escape(spec['rule'])}}},")
    parts.append("}")
    return "\n".join(parts)


def tex_db_card(db: dict[str, Any], table_no: str) -> str:
    fields = [
        {
            "name": col["name"],
            "java_type": "",
            "label": col["type"],
            "desc": ("必填。" if col["required"] else "") + col["desc"],
            "required": False,
        }
        for col in db["columns"]
    ]
    return "\n".join(
        [
            rf"\EaDbCard{{{db['id']}}}{{%",
            f"  title={{{tex_escape(db['title'])}}},",
            f"  tablename={{{tex_name(db['name'])}}},",
            f"  table={{{table_no}}},",
            f"  rule={{{tex_escape(db['title'])}}},",
            "  data={",
            tex_field_rows(fields),
            "  },",
            "  datarows={1},",
            "}",
        ]
    )


def tex_elf_card(record: dict[str, Any]) -> str:
    return "\n".join(
        [
            rf"\EaElfCard{{{record['id']}}}{{%",
            f"  title={{{tex_escape(record['title'])}}},",
            f"  table={{{record['table']}}},",
            f"  rule={{{tex_escape(record['rule'])}}},",
            "  data={",
            tex_field_rows(record["data"]),
            "  },",
            "  datarows={1},",
            "}",
        ]
    )


def trim_white(image: Image.Image) -> Image.Image:
    rgb = image.convert("RGB")
    bg = Image.new("RGB", rgb.size, "white")
    diff = ImageChops.difference(rgb, bg)
    bbox = diff.getbbox()
    if not bbox:
        return rgb
    left, top, right, bottom = bbox
    return rgb.crop(
        (
            max(0, left - 8),
            max(0, top - 8),
            min(rgb.width, right + 8),
            min(rgb.height, bottom + 8),
        )
    )


def split_card(
    image: Image.Image, max_height: int = 1600
) -> list[Image.Image]:
    if image.height <= max_height:
        return [image]

    gray = image.convert("L")
    pixels = gray.load()
    width, height = gray.size
    line_rows: list[int] = []
    threshold = max(20, int(width * 0.55))
    for y in range(height):
        dark = 0
        for x in range(width):
            if pixels[x, y] < 100:
                dark += 1
                if dark >= threshold:
                    line_rows.append(y)
                    break

    segments = []
    start = 0
    while start < height:
        hard_end = min(height, start + max_height)
        if hard_end >= height:
            end = height
        else:
            candidates = [
                y
                for y in line_rows
                if start + max_height // 2 <= y <= hard_end
            ]
            end = candidates[-1] + 1 if candidates else hard_end
        segments.append(image.crop((0, start, width, end)))
        start = end
    return segments


def render_tex_card(
    card_name: str,
    card_tex: str,
    tex_dir: Path,
    prefix: str,
) -> list[str]:
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "-", card_name)
    work_dir = WORK / "card-render" / safe_name
    work_dir.mkdir(parents=True, exist_ok=True)
    tex_path = work_dir / "card.tex"
    tex_path.write_text(
        "\n".join(
            [
                r"\documentclass[UTF8,10pt,fontset=fandol]{ctexart}",
                r"\usepackage[paperwidth=210mm,paperheight=1000mm,left=10mm,right=10mm,top=10mm,bottom=10mm]{geometry}",
                r"\usepackage{api-card}",
                r"\pagestyle{empty}",
                r"\setlength{\parindent}{0pt}",
                r"\begin{document}",
                card_tex,
                r"\end{document}",
            ]
        ),
        encoding="utf-8",
    )
    env = dict(**__import__("os").environ)
    env["TEXINPUTS"] = f"{tex_dir}:{env.get('TEXINPUTS', '')}"
    pdf_path = work_dir / "card.pdf"
    log_path = work_dir / "card.log"
    pdf_path.unlink(missing_ok=True)
    log_path.unlink(missing_ok=True)
    compiled = subprocess.run(
        [
            "xelatex",
            "-interaction=nonstopmode",
            "-halt-on-error",
            "card.tex",
        ],
        cwd=work_dir,
        env=env,
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    log_text = log_path.read_text(encoding="utf-8", errors="replace") if log_path.exists() else compiled.stdout
    if compiled.returncode != 0 and (
        not pdf_path.exists() or "Output written on card.pdf" not in log_text
    ):
        raise RuntimeError(
            f"xelatex failed for {card_name}:\n{log_text[-3000:]}"
        )
    subprocess.run(
        [
            "pdftoppm",
            "-r",
            "144",
            "-png",
            str(work_dir / "card.pdf"),
            str(work_dir / "page"),
        ],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    pages = sorted(work_dir.glob("page-*.png"))
    if not pages:
        raise RuntimeError(f"No rendered pages for {card_name}")
    trimmed = [trim_white(Image.open(page)) for page in pages]
    stitched = Image.new(
        "RGB",
        (max(img.width for img in trimmed), sum(img.height for img in trimmed)),
        "white",
    )
    y = 0
    for img in trimmed:
        stitched.paste(img, (0, y))
        y += img.height
    segments = split_card(stitched)
    rels = []
    for idx, segment in enumerate(segments, 1):
        file_name = f"{prefix}-{safe_name}-{idx:02d}.png"
        segment.save(OUT / "接口表/images" / file_name, optimize=True)
        rels.append(f"images/{file_name}")
    return rels


def render_images(specs: list[dict[str, Any]], dbs: list[dict[str, Any]], masters: list[dict[str, Any]]) -> None:
    tex_dir = OUT / "接口表/tex"
    card_dir = tex_dir / "cards"
    card_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(
        Path(
            "/home/nc/projects/work/cuc-ah-ucac-pso-hx-repository/"
            "doc/hx-assessment/api-tables/tex/api-card.sty"
        ),
        tex_dir / "api-card.sty",
    )
    for spec in specs:
        card_tex = tex_http_card(spec)
        (card_dir / f"{spec['id']}.tex").write_text(card_tex, encoding="utf-8")
        spec["image_rels"] = render_tex_card(
            spec["id"], card_tex, tex_dir, "api"
        )
        spec["image_rel"] = spec["image_rels"][0]
    for idx, db in enumerate(dbs, 1):
        card_tex = tex_db_card(db, f"7.1-{idx}")
        (card_dir / f"{db['name']}.tex").write_text(card_tex, encoding="utf-8")
        db["image_rels"] = render_tex_card(
            db["name"], card_tex, tex_dir, "data"
        )
        db["image_rel"] = db["image_rels"][0]
    for master in masters:
        card_tex = tex_elf_card(master)
        (card_dir / f"{master['id']}.tex").write_text(card_tex, encoding="utf-8")
        master["image_rels"] = render_tex_card(
            master["id"], card_tex, tex_dir, "api"
        )
        master["image_rel"] = master["image_rels"][0]


def render_cover_backgrounds() -> None:
    bg_dir = WORK / "backgrounds"
    bg_dir.mkdir(parents=True, exist_ok=True)
    cover_html = bg_dir / "cover.html"
    back_html = bg_dir / "backcover.html"
    cover_svg = f"""<!doctype html><html><head><meta charset="utf-8">
    <style>html,body{{margin:0;width:1588px;height:2246px;overflow:hidden}}
    body{{background:{COLORS['panel']};font-family:'Noto Sans CJK SC'}}</style></head>
    <body><svg width="1588" height="2246" xmlns="http://www.w3.org/2000/svg">
    <rect width="1588" height="2246" fill="{COLORS['primary']}"/>
    <path d="M0 250 C390 80 860 500 1588 210 L1588 0 L0 0Z" fill="#405B54"/>
    <path d="M0 1780 C360 1600 950 1940 1588 1680 L1588 2246 L0 2246Z" fill="#2B403B"/>
    <rect x="126" y="126" width="1336" height="1994" fill="none" stroke="{COLORS['accent']}" stroke-width="2" opacity=".7"/>
    <rect x="150" y="150" width="1288" height="1946" fill="none" stroke="{COLORS['accent']}" stroke-width=".7" opacity=".45"/>
    <circle cx="1320" cy="390" r="175" fill="none" stroke="{COLORS['accent']}" stroke-width="1.2" opacity=".25"/>
    <circle cx="1320" cy="390" r="135" fill="none" stroke="{COLORS['accent']}" stroke-width=".7" opacity=".25"/>
    <line x1="350" y1="810" x2="1238" y2="810" stroke="{COLORS['accent']}" stroke-width="1" opacity=".5"/>
    </svg></body></html>"""
    back_svg = f"""<!doctype html><html><head><meta charset="utf-8">
    <style>html,body{{margin:0;width:1588px;height:2246px;overflow:hidden}}
    body{{background:{COLORS['panel']};font-family:'Noto Sans CJK SC'}}</style></head>
    <body><svg width="1588" height="2246" xmlns="http://www.w3.org/2000/svg">
    <rect width="1588" height="2246" fill="{COLORS['panel']}"/>
    <rect x="126" y="126" width="1336" height="1994" fill="none" stroke="{COLORS['primary']}" stroke-width="2" opacity=".55"/>
    <rect x="150" y="150" width="1288" height="1946" fill="none" stroke="{COLORS['accent']}" stroke-width=".8" opacity=".5"/>
    <path d="M0 0 H1588 V330 C1080 470 470 180 0 380Z" fill="{COLORS['primary']}" opacity=".10"/>
    <path d="M0 2246 H1588 V1960 C1120 1830 430 2080 0 1910Z" fill="{COLORS['primary']}" opacity=".10"/>
    <line x1="650" y1="1080" x2="938" y2="1080" stroke="{COLORS['accent']}" stroke-width="1"/>
    <circle cx="794" cy="1080" r="5" fill="{COLORS['accent']}"/>
    </svg></body></html>"""
    cover_html.write_text(cover_svg, encoding="utf-8")
    back_html.write_text(back_svg, encoding="utf-8")
    for src, name in [(cover_html, "cover_bg.png"), (back_html, "backcover_bg.png")]:
        subprocess.run(
            [
                "npx",
                "--yes",
                "playwright",
                "screenshot",
                "--channel",
                "chrome",
                "--viewport-size",
                "1588,2246",
                "--full-page",
                src.as_uri(),
                str(bg_dir / name),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
        )


def create_diagram() -> Path:
    width, height = 1500, 900
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(FONT_BOLD, 28)
    body_font = ImageFont.truetype(FONT_REGULAR, 23)
    small_font = ImageFont.truetype(FONT_REGULAR, 18)

    bar_x, bar_y, bar_w, bar_h = 60, 45, 1380, 58
    d.rectangle((bar_x, bar_y, bar_x + bar_w, bar_y + bar_h), fill="#3D6B82")
    title = "员工考核系统改造总体功能结构"
    tw = d.textbbox((0, 0), title, font=title_font)[2]
    d.text((width // 2 - tw // 2, bar_y + 12), title, font=title_font, fill="white")

    center = (555, 160, 945, 220)
    d.rectangle(center, fill="#D9E8EF", outline="#789CB1", width=2)
    center_text = "员工年度考核应用"
    cw = d.textbbox((0, 0), center_text, font=body_font)[2]
    d.text((750 - cw // 2, 177), center_text, font=body_font, fill="#263746")

    boxes = [
        (105, 375, 300, 300, "管理端", ["考核模板", "年度考核计划", "人员、主考核与分管", "发布凭证", "评分结构与汇总等次"], "#DCEAF0"),
        (440, 375, 300, 300, "匿名评分端", ["账号登录", "评分会话", "评分工作台", "按比例分配", "保存与最终提交"], "#D8EAE4"),
        (775, 375, 300, 300, "员工实名端", ["年度切换", "登记表查询", "草稿保存", "首次提交", "Word导出"], "#F2E8D2"),
        (1110, 375, 300, 300, "边界外主系统", ["人员", "部门", "实名身份", "管理员角色", "只读基础信息"], "#EADCE4"),
    ]
    for x, y, w, h, title, lines, fill in boxes:
        d.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=fill, outline="#7C98A8", width=2)
        label = title
        lw = d.textbbox((0, 0), label, font=body_font)[2]
        d.text((x + w // 2 - lw // 2, y + 24), label, font=body_font, fill="#263746")
        yy = y + 84
        for line in lines:
            lw = d.textbbox((0, 0), line, font=small_font)[2]
            d.text((x + w // 2 - lw // 2, yy), line, font=small_font, fill="#37474F")
            yy += 36

    def arrow(x1: int, y1: int, x2: int, y2: int) -> None:
        d.line((x1, y1, x2, y2), fill="#3D6B82", width=3)
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 14
        for offset in (2.55, -2.55):
            xa = x2 + size * math.cos(angle + offset)
            ya = y2 + size * math.sin(angle + offset)
            d.line((x2, y2, xa, ya), fill="#3D6B82", width=3)

    arrow(640, 220, 255, 375)
    arrow(710, 220, 590, 375)
    arrow(790, 220, 925, 375)
    arrow(860, 220, 1260, 375)
    d.text((146, 720), "本系统维护考核业务数据；主系统只读提供人员、部门和实名身份。", font=small_font, fill="#455A64")
    path = WORK / "system-structure.png"
    img.save(path)
    return path


def section_texts() -> list[dict[str, Any]]:
    return [
        {
            "heading": "1 项目概述",
            "level": 1,
            "paras": [
                "本项目基于员工年度考核 v3-new-v2 阶段一计数模型，对可提交的业务需求、数据模型、页面功能、接口契约和验收约束进行完整说明。需求正文以当前后端固定提交和前端固定提交的实现事实为依据，功能规模只采用阶段一 JSON 与证据 XLSX 的已确认结果。",
                "系统面向考核管理员、实名员工和匿名评分人提供三个相互隔离的用户边界。管理端负责模板、计划、人员快照、主考核与分管、评分结构和结果处理；实名员工端只允许本人访问年度考核登记表；匿名评分端使用独立凭证和会话完成席位评分。",
            ],
        },
        {
            "heading": "1.1 系统总体功能结构",
            "level": 2,
            "paras": [
                "系统由管理端、实名员工端、匿名评分端和边界外主系统共同构成。本应用维护年度考核业务数据，主系统维护人员、部门、实名身份和管理员角色。",
            ],
        },
        {
            "heading": "1.2 项目预算工作量",
            "level": 2,
            "paras": [
                "按 NESMA 详细功能点分析，本版共形成 13 个数据功能和 37 个事务功能。其中 11 个 ILF、2 个 ELF；22 个 EI、10 个 EO、5 个 EQ。数据功能合计 87 FP，事务功能合计 180 FP，总计 267 UFP。",
                "上述数量、类型、复杂度与 FP 均直接取自 v3-new-v2 阶段一模型，不在本阶段重新计数。发布就绪核对、等次填写表导出、退出评分、人员导入模板下载，以及数据割接、迁移、环境切换、线下凭证发放和线下等次评定不增加功能点。",
            ],
        },
        {
            "heading": "2 文档目的、范围与依据",
            "level": 1,
            "paras": [
                "本文用于说明系统已交付范围内的业务能力、数据对象、页面入口、接口输入输出以及可验收规则，为需求确认、测试设计和功能规模复核提供统一依据。",
                "文档覆盖模板与计划全生命周期、人员文件预检与确认、总监主考核和领导分管范围、评分结构、匿名账号与凭证、评分任务、评分汇总与等次导入、年度考核登记表。通用登录、菜单、布局、分页、日志、数据库约束和缓存等技术机制不重复描述为业务功能。",
                "阶段一 JSON 是唯一计数事实来源。旧版 spec-book、旧功能清单和旧接口 YAML 仅用于理解版式和章节组织，其旧总计数、旧编号、旧类型和旧数据功能分类不得继承。",
            ],
        },
        {
            "heading": "3 术语和缩写",
            "level": 1,
            "paras": [
                "年度考核计划：某一年度发起的一次员工年度考核安排，是计划任职、评分结构、评分任务与结果的共同业务边界。",
                "计划人员快照：计划内由用户导入并固化的员工编号和姓名记录，不是跨计划复用的本地人员主数据。",
                "计划部门快照：计划内由导入文件固化的考核部门名称，是任职、评分席位与范围的计划内身份。",
                "总监主考核 BU：总监从本人总监任职中选择的唯一受考核归属部门，也是其受评分和员工组自评的归属。",
                "领导普通员工分管范围：领导以总监、正职或副职身份评价普通员工成员的最终集合，显式配置可替换默认范围。",
                "评分席位：计划部门内按评分职级形成的评分资格；内部标识没有业务含义，不在结果中展示。",
                "匿名评分账号：计划内临时认证凭证，拥有一个或多个席位授权，状态为待激活、可用或已最终提交。",
                "评分会话：匿名账号通过密码兑换的数据库权威临时会话，原始令牌只通过专用 HttpOnly Cookie 传递。",
                "最终提交：匿名账号将全部授权席位的完整评分草稿一次性转为正式评分的不可逆操作。",
                "评分汇总分：系统按评分组计算并固化的分数，不直接换算年度考核等次。",
                "年度考核等次：客户线下综合评定后导入系统的优秀、良好、称职、待改进和不称职。",
                "年度考核登记表：按员工和年度唯一的独立年度档案，承载基础信息快照、年度主要工作摘要和考核等次。",
                "NESMA：系统与软件工程 功能规模测量 NESMA 方法 GB/T 42588—2023。UFP 为未调整功能点。",
            ],
        },
        {
            "heading": "4 范围定义与角色",
            "level": 1,
            "paras": [
                "功能用户包括考核管理员、实名员工和匿名评分人。管理员使用 admin 或 ea-admin 角色访问计划和模板；全部实名登录员工均可访问本人登记表；匿名评分人使用独立账号，不进入实名菜单。",
                "管理端和实名端使用主系统实名身份；主系统同时维护人员、部门和管理员角色。本应用只读读取必要的身份与组织事实，并在计划或登记表业务发生时固化快照，不回流主系统。",
                "应用边界内的持久业务数据包括模板与指标、计划与计划快照、评分结构与账号授权、评分任务与记录、评分汇总和年度考核登记表。评分会话、操作日志、年度守卫和实时配额属于技术控制数据，不单列为数据功能。",
            ],
        },
        {
            "heading": "5 总体业务规则与状态",
            "level": 1,
            "paras": [
                "模板必须为“分组—叶指标”两层结构，全部叶指标满分合计为100。草稿可修改，发布后只读；停用只阻止新计划选择，不影响已引用计划。",
                "同一年度最多存在一个未终止计划。计划从草稿发布为进行中，之后只能完成或终止。发布时冻结评分结构、生成账号授权与评分任务；发布后不允许回退或修改结构。",
                "人员导入采用“先预检、后确认”两步一致流程。确认时必须提交同一文件、预检摘要和客户端读取的草稿版本；任一不匹配即零写入。",
                "评分账号与席位授权分离。员工匿名账号不保存实际使用人；领导匿名账号可在内部关联计划人员。一个席位只授权一个账号，融合领导账号可拥有多个领导席位。",
                "评分草稿按任务版本保存；最终提交以账号为边界，对全部授权席位一次性校验、提交、停用账号并撤销会话。",
                "计划完成时先计算组内平均分与组加权分，再汇总个人评分汇总分并固化。评分汇总分不直接映射年度考核等次；等次由管理员文件全量导入。",
                "登记表按“员工＋年度”唯一。员工首次保存时固化主系统基础资料，提交后工作摘要只读；管理员仍可通过等次导入更正等次字段。",
            ],
        },
    ]


def data_model_content() -> list[dict[str, Any]]:
    return [
        {
            "heading": "7 数据模型",
            "level": 1,
            "paras": ["本章按用户可识别的数据功能组织。每个模型说明业务含义、字段分组、来源流向与关键加工规则；物理表卡片见各节末尾的图片和接口表 YAML。"],
        },
        {
            "heading": "7.1.1 考核模板与计划指标",
            "level": 2,
            "nesma": ["DF-04", "DF-05"],
            "paras": [
                "考核模板承载模板名称、说明、状态、可用于新计划标记和两层指标结构。叶指标记录名称、说明、同级排序和满分；分组满分由叶指标求和展示，不独立保存。计划创建后保存所选模板的指标快照，原模板后续变化不覆盖已建计划。",
                "年度考核计划记录考核年度、模板、生命周期状态、备注、发起人、草稿版本及发布、完成、终止时间，并保存四项默认评分组权重。计划指标快照与模板指标解耦，保证同一计划执行期间口径稳定。",
            ],
        },
        {
            "heading": "7.1.2 计划人员、部门、任职、权重与分管范围",
            "level": 2,
            "nesma": ["DF-01", "DF-02", "DF-08"],
            "paras": [
                "组织人员导入快照由计划人员、计划部门和计划任职三个记录组构成。人员以员工编号和主系统人员标识识别，部门以导入名称识别，任职记录计划人员在考核部门中的职级。",
                "部门职级权重记录导入文件给出的按部门和职级覆盖权重；未填写项继承计划默认权重。领导普通员工分管范围记录评分主体、部门、任职角色、员工成员和范围来源。默认范围由任职推导，显式配置整体替换某领导在该部门的默认普通员工范围。",
            ],
        },
        {
            "heading": "7.1.3 总监主考核",
            "level": 2,
            "nesma": ["DF-03"],
            "paras": [
                "总监主考核记录总监唯一主考核 BU 以及逐总监指定的正副职评分名单。名单人员必须以其已经在文件中存在的正职或副职任职确认角色；同一 BU 的不同总监可以使用不同名单。",
                "总监跨多个 BU 任职时只形成一个被考核人，其主考核结果归属于主考核 BU；分管范围仅在领导以分管身份评价普通员工时生效，不改变总监的唯一被考核归属。",
            ],
        },
        {
            "heading": "7.1.4 评分结构",
            "level": 2,
            "nesma": ["DF-06"],
            "paras": [
                "评分结构包含被考核人、评分组、评分关系和评分席位。系统依据计划人员、部门、职级、主考核名单和分管范围编译结构：普通员工归属其员工任职部门，总监归属主考核BU。",
                "评分组包含条件性存在的总监组、正职组、副职组和员工组。基础权重由计划默认值与部门职级覆盖共同确定；副职缺席转正职，总监缺席转员工，发布时有效权重合计100%。",
                "评分席位表达应当评价谁，不承担登录认证。内部席位标识没有业务含义，不在结果展示、导出或审计明细中披露。",
            ],
        },
        {
            "heading": "7.1.5 匿名评分账号与席位授权",
            "level": 2,
            "nesma": ["DF-07"],
            "paras": [
                "匿名评分账号记录匿名账号、初始密码摘要、账号类型、状态、最终提交时间及领导发放对象。账号状态为待激活、可用或已最终提交。",
                "账号席位授权将每个评分席位恰好连接到一个匿名账号。员工账号不保存实际使用人；领导账号关联计划人员并可授权多个领导席位。凭证只在发布或全量重置成功响应中一次性导出。",
            ],
        },
        {
            "heading": "7.1.6 评分任务与指标所得分",
            "level": 2,
            "nesma": ["DF-09"],
            "paras": [
                "评分任务在发布时依据评分关系生成，记录被考核人、任务状态、草稿版本和评分总分。指标所得分记录每项叶子指标的当前分值，全部指标填写完成时才形成任务总分。",
                "草稿保存不接受客户端总分作为权威，服务端依据指标分值重新求和。最终提交以账号为边界，统一提交全部授权席位并撤销会话。",
            ],
        },
        {
            "heading": "7.1.7 评分汇总",
            "level": 2,
            "nesma": ["DF-10"],
            "paras": [
                "评分汇总包含评分组汇总和被考核人汇总分。计划完成时，系统按评分组计算六位组内平均分和组加权分，再汇总为两位评分汇总分，并记录计算时间。",
                "汇总只从正式提交任务生成，完成后只读。审计明细按“评分1、评分2”等组内匿名序号展示，不披露账号、席位或评分人员身份。",
            ],
        },
        {
            "heading": "7.1.8 年度考核登记表",
            "level": 2,
            "nesma": ["DF-11"],
            "paras": [
                "年度考核登记表按员工编号和考核年度唯一，保存姓名、现单位、现部门、年度主要工作摘要、年度考核等次、登记状态、提交时间及草稿版本。",
                "员工首次保存时从当前实名身份和主系统部门固化快照；等次导入代建时从计划人员与受考核任职固化。提交后员工不能修改摘要，管理员可通过全量等次导入更正等次。",
            ],
        },
        {
            "heading": "7.1.9 主系统人员与部门",
            "level": 2,
            "nesma": ["DF-12", "DF-13"],
            "paras": [
                "主系统人员与部门是边界外只读数据。本应用在人员导入时校验员工标识、人员状态和部门，在登记表操作时读取当前员工编号、姓名、现单位和现部门。",
                "本应用不对这两组数据执行新增、修改或删除，只在业务发生时固化计划快照或登记表快照，后续主系统变化不反向覆盖已生效业务事实。",
            ],
        },
    ]


def page_content() -> list[dict[str, Any]]:
    return [
        {
            "heading": "8 界面功能模块",
            "level": 1,
            "paras": ["本章按用户入口组织功能。功能清单中的每个可计数事务过程均在此有明确归属；接口图片在第九章给出。"],
        },
        {
            "heading": "8.1 角色与入口隔离",
            "level": 2,
            "paras": ["管理端计划和模板入口只对 admin 与 ea-admin 开放；登记表入口对全部实名登录员工开放并限定本人；匿名评分入口 `/rating` 独立于实名菜单和实名会话。"],
        },
        {
            "heading": "8.2 考核模板页面",
            "level": 2,
            "pages": [
                ("8.2.1 查询考核模板", "管理员按名称和状态查看模板列表，并打开详情查看两层指标树。", ["TF-08"]),
                ("8.2.2 新建考核模板", "管理员录入模板名称和说明；系统创建草稿并记录审计信息。", ["TF-01"]),
                ("8.2.3 修改模板资料", "仅草稿模板可修改基础资料；发布后基本信息只读。", ["TF-02"]),
                ("8.2.4 查看模板详情", "展示模板基础信息和分组、叶指标结构。", ["TF-08"]),
                ("8.2.5 配置指标", "管理员整体保存两层指标结构，系统校验叶指标满分合计为100。", ["TF-03"]),
                ("8.2.6 发布模板", "校验结构后发布，发布后不可继续修改。", ["TF-04"]),
                ("8.2.7 停用模板", "停用后不再出现在新计划可选模板中，已引用计划不受影响。", ["TF-05"]),
                ("8.2.8 复制模板", "从已有模板复制独立草稿并允许输入新名称。", ["TF-06"]),
                ("8.2.9 删除草稿模板", "仅未发布模板可删除。", ["TF-07"]),
            ],
        },
        {
            "heading": "8.3 考核计划列表",
            "level": 2,
            "pages": [
                ("8.3.1 查询计划", "按年度和状态查看计划、模板、版本及生命周期时间。", ["TF-12"]),
                ("8.3.2 新建计划", "选择年度和已发布模板创建草稿；同年度未终止计划唯一。", ["TF-09"]),
                ("8.3.3 删除草稿计划", "删除从未发布的计划及草稿结构。", ["TF-11"]),
            ],
        },
        {
            "heading": "8.4 计划工作台—概览与生命周期",
            "level": 2,
            "pages": [
                ("8.4.1 修改草稿计划", "按草稿版本修改模板和备注，并重建未发布结构。", ["TF-10"]),
                ("8.4.2 发布计划", "执行完整就绪门禁，冻结结构、生成任务和账号并交付凭证。", ["TF-19"]),
                ("8.4.3 重置凭证", "批量重置全部未提交账号并重新交付。", ["TF-20"]),
                ("8.4.4 查询评分进度", "展示任务与账号进度并说明完成条件。", ["TF-23"]),
                ("8.4.5 终止计划", "不可逆停止后续评分并撤销全部会话。", ["TF-21"]),
                ("8.4.6 完成计划", "全部任务提交后固化评分汇总并结束计划。", ["TF-22"]),
            ],
        },
        {
            "heading": "8.5 计划工作台—人员操作",
            "level": 2,
            "pages": [
                ("8.5.1 下载导入模板", "下载两张固定表头的空 Excel 模板；该支撑动作不单列功能点。", []),
                ("8.5.2 核对人员文件", "只读校验人员、部门、职级、权重和主考核，返回统计、替换影响与结构预览。", ["TF-13"]),
                ("8.5.3 确认替换人员配置", "以同一文件、摘要和版本确认，原子替换计划人员快照并重建结构。", ["TF-14"]),
            ],
        },
        {
            "heading": "8.6 计划工作台—任职清单",
            "level": 2,
            "pages": [
                ("8.6.1 查询任职清单", "按部门和职级筛选，分页查看任职快照和受考核任职标记。", ["TF-15"]),
            ],
        },
        {
            "heading": "8.7 计划工作台—主考核与分管",
            "level": 2,
            "pages": [
                ("8.7.1 查询主考核与分管", "展示总监主考核卡片和当前页领导最终员工成员。", ["TF-16"]),
                ("8.7.2 保存分管范围", "显式选择员工成员并重建普通员工分管范围。", ["TF-17"]),
                ("8.7.3 导出核对表", "导出完整人员结构与领导分管两张核对表。", ["TF-27"]),
            ],
        },
        {
            "heading": "8.8 计划工作台—结构预览",
            "level": 2,
            "pages": [
                ("8.8.1 核对发布就绪状态", "全计划派生就绪查询，作为发布门禁支撑，不单列功能点。", []),
                ("8.8.2 查询评分结构", "按部门分页查看评分来源、有效权重、缺组转移和预期任务。", ["TF-18"]),
            ],
        },
        {
            "heading": "8.9 计划工作台—评分汇总与等次",
            "level": 2,
            "pages": [
                ("8.9.1 查询评分汇总", "按部门、关键词、等次和排序查看已固化结果。", ["TF-24"]),
                ("8.9.2 导出等次填写表", "导出固定列供线下填写；该导出归入汇总与导入流程，不单列。", []),
                ("8.9.3 导入年度考核等次", "全量校验后原子更新登记表等次。", ["TF-26"]),
            ],
        },
        {
            "heading": "8.10 评分汇总审计明细",
            "level": 2,
            "pages": [
                ("8.10.1 查询评分汇总明细", "按匿名序号展示指标、任务、组均分、组加权和组成项。", ["TF-25"]),
            ],
        },
        {
            "heading": "8.11 年度考核登记表",
            "level": 2,
            "pages": [
                ("8.11.1 查询可选年度", "返回本人可切换年度及默认年度。", ["TF-28"]),
                ("8.11.2 查询登记表", "返回本人档案或临时空表。", ["TF-29"]),
                ("8.11.3 保存草稿", "首次保存固化身份快照，后续按版本更新摘要。", ["TF-30"]),
                ("8.11.4 提交登记表", "校验摘要后首次提交并使员工内容只读。", ["TF-31"]),
                ("8.11.5 导出 Word", "按客户固定模板导出已提交登记表。", ["TF-32"]),
            ],
        },
        {
            "heading": "8.12 匿名评分工作台",
            "level": 2,
            "pages": [
                ("8.12.1 匿名登录与恢复", "使用账号密码兑换会话，或从有效 Cookie 恢复任务清单。", ["TF-33", "TF-34"]),
                ("8.12.2 查询席位草稿", "一次读取席位全部任务草稿、公共指标和配额边界。", ["TF-34"]),
                ("8.12.3 查询单任务草稿", "用于单行恢复或重新失败重试。", ["TF-34"]),
                ("8.12.4 按比例分配", "按最大余数法生成指标建议分，不写库。", ["TF-35"]),
                ("8.12.5 保存手填草稿", "保存完整指标分值并重算任务总分。", ["TF-36"]),
                ("8.12.6 退出评分", "撤销技术会话并清除 Cookie，按技术动作排除计数。", []),
                ("8.12.7 最终提交", "员工或领导账号统一提交全部授权席位。", ["TF-37"]),
            ],
        },
        {
            "heading": "8.13 主系统数据接口",
            "level": 2,
            "pages": [
                ("8.13.1 导入校验读取", "人员导入前只读校验主系统人员和部门。", []),
                ("8.13.2 当前实名人员读取", "登记表操作读取本人姓名、现单位和现部门。", []),
            ],
        },
    ]


def interface_intro() -> dict[str, Any]:
    return {
        "heading": "9 接口契约",
        "level": 1,
        "paras": [
            "本章接口图片由 `接口表/http/*.yaml` 和 `接口表/data/*.yaml` 生成。HTTP 成功响应遵循若依既有信封：AjaxResult 使用 code、msg、data；TableDataInfo 使用 code、msg、total、rows。文件下载接口不虚构 JSON；成功时直接输出对应文件。",
            "接口表只描述成功载荷，不收录失败码块。匿名评分令牌只从专用 Cookie 读取，不进入路径、查询参数或响应业务字段。接口中的 NESMA 编号直接引用阶段一计数模型。",
        ],
    }


def legacy_front_sections(model: dict[str, Any]) -> list[dict[str, Any]]:
    terms = [
        ("年度考核计划", "某一年度一次员工年度考核，是计划任职、评分关系、评分任务和评分结果的共同载体。"),
        ("被考核对象", "计划内实际接受年度考核评分的人员，每名计划人员在一个计划内最多一个。"),
        ("评分席位", "计划部门内按职级形成的一份评分资格；内部序号没有业务含义。"),
        ("匿名评分账号", "计划内临时认证凭证，通过账号席位授权访问一个或多个评分席位。"),
        ("强匿名员工账号", "不保存实际使用人归属的员工组匿名账号，可评价本部门全部被考核员工。"),
        ("评分会话", "匿名账号登录后取得的数据库权威访问会话。"),
        ("最终提交", "匿名账号将全部授权席位的完整评分草稿一次性转为正式评分的不可逆操作。"),
        ("评分汇总分", "系统完成计划时按评分组计算并固化的分数，不直接换算年度考核等次。"),
        ("年度考核等次", "由管理员线下综合评定后导入的优秀、良好、称职、待改进或不称职。"),
        ("年度考核登记表", "按员工和年度唯一、独立于计划的年度档案。"),
        ("NESMA", "系统与软件工程 功能规模测量 NESMA 方法 GB/T 42588—2023。"),
        ("UFP", "未调整功能点。本版计数结果为 267 UFP。"),
    ]
    return [
        {"heading": "1 项目概述", "level": 1, "paras": []},
        {
            "heading": "1.1 项目背景",
            "level": 2,
            "paras": [
                "M域信创改造，改造对象是安徽众炎软件开发维护的中国联通安徽省分公司人力资源项目（HR支撑平台）——员工年度考核系统。改造范围是全系统使用现代技术栈重新设计后台数据服务和前台用户界面，同时干净地割接生产数据，最终通过大额成本管控主项目进行生产部署。",
            ],
        },
        {
            "heading": "1.2 系统总体功能结构",
            "level": 2,
            "paras": [
                "系统由管理端、匿名评分端和员工实名端组成，并只读使用边界外主系统维护的人员、部门实名身份和管理员角色。",
            ],
            "images": [
                {
                    "path": str((WORK / "system-structure.png").resolve()),
                    "caption": "图 1：系统总体功能结构",
                }
            ],
        },
        {"heading": "1.3 往期建设情况", "level": 2, "paras": ["无"]},
        {
            "heading": "1.4 本项目预算工作量",
            "level": 2,
            "paras": [
                "工作量计数对象为员工年度考核从无到有完整重新实现的功能范围，采用 NESMA 详细功能点分析，项目类型为新开发项目。计数本系统边界为验收时代码库基线包含的管理端、匿名评分端、员工端及模板、计划、导入、评分、汇总、等次和年度考核登记表业务功能；排除通用平台、其他业务域、技术实现和非功能需求。已核验结果包括13个数据功能和37个事务功能，数据功能87 FP、事务功能180 FP，合计267 UFP。",
            ],
        },
        {"heading": "1.5 本项目投资计划预算", "level": 2, "paras": ["无"]},
        {
            "heading": "2 文档目的",
            "level": 1,
            "paras": [
                "本文基于开发人员的已核验 NESMA 详细 FPA 结论，用于说明员工年度考核系统改造的已交付业务范围、数据模型、页面功能、接口契约和可验收约束，从而为需求支撑提供依据。",
            ],
        },
        {
            "heading": "3 文档范围",
            "level": 1,
            "paras": [
                "本文覆盖管理员使用的考核模板与计划全生命周期、计划人员任职导入、匿名凭证交付、匿名评分、评分汇总、年度考核等次处理，以及员工本人年度考核登记表。本文不覆盖通用平台登录、通用布局和分页能力，干部考核及其他业务域，当前生产系统，也不把数据库约束、触发器、事务、缓存限流等技术机制重复计为业务功能。",
            ],
        },
        {
            "heading": "4 术语和缩写解释",
            "level": 1,
            "paras": [],
            "tables": [
                {
                    "headers": ["术语或缩写", "定义"],
                    "widths": [2600, 5700],
                    "rows": terms,
                }
            ],
        },
        {
            "heading": "5 参考文档",
            "level": 1,
            "paras": [
                "《系统与软件工程功能规模测量 NESMA 方法 GB/T42588—2023》",
                "《2025年中国软件行业基准数据（CSBMK® -202510）》",
                "GB∕T36964-2018 软件工程 软件开发成本度量规范",
            ],
        },
        {
            "heading": "6 范围定义",
            "level": 1,
            "paras": [
                "员工年度考核应用负责维护模板、计划结构与快照、被考核对象、匿名账号与会话、评分任务与记录、汇总结果、年度考核登记表和业务操作日志。大额成本管控主系统负责维护人员、部门及实名登录身份；本系统只在计划导入和员工登记表操作时读取必要信息。",
                "系统用户包括年度考核管理员、匿名评分用户（领导匿名账号和强匿名员工账号）、员工实名用户，以及向本系统提供人员、部门和实名身份信息的主系统。",
            ],
        },
    ]


def model_paragraphs(items: list[tuple[str, str, list[str]]]) -> list[dict[str, Any]]:
    sections = []
    for heading, base, paragraphs in items:
        sections.extend(
            [
                {"heading": f"{base} {heading}", "level": 4, "paras": []},
                {
                    "heading": f"{base}.1 需求描述",
                    "level": 5,
                    "paras": [paragraphs[0]],
                },
                {
                    "heading": f"{base}.2 数据来源及流向",
                    "level": 5,
                    "paras": [paragraphs[1]],
                },
                {
                    "heading": f"{base}.3 加工逻辑",
                    "level": 5,
                    "paras": [paragraphs[2]],
                },
            ]
        )
    return sections


def legacy_data_sections(dbs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = [
        {"heading": "7 功能需求", "level": 1, "paras": []},
        {"heading": "7.1 数据模型", "level": 2, "paras": []},
        {"heading": "7.1.1 考核配置与计划结构", "level": 3, "paras": []},
    ]
    sections.extend(
        model_paragraphs(
            [
                (
                    "考核模板",
                    "7.1.1.1",
                    [
                        "新增模型。考核模板承载年度考核采用的模板基本信息和两层指标结构（评分表格），是管理员在考核模板页面上直接创建和维护的独立业务对象。模型内包含模板基本信息（模板名称、模板状态、模板启用状态、模板说明、模板创建人、模板创建时间、模板更新人、模板更新时间）和模板指标结构（指标层级、指标名称、指标说明、指标排序、指标满分）。物理表见表7.1-1、表7.1-2。",
                        "模板基本信息由管理员在模板页面录入；模板创建人、更新人取当前实名管理员身份；模板指标结构由管理员在指标配置操作中整体维护。模型数据流向考核模板页面的列表与详情展示，并作为输入流向年度考核计划的创建、草稿更新和发布过程。",
                        "系统在模板新建、基础信息修改、指标替换、发布、停用和复制时写入或更新模型数据。指标必须为“分组—叶子”两层结构，全部叶指标满分合计为100；草稿可编辑和删除，发布后只读，已发布模板可停用但不物理删除。",
                    ],
                ),
                (
                    "年度考核计划及计划指标快照",
                    "7.1.1.2",
                    [
                        "新增模型。年度考核计划承载一次年度评分安排及其计划基本信息、人员快照、部门快照、任职快照、计划指标快照、评分组、评分席位和评分关系，是系统内最大的复合业务对象。物理表见表7.1-3至表7.1-14。",
                        "计划基本信息由管理员创建和维护；人员、部门、任职和权重来自管理员导入的完整人员任职文件；计划指标快照取自所选考核模板；评分结构由系统依据计划部门、职级、主考核和分管范围自动编译生成。",
                        "同一年度最多存在一个未终止计划。计划草稿可修改和删除；发布时冻结评分结构并激活账号与评分任务；完成时固化评分汇总；终止时停止后续评分。发布后评分结构不允许回退或修改。",
                    ],
                ),
                (
                    "组织人员导入快照",
                    "7.1.1.3",
                    [
                        "新增模型。组织人员导入快照由计划人员、计划部门和计划任职三个记录组构成。人员由员工编号、主系统人员标识、姓名识别；部门仅保存导入文件中的考核部门名称和系统生成的计划内标识；任职记录计划人员在考核部门中的职级。物理表见表7.1-5至表7.1-7。",
                        "三类数据均来自管理员上传的完整人员文件，在计划内固化为不可随主系统变化覆盖的快照。数据流向任职清单、评分结构编译、主考核与分管核对以及评分汇总。",
                        "确认替换时必须同时匹配同一文件、预检摘要和客户端草稿版本。任一行校验失败则整份文件零写入；确认成功后原子替换人员快照并重建未发布评分结构。",
                    ],
                ),
                (
                    "部门职级权重",
                    "7.1.1.4",
                    [
                        "新增模型。部门职级权重记录导入文件按考核部门和职级给出的基础权重覆盖值。未填写项继承计划默认权重。物理表见表7.1-8。",
                        "权重由管理员在人员文件中填写并随计划人员导入固化，流向评分结构编译、结构预览和评分汇总计算。",
                        "同一部门同一职级的非空权重必须一致；发布前按实际存在的评分组解析有效权重，副职缺席转正职，总监缺席转员工，最终权重合计100%。",
                    ],
                ),
                (
                    "总监主考核",
                    "7.1.1.5",
                    [
                        "新增模型。总监主考核记录总监唯一主考核BU和逐总监指定的正副职评分名单。名单人员必须以已在人员文件中存在的正职或副职任职确认角色。物理表见表7.1-9、表7.1-10。",
                        "主考核BU和评分人名单均由管理员在人员文件中配置，本系统只读展示并固化。数据流向评分结构、主考核与分管页面、核对表导出和评分汇总明细。",
                        "主考核BU必须是总监本人已有总监任职的部门。同一BU的不同总监可以使用不同名单；评分人不能是总监本人；总监跨多个BU任职时只形成一个被考核人。",
                    ],
                ),
                (
                    "领导普通员工分管范围",
                    "7.1.1.6",
                    [
                        "新增模型。领导普通员工分管范围记录评分主体、考核部门、任职角色、普通员工成员和范围来源。物理表见表7.1-11、表7.1-12。",
                        "默认范围按领导任职和普通员工任职推导；管理员可在主考核与分管页显式选择成员替换默认范围。数据流向评分结构、结构预览和评分汇总。",
                        "显式名单替换对应领导在该部门的默认员工范围，空名单表示不再评价任何普通员工。保存后重建未发布结构；重复人员导入时清除旧派生配置并按新任职重建默认范围。",
                    ],
                ),
                (
                    "被考核对象与评分结构",
                    "7.1.1.7",
                    [
                        "新增模型。评分结构包含被考核人、评分组、评分关系和评分席位。普通员工归属其员工任职部门，总监归属主考核BU。物理表见表7.1-13、表7.1-15至表7.1-17。",
                        "结构由系统依据计划人员、部门、职级、主考核名单和分管范围编译，不接受用户直接维护。数据流向结构预览、发布计划、匿名评分任务和评分汇总。",
                        "一个计划人员最多形成一个被考核人；一个评分席位恰好授权一个匿名账号；内部席位标识不展示、不导出。发布时固化结构，完成后永久只读。",
                    ],
                ),
            ]
        )
    )
    sections.append({"heading": "7.1.2 匿名评分执行", "level": 3, "paras": []})
    sections.extend(
        model_paragraphs(
            [
                (
                    "匿名评分账号与席位授权",
                    "7.1.2.1",
                    [
                        "新增模型。匿名评分账号承载计划内临时认证凭证及账号席位授权，记录匿名账号、初始密码摘要、账号类型、状态、最终提交时间和领导发放对象。物理表见表7.1-18、表7.1-19。",
                        "账号和授权由系统在计划编译阶段生成，发布或凭证重置时生成一次性交付密码。数据流向匿名登录、评分工作台、凭证重置、最终提交和计划生命周期处理。",
                        "员工匿名账号不保存实际使用人；领导账号关联计划人员并可授权多个领导席位。最终提交成功后账号停用，密码只保存不可逆摘要，凭证只在成功响应中一次性导出。",
                    ],
                ),
                (
                    "评分任务与指标所得分",
                    "7.1.2.2",
                    [
                        "新增模型。评分任务承载一个评分席位对被考核人的完整任务，指标所得分记录各叶子指标的当前分值。评分任务保存草稿版本、任务状态和评分总分。物理表见表7.1-20、表7.1-22。",
                        "任务骨架由发布操作生成；指标分值由匿名评分人填写或按比例分配后保存。数据流向评分工作台、最终提交、评分进度、完成计划和评分汇总明细。",
                        "草稿总分由服务端依据指标所得分重新求和；最终提交统一校验账号全部授权席位后转为正式评分，不允许逐席位提交或撤回。",
                    ],
                ),
                (
                    "评分会话",
                    "7.1.2.3",
                    [
                        "技术支撑模型。评分会话记录匿名账号登录后取得的数据库权威会话，字段包括令牌摘要、签发时间、最后访问时间、空闲过期时间、绝对过期时间和撤销事实。物理表见表7.1-21。",
                        "会话由匿名账号兑换评分凭证成功后创建，原始令牌只通过专用HttpOnly Cookie传递，不进入路径、查询参数或响应业务字段。",
                        "一个账号同时只能有一个有效会话，新登录撤销旧会话；退出、凭证重置、最终提交、计划完成或终止时撤销会话。该模型属于技术控制数据，不单列数据功能。",
                    ],
                ),
            ]
        )
    )
    sections.append({"heading": "7.1.3 评分结果与年度档案", "level": 3, "paras": []})
    sections.extend(
        model_paragraphs(
            [
                (
                    "评分汇总",
                    "7.1.3.1",
                    [
                        "新增模型。评分汇总包含评分组汇总和被考核人汇总分，记录组内平均分、组加权分、个人评分汇总分和计算时间。物理表见表7.1-23、表7.1-24。",
                        "汇总全部由已完成计划的正式评分任务计算生成，不接受人工录入。数据流向汇总等次页面、评分汇总明细和年度考核等次导入。",
                        "系统先计算各评分组六位平均分和加权分，再汇总为两位个人评分汇总分。总监员工组按各分管BU员工均分与一份自评分等权组成。汇总分不直接映射年度考核等次。",
                    ],
                ),
                (
                    "年度考核登记表",
                    "7.1.3.2",
                    [
                        "新增模型。年度考核登记表按员工和年度唯一，承载姓名、现单位、现部门快照、年度主要工作摘要、年度考核等次、登记状态、提交时间和草稿版本。物理表见表7.1-25。",
                        "员工首次保存时从当前实名身份固化人员与部门快照；管理员导入等次时可为不存在档案的被考核人代建。数据流向员工登记表查询、保存、提交、Word导出和等次处理。",
                        "草稿可反复保存；首次提交后员工工作摘要永久只读；管理员仍可通过等次全量导入更正等次。登记表不依赖年度考核计划完成。",
                    ],
                ),
            ]
        )
    )
    sections.append({"heading": "7.1.4 审计与主系统数据", "level": 3, "paras": []})
    sections.extend(
        model_paragraphs(
            [
                (
                    "员工考核操作日志",
                    "7.1.4.1",
                    [
                        "技术支撑模型。员工考核操作日志记录计划、凭证、结果和匿名会话相关的关键操作事实，包含业务对象、操作类型、操作人、前后摘要、来源地址和操作时间。物理表见表7.1-26。",
                        "日志由系统在业务操作过程中自动写入，不接受人工录入；匿名评分安全事件允许操作人为空。日志只追加、不修改、不删除。",
                        "日志用于审计和问题追溯，无独立用户查询或维护入口，按阶段一决定不列为数据功能。",
                    ],
                ),
                (
                    "主系统人员",
                    "7.1.4.2",
                    [
                        "引用模型（边界外数据，本系统只读）。主系统人员由大额成本管控主系统维护，本系统读取员工编号、人员姓名、现单位和人员状态。",
                        "人员数据在计划导入时用于身份校验，在员工登记表操作时用于当前身份和基础资料读取，并在业务发生时固化为计划人员快照或登记表快照。",
                        "本系统不对人员执行新增、修改或删除。计划内快照和登记表快照固化后不再随主系统变化覆盖。",
                    ],
                ),
                (
                    "主系统部门",
                    "7.1.4.3",
                    [
                        "引用模型（边界外数据，本系统只读）。主系统部门由边界外主系统维护；员工登记表读取当前现部门。计划考核部门名称由人员导入文件提供，不要求与主系统部门标识一致。",
                        "主系统部门数据只流向当前员工登记表快照。计划部门快照来自导入文件，并与主系统实时部门变化隔离。",
                        "本系统不对部门执行新增、修改或删除，只在业务发生时读取必要字段并固化快照。",
                    ],
                ),
            ]
        )
    )
    sections.append({"heading": "7.1.5 物理表结构", "level": 3, "paras": []})
    table_groups = [
        ("7.1.5.1 考核配置与计划结构", 0, 14),
        ("7.1.5.2 匿名评分执行", 14, 22),
        ("7.1.5.3 评分结果与年度档案", 22, 25),
        ("7.1.5.4 审计", 25, 26),
    ]
    for heading, start, end in table_groups:
        sections.append(
            {
                "heading": heading,
                "level": 4,
                "paras": [],
                "images": [
                    {
                        "path": str((OUT / "接口表" / rel).resolve()),
                        "caption": "",
                    }
                    for db in dbs[start:end]
                    for rel in db["image_rels"]
                ],
            }
        )
    return sections


def spec_image_entries(specs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    entries = []
    for spec in specs:
        for rel in spec["image_rels"]:
            entries.append(
                {
                    "path": str((OUT / "接口表" / rel).resolve()),
                    "caption": "",
                }
            )
    return entries


def legacy_ui_sections(
    specs: list[dict[str, Any]], masters: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    by_id = {item["id"]: item for item in specs + masters}
    groups = [
        (
            "7.2 界面功能模块",
            "",
            [],
        ),
        (
            "7.2.1 考核模板页面",
            "实名管理端“考核模板”入口，允许超级管理员和员工考核管理员角色使用。页面共同展示模板列表、模板详情和两层指标树；草稿可编辑，发布后只读，已发布模板可停用或复制。",
            [
                ("7.2.1.1 查询考核模板列表", "管理员按模板名称、模板状态组合筛选，系统检索并直接呈现符合条件的模板名称、状态、启用状态、创建时间及结果总数。接口契约见表7-1。", ["template.list"]),
                ("7.2.1.2 创建考核模板", "管理员填写模板名称和说明，系统创建初始为草稿的模板并记录当前管理员审计信息。接口契约见表7-2。", ["template.create"]),
                ("7.2.1.3 更新考核模板基本信息", "管理员修改目标模板的名称和说明，系统校验其处于草稿状态后更新并记录更新人与时间。接口契约见表7-3。", ["template.update"]),
                ("7.2.1.4 查询考核模板详情", "管理员选定目标模板，系统读取并返回模板基本信息和完整两层指标树，不做任何计算或维护。接口契约见表7-4。", ["template.detail"]),
                ("7.2.1.5 配置考核模板指标", "管理员整体提交两层指标结构，系统校验无第三层、无空分组、满分精度合法后替换草稿模板的指标结构。接口契约见表7-5。", ["template.configure-indicators"]),
                ("7.2.1.6 发布考核模板", "管理员对目标草稿模板执行发布，系统校验结构完整且全部叶子满分合计为100后，将模板转为永久只读的已发布状态。接口契约见表7-6。", ["template.publish"]),
                ("7.2.1.7 停用考核模板", "管理员对已发布且启用的模板执行停用，使其仅不再被新计划选择。接口契约见表7-7。", ["template.disable"]),
                ("7.2.1.8 复制考核模板", "管理员指定来源模板和新模板名称，系统复制完整指标并创建相互独立的草稿副本。接口契约见表7-8。", ["template.copy"]),
                ("7.2.1.9 删除草稿考核模板", "管理员删除未发布模板，系统删除模板及其指标。接口契约见表7-9。", ["template.delete"]),
            ],
        ),
        (
            "7.2.2 考核计划列表",
            "实名管理端“考核计划”入口，允许超级管理员和员工考核管理员角色使用。页面按年度和状态呈现计划，提供创建草稿和删除从未发布草稿的入口。",
            [
                ("7.2.2.1 查询年度考核计划列表", "管理员按考核年度、计划状态组合筛选，系统检索并呈现各计划的模板、状态、草稿版本及生命周期时间。接口契约见表7-10。", ["plan.list"]),
                ("7.2.2.2 创建年度考核计划", "管理员选择考核年度、已发布模板和备注，系统校验同年度最多一个未终止计划后创建草稿。接口契约见表7-11。", ["plan.create"]),
                ("7.2.2.3 删除年度考核计划草稿", "管理员删除从未发布的计划，系统在同一事务内删除计划及全部草稿结构。接口契约见表7-12。", ["plan.delete"]),
            ],
        ),
        (
            "7.2.3 计划工作台—概览",
            "计划工作台概览集中呈现计划基础信息、草稿版本、发布就绪结论、评分进度和生命周期操作。按钮随草稿、进行中、已完成、已终止状态变化。",
            [
                ("7.2.3.1 更新年度考核计划草稿", "管理员携带草稿版本修改所选模板和计划备注，系统校验后重新编译生成计划结构。接口契约见表7-13。", ["plan.update-draft"]),
                ("7.2.3.2 核对计划发布就绪状态", "管理员触发就绪核对，系统检查模板、人员、席位、评分关系、权重、账号授权和任务预期数。该查询作为发布门禁支撑，不单列功能点。接口契约见表7-14。", ["plan.readiness"]),
                ("7.2.3.3 发布计划并交付匿名凭证", "管理员执行发布，系统完成并发校验后激活匿名账号、生成评分任务、冻结结构并一次性流式交付凭证。接口契约见表7-15。", ["plan.publish"]),
                ("7.2.3.4 重置未提交匿名凭证", "管理员对进行中计划执行全量凭证重置，系统更新全部未提交账号密码摘要、撤销会话并重新交付凭证。接口契约见表7-16。", ["plan.credentials-reset"]),
                ("7.2.3.5 终止年度考核计划", "管理员对进行中计划执行不可逆终止，系统阻断后续评分并撤销全部会话。接口契约见表7-17。", ["plan.terminate"]),
                ("7.2.3.6 完成计划并固化评分汇总", "全部应有任务正式提交后，管理员执行完成，系统原子计算并写入组级和个人评分汇总。接口契约见表7-18。", ["plan.complete"]),
                ("7.2.3.7 查询计划评分进度", "管理员查询目标计划进度，系统汇总任务和账号提交状态，并派生是否可完成及阻断原因。接口契约见表7-19。", ["plan.rating-progress"]),
            ],
        ),
        (
            "7.2.4 计划工作台—人员操作",
            "人员操作页提供固定表头导入模板和完整人员文件上传。仅草稿计划允许导入，页面显示校验统计、行级错误和替换影响。",
            [
                ("7.2.4.1 获取计划人员导入模板", "系统输出两张固定表头的空Excel模板，不读取持久业务数据。接口契约见表7-20。", ["plan.personnel-import-template"]),
                ("7.2.4.2 核对人员文件", "管理员上传双表人员文件，系统只读核对人员身份、部门职级、权重、主考核和替换影响。接口契约见表7-21。", ["plan.personnel-precheck"]),
                ("7.2.4.3 确认替换人员配置", "管理员以同一文件、预检摘要和草稿版本确认，系统原子替换人员快照并重建未发布结构。接口契约见表7-22。", ["plan.personnel-import"]),
            ],
        ),
        (
            "7.2.5 计划工作台—任职清单",
            "任职清单按计划部门和考核职级展示固化的人员任职快照及被考核任职标记。",
            [
                ("7.2.5.1 查询计划人员任职快照", "管理员按计划部门、考核职级筛选并分页查询，系统只读返回任职快照及受考核任职标识。接口契约见表7-23。", ["plan.personnel"]),
            ],
        ),
        (
            "7.2.6 计划工作台—主考核与分管",
            "主考核与分管页展示文件确定的总监主考核、正副职来源以及当前页领导的最终员工范围。页面可导出完整核对表。",
            [
                ("7.2.6.1 查询主考核与分管", "管理员按领导任职、来源和关键词筛选，系统在同一只读快照中返回总监主考核卡片和当前页领导成员。接口契约见表7-24。", ["plan.roster-workbench"]),
                ("7.2.6.2 保存领导分管范围", "管理员显式选择普通员工成员，系统替换指定领导任职的分管范围并重建结构。接口契约见表7-25。", ["plan.assessment-scope-save"]),
                ("7.2.6.3 导出主考核与分管核对表", "系统按完整一致快照导出人员结构和领导分管两张核对表。接口契约见表7-26。", ["plan.roster-check-export"]),
            ],
        ),
        (
            "7.2.7 计划工作台—结构预览",
            "结构预览按考核部门分页展示被考核人、评分来源、有效权重、缺组转移和预期任务。",
            [
                ("7.2.7.1 查询评分结构", "管理员选择部门并翻页，系统读取本页结构；全计划阻断原因由发布就绪查询返回。接口契约见表7-27。", ["plan.structure-preview"]),
            ],
        ),
        (
            "7.2.8 计划工作台—评分汇总",
            "已完成计划的评分汇总页支持部门、人员关键词、年度考核等次、排序和分页，并提供等次填写表导出和完整文件导入。",
            [
                ("7.2.8.1 查询评分汇总", "管理员按部门、关键词、等次和排序分页查询，系统呈现已固化汇总分、等次和稳定评分组槽位。接口契约见表7-28。", ["plan.rating-results"]),
                ("7.2.8.2 导出年度考核等次填写表", "管理员导出固定六列文件供线下填写；该导出归入汇总和导入流程，不单列功能点。接口契约见表7-29。", ["plan.grade-fill-export"]),
                ("7.2.8.3 导入年度考核等次", "管理员上传完整等次文件，系统全量校验后原子回填或更正登记表等次。接口契约见表7-30。", ["plan.grade-import"]),
            ],
        ),
        (
            "7.2.9 评分汇总审计明细",
            "评分汇总明细展示一名被考核人的完整评分审计链，不披露账号、席位或评分人员身份。",
            [
                ("7.2.9.1 查询个人评分汇总审计明细", "管理员选定被考核对象，系统重算总监员工组组成项并返回匿名化评分明细。接口契约见表7-31。", ["plan.rating-summary-detail"]),
            ],
        ),
        (
            "7.2.10 匿名评分工作台",
            "匿名评分入口独立于实名菜单。匿名账号登录后只能访问其获授权席位和任务；领导融合账号可在多个领导席位之间切换。",
            [
                ("7.2.10.1 匿名账号登录并兑换评分会话", "评分人输入匿名账号和密码，系统校验凭证、建立独立数据库会话并返回工作台初始化信息。接口契约见表7-32。", ["rating.session.exchange"]),
                ("7.2.10.2 恢复匿名评分工作台", "系统从有效评分Cookie恢复账号获授权的任务清单和席位配额边界。接口契约见表7-33。", ["rating.session.current"]),
                ("7.2.10.3 退出匿名评分会话", "系统撤销技术会话并清除Cookie，无令牌时幂等成功；该动作不单列功能点。接口契约见表7-34。", ["rating.session.logout"]),
                ("7.2.10.4 查询评分任务与席位草稿", "评分人进入席位时一次读取全部任务草稿和公共指标，单任务查询用于单行恢复。接口契约见表7-35、表7-36。", ["rating.task.get-draft", "rating.seat.get-drafts"]),
                ("7.2.10.5 按总分分配指标建议分值", "评分人输入期望总分，系统按最大余数法生成各叶子指标建议分，不保存草稿。接口契约见表7-37。", ["rating.task.allocate"]),
                ("7.2.10.6 保存评分任务草稿", "评分人提交完整指标分值集合，系统校验会话、授权、版本和分值合法性后整体保存。接口契约见表7-38。", ["rating.task.save-draft"]),
                ("7.2.10.7 最终提交匿名账号全部评分", "全部授权席位填写完整后，系统统一校验并一次性转为正式评分、停用账号并撤销会话。接口契约见表7-39、表7-40。", ["rating.submission.employee", "rating.submission.leader"]),
            ],
        ),
        (
            "7.2.11 年度考核登记表",
            "员工实名端入口面向全部实名登录员工开放，按年度展示本人登记表。草稿可保存，首次提交后工作摘要只读，并可导出已提交登记表。",
            [
                ("7.2.11.1 查询登记表可用年度", "系统结合当前年度、上一年度和本人历史档案返回可选年度及默认年度。接口契约见表7-41。", ["registration.years"]),
                ("7.2.11.2 查询年度考核登记表", "系统返回本人已有登记表或未落库的临时空表。接口契约见表7-42。", ["registration.detail"]),
                ("7.2.11.3 保存年度考核登记表草稿", "员工按草稿版本保存工作摘要，首次保存固化当前身份和部门快照。接口契约见表7-43。", ["registration.save-draft"]),
                ("7.2.11.4 提交年度考核登记表", "系统校验摘要后首次提交，员工填写内容永久只读。接口契约见表7-44。", ["registration.submit"]),
                ("7.2.11.5 导出年度考核登记表Word", "员工导出已提交登记表，系统替换客户固定Word模板中的身份、摘要和等次标记。接口契约见表7-45。", ["registration.export"]),
            ],
        ),
        (
            "7.2.12 主系统数据接口",
            "本节登记员工考核后端与边界外主系统之间的只读数据契约。",
            [
                ("7.2.12.1 主系统人员与部门共享数据读取", "人员导入前通过共享数据只读校验员工和部门，任一标识无效时整次导入零写入。接口契约见表7-46。", ["master-data.import-validate"]),
                ("7.2.12.2 当前实名人员与部门服务读取", "登记表操作读取当前员工编号、姓名、现单位和现部门，字段缺失时拒绝操作。接口契约见表7-47。", ["master-data.current-employee"]),
            ],
        ),
    ]

    sections: list[dict[str, Any]] = []
    for heading, intro, features in groups:
        level = 1 if heading.startswith("7 ") else (2 if heading.count(".") == 1 else 3)
        sections.append(
            {
                "heading": heading,
                "level": level,
                "paras": [intro] if intro else [],
            }
        )
        for feature_heading, paragraph, api_ids in features:
            sections.append(
                {
                    "heading": feature_heading,
                    "level": 4,
                    "paras": [paragraph],
                    "images": [
                        entry
                        for api_id in api_ids
                        for entry in spec_image_entries([by_id[api_id]])
                    ],
                }
            )
    return sections


def legacy_tail_sections() -> list[dict[str, Any]]:
    return [
        {
            "heading": "8 界面需求",
            "level": 1,
            "paras": [],
        },
        {
            "heading": "8.1 角色与入口隔离",
            "level": 2,
            "paras": [
                "管理端计划和模板入口只对超级管理员和员工考核管理员开放；年度考核登记表对全部实名登录员工开放且限定本人数据；匿名评分入口 `/rating` 路由独立于实名菜单和实名登录状态。管理员也可填写本人登记表，但不获得他人登记表访问权。",
            ],
        },
        {
            "heading": "8.2 状态与只读反馈",
            "level": 2,
            "paras": [
                "模板和计划页面应依据当前生命周期状态显示可执行操作。发布后的模板、已发布计划评分结构、已完成评分记录以及员工已提交的年度主要工作摘要均以只读方式呈现，不提供回退或撤回入口。",
            ],
        },
        {
            "heading": "8.3 并发冲突反馈",
            "level": 2,
            "paras": [
                "计划草稿、评分草稿和登记表草稿写入均携带客户端读取的版本。版本不一致时，系统拒绝旧请求并提示用户刷新后重试，不覆盖较新数据。",
            ],
        },
        {
            "heading": "8.4 匿名信息保护",
            "level": 2,
            "paras": [
                "匿名评分页面、地址栏和浏览器可读存储不得展示会话令牌或明文密码。评分汇总审计明细只使用评分1、评分2等组内匿名序号，不展示账号、评分席位或评分人员身份。",
            ],
        },
        {
            "heading": "8.5 文件交互反馈",
            "level": 2,
            "paras": [
                "人员任职和年度考核等次文件校验失败时，页面显示文件级或行级错误，且不得呈现部分成功。匿名凭证只在发布或全量重置成功响应中下载；下载失败后页面不得盲目重复发布。",
            ],
        },
        {
            "heading": "9 非功能性需求",
            "level": 1,
            "paras": [],
        },
        {
            "heading": "9.1 访问控制与会话安全",
            "level": 2,
            "paras": [
                "实名管理接口采用超级管理员或员工考核管理员角色控制，员工登记表接口要求实名认证并由当前登录身份限定本人数据。匿名评分使用独立数据库权威会话，原始会话令牌仅通过专用HttpOnly、SameSite=Strict Cookie传递，生产配置启用Secure；匿名写请求执行同源校验。",
            ],
        },
        {
            "heading": "9.2 凭证与敏感信息保护",
            "level": 2,
            "paras": [
                "匿名密码只保存不可逆摘要。初始密码和重置密码仅在当次成功响应的内存中用于生成凭证文件，响应结束后清除；系统不得将明文或可逆密文写入数据库、缓存、临时文件或日志。员工强匿名账号不保存实际使用人归属。",
            ],
        },
        {
            "heading": "9.3 审计可追溯性",
            "level": 2,
            "paras": [
                "发布、凭证重置、终止、完成、等次导入、草稿删除和匿名会话安全事件应留下不可由普通业务操作修改的审计事实。操作日志只追加，不参与功能点计数。",
            ],
        },
        {"heading": "10 其它需求（可选）", "level": 1, "paras": ["无"]},
    ]


def build_document_model(
    model: dict[str, Any],
    rows: list[dict[str, Any]],
    mappings: list[dict[str, Any]],
    specs: list[dict[str, Any]],
    dbs: list[dict[str, Any]],
    masters: list[dict[str, Any]],
) -> dict[str, Any]:
    sections = []
    sections.extend(legacy_front_sections(model))
    sections.extend(legacy_data_sections(dbs))
    sections.extend(legacy_ui_sections(specs, masters))
    sections.extend(legacy_tail_sections())
    return {
        "title": "员工考核系统改造",
        "subtitle": "需求说明书",
        "version": "v3-new-v2",
        "date": "2026年10月",
        "counts": {
            "ilf": 11,
            "elf": 2,
            "ei": 22,
            "eo": 10,
            "eq": 5,
            "ufp": 267,
        },
        "sections": sections,
        "mappings": mappings,
        "features": rows,
    }


def main() -> None:
    model = load_json()
    specs = enrich_endpoints(endpoint_specs())
    dbs = parse_db_tables()
    rows = feature_rows(model, specs)
    mappings = mapping_rows(model, rows, specs, dbs)
    write_xlsx_outputs(model, rows, mappings)
    masters = write_yaml(model, specs, dbs)
    render_cover_backgrounds()
    create_diagram()
    render_images(specs, dbs, masters)
    doc_model = build_document_model(model, rows, mappings, specs, dbs, masters)
    payload = {
        "model": model,
        "specs": specs,
        "dbs": dbs,
        "masters": masters,
        "features": rows,
        "mappings": mappings,
        "document": doc_model,
    }
    (WORK / "delivery-model.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "out": str(OUT),
                "apis": len(specs),
                "db_tables": len(dbs),
                "features": len(rows),
                "mappings": len(mappings),
                "images": len(list((OUT / "接口表/images").glob("*.png"))),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
