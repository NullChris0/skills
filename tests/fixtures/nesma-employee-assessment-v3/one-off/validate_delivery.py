#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import yaml
from openpyxl import load_workbook
from lxml import etree


OUT = Path("/home/nc/projects/work/hx-employee-assessment-delivery-v3-new-v2")
WORK = Path("/tmp/hx-delivery-work")
BACKEND = WORK / "backend-baseline/hx_assessment"
COUNT = (
    Path("/home/nc/projects/work/cuc-ah-ucac-pso-hx-repository")
    / "doc/hx-assessment/NESMA估算/员工年度考核-v3-new-v2"
)


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def load_model() -> dict:
    return json.loads((COUNT / "计数模型.json").read_text(encoding="utf-8"))


def check_counts(model: dict) -> dict:
    dfs = model["data_functions"]
    tfs = model["transactions"]
    assert len(dfs) == 13
    assert len(tfs) == 37
    assert sum(x["type"] == "ILF" for x in dfs) == 11
    assert sum(x["type"] == "ELF" for x in dfs) == 2
    assert sum(x["type"] == "EI" for x in tfs) == 22
    assert sum(x["type"] == "EO" for x in tfs) == 10
    assert sum(x["type"] == "EQ" for x in tfs) == 5
    dfp = sum(x["fp"] for x in dfs)
    tfp = sum(x["fp"] for x in tfs)
    assert (dfp, tfp, dfp + tfp) == (87, 180, 267)
    assert "assumptions" not in model
    return {
        "data_functions": len(dfs),
        "transactions": len(tfs),
        "data_fp": dfp,
        "transaction_fp": tfp,
        "ufp": dfp + tfp,
    }


def check_mapping(model: dict) -> dict:
    wb = load_workbook(OUT / "映射表.xlsx", data_only=True)
    ws = wb["NESMA映射"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 50, len(rows)
    ids = [row[0] for row in rows]
    expected = [x["id"] for x in model["data_functions"]] + [
        x["id"] for x in model["transactions"]
    ]
    assert ids == expected
    assert len(set(ids)) == 50
    assert all(row[7] == "已覆盖" for row in rows)
    assert all(row[4] for row in rows), "功能清单行映射不能为空"
    assert all(row[5] for row in rows), "章节/页面/数据模型映射不能为空"

    summary = wb["覆盖汇总"]
    values = {row[0]: row[1] for row in summary.iter_rows(min_row=2, values_only=True)}
    assert values["数据功能"] == 13
    assert values["事务功能"] == 37
    assert values["UFP总数"] == 267
    assert values["未覆盖"] == 0
    return {"rows": len(rows), "unique": len(set(ids))}


def check_features(model: dict) -> dict:
    wb = load_workbook(OUT / "功能清单.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 37
    expected_cols = [
        "序号",
        "系统名称",
        "一级模块",
        "二级模块",
        "功能需求描述",
        "外部接口需求描述",
        "功能（新增/修改）",
        "备注",
    ]
    assert [x.value for x in ws[1]] == expected_cols
    remarks = "\n".join(str(row[7]) for row in rows)
    for item in model["transactions"]:
        assert remarks.count(item["id"]) == 1, item["id"]
    last_system = ""
    last_module = ""
    for row in rows:
        last_system = row[1] or last_system
        last_module = row[2] or last_module
        assert last_system and last_module and row[3] and row[4]
    return {"rows": len(rows), "tf_mentions": 37}


def baseline_endpoints() -> tuple[set[tuple[str, str, str]], list[dict]]:
    endpoint_set: set[tuple[str, str, str]] = set()
    records: list[dict] = []
    controller_root = BACKEND / "src/main/java/com/asiainfo/ea/controller"
    for path in sorted(controller_root.rglob("*Controller.java")):
        source = path.read_text(encoding="utf-8")
        class_match = re.search(r"public\s+class\s+(\w+)", source)
        if not class_match:
            fail(f"无法解析控制器类名：{path}")
        class_name = class_match.group(1)
        class_map = re.search(r'@RequestMapping\("([^"]+)"\)', source)
        if not class_map:
            fail(f"无法解析控制器前缀：{path}")
        base = class_map.group(1).rstrip("/")
        pattern = re.compile(
            r"@(Get|Post|Put|Delete)Mapping"
            r"(?:\s*\((.*?)\))?\s+"
            r"public\s+[^{;]+?\s+(\w+)\s*\(",
            re.S,
        )
        for match in pattern.finditer(source):
            verb = match.group(1).upper()
            annotation_args = match.group(2) or ""
            method = match.group(3)
            path_match = re.search(r'"([^"]*)"', annotation_args)
            method_path = path_match.group(1) if path_match else ""
            full_path = base + method_path
            full_path = "/" + full_path.strip("/")
            endpoint_set.add((verb, full_path, f"{class_name}#{method}"))
            records.append(
                {
                    "method": verb,
                    "path": full_path,
                    "controller": f"{class_name}#{method}",
                }
            )
    return endpoint_set, records


def all_yaml_records() -> list[dict]:
    records = []
    for path in sorted((OUT / "接口表").rglob("*.yaml")):
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        data["_path"] = str(path.relative_to(OUT))
        records.append(data)
    return records


def check_yaml(model: dict) -> dict:
    records = all_yaml_records()
    http = [x for x in records if x.get("kind") == "http"]
    elf = [x for x in records if x.get("kind") == "elf"]
    db = [x for x in records if "columns" in x]
    assert len(http) == 45, len(http)
    assert len(elf) == 2
    assert len(db) == 26
    assert len({x["id"] for x in records}) == len(records)

    baseline_set, _ = baseline_endpoints()
    yaml_set = {
        (x["method"], x["path"], x["controller"])
        for x in http
    }
    assert yaml_set == baseline_set, {
        "missing": sorted(baseline_set - yaml_set),
        "extra": sorted(yaml_set - baseline_set),
    }

    valid_ids = {
        x["id"] for x in model["data_functions"] + model["transactions"]
    }
    uncounted = 0
    for rec in http + elf:
        refs = rec.get("nesma")
        refs = refs if isinstance(refs, list) else [refs]
        for ref in refs:
            if ref == "uncounted":
                uncounted += 1
            elif ref not in valid_ids:
                fail(f"{rec['_path']} 引用了无效 NESMA：{ref}")
        if rec.get("data_java") and not rec.get("data"):
            fail(f"{rec['_path']} 声明了 data_java 但 data 为空")
        if rec.get("body_java") and not rec.get("body"):
            fail(f"{rec['_path']} 声明了 body_java 但 body 为空")
    session_exchange = next(
        x for x in http if x["id"] == "rating.session.exchange"
    )
    assert "sessionToken" not in {
        item["name"] for item in session_exchange.get("data", [])
    }

    image_dir = OUT / "接口表/images"
    assert len(list(image_dir.glob("data-*.png"))) == 26
    assert len(list(image_dir.glob("api-*.png"))) == 47
    return {
        "http": len(http),
        "elf": len(elf),
        "db": len(db),
        "uncounted": uncounted,
        "images": len(list(image_dir.glob("*.png"))),
    }


def docx_text(path: Path) -> str:
    result = subprocess.run(
        ["pandoc", str(path), "-t", "plain"],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    return result.stdout


def check_docx(model: dict, yaml_result: dict) -> dict:
    path = OUT / "需求说明书.docx"
    assert path.exists() and path.stat().st_size > 1_000_000
    with ZipFile(path) as z:
        names = z.namelist()
        assert "word/document.xml" in names
        assert "word/styles.xml" in names
        assert "word/settings.xml" in names
        xml = z.read("word/document.xml").decode("utf-8")
        media = [name for name in names if name.startswith("word/media/")]
        assert len(media) == 74, len(media)
        assert "<w:toc" not in xml  # TOC is a field, not a literal string.
        assert " TOC \\o " in xml
        for text in ["图 1：系统总体功能结构", "表7-1", "表7-47", "表7.1-1", "表7.1-26"]:
            assert text in xml, text

    tree = etree.fromstring(xml.encode("utf-8"))
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }
    text = "\n".join(tree.xpath("//w:t/text()", namespaces=ns))
    prohibited_patterns = [
        r"(?<!\d)256 UFP",
        r"(?<!\d)256",
        r"(?<!\d)81 FP",
        r"(?<!\d)175 FP",
        r"(?<!\d)80 FP",
        r"(?<!\d)187 FP",
        r"assumptions",
        r"\[Company Name\]",
        r"TODO",
        r"文档控制",
        r"附录 A NESMA",
    ]
    for pattern in prohibited_patterns:
        if re.search(pattern, text):
            fail(f"文档残留禁止内容：{pattern}")

    required_headings = [
        "1 项目概述",
        "1.1 项目背景",
        "1.2 系统总体功能结构",
        "1.3 往期建设情况",
        "1.4 本项目预算工作量",
        "1.5 本项目投资计划预算",
        "2 文档目的",
        "3 文档范围",
        "4 术语和缩写解释",
        "5 参考文档",
        "6 范围定义",
        "7 功能需求",
        "7.1 数据模型",
        "7.1.5 物理表结构",
        "7.2 界面功能模块",
        "7.2.1.1 查询考核模板列表",
        "7.2.10.7 最终提交匿名账号全部评分",
        "7.2.12.2 当前实名人员与部门服务读取",
        "8 界面需求",
        "9 非功能性需求",
        "10 其它需求（可选）",
    ]
    for heading in required_headings:
        assert heading in text, f"缺少旧稿章节：{heading}"
    for table in range(1, 48):
        assert f"表7-{table}" in text, f"缺少接口表引用 7-{table}"
    assert len(list((OUT / "接口表/images").glob("data-*.png"))) == 26
    assert "第 页 / 共" not in text
    assert "v3-new-v2" not in text

    return {
        "chars": len(text.replace("\n", "").replace(" ", "")),
        "media": len(media),
        "interface_tables": 47,
        "tables_7": 26,
    }


def main() -> None:
    model = load_model()
    result = {
        "counts": check_counts(model),
        "mapping": check_mapping(model),
        "features": check_features(model),
        "yaml": check_yaml(model),
    }
    result["docx"] = check_docx(model, result["yaml"])
    result["status"] = "PASS"
    (WORK / "validation.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
